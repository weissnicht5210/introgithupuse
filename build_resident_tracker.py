"""전공의·인턴 파견 모니터링 엑셀 (Excel 2019 이상 호환 — Microsoft 365에서도 동작)

- 입력 시트는 엑셀 '표'(Table): 아래에 붙여넣으면 표와 계산 열이 자동으로 늘어난다(행 수 제한 없음).
- 출력 시트는 정해진 행 수만큼 미리 수식을 넣은 목록(INDEX/MATCH): 표시 한도를 넘으면 경고가 뜬다.
- Excel 2019에 없는 함수(FILTER·SORT·UNIQUE·SEQUENCE·XLOOKUP·XMATCH·LET)는 쓰지 않는다.
- DATA=<pickle> 로 실행하면 기존 파일에서 옮긴 입력 자료로 채운다(실제 자료는 저장소에 넣지 않음).
"""
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill

from xlsx2019 import (CI, CL, FONT, add_name, border, center, cse, dv, f_base, f_bold, f_input, f_note, f_title,
                      fill_aux, fill_key, head_cell, make_table, q, save)

DATA = None  # 실제 자료로 채울 때: DATA=<pickle 경로> (입력 시트 값만 담은 사전)
if os.environ.get("DATA"):
    import pickle
    with open(os.environ["DATA"], "rb") as fh:
        DATA = pickle.load(fh)
OUT = os.environ.get("OUT", "전공의_파견_모니터링.xlsx")
MAXR = 1048576          # 서식·검증 적용 끝 행
GURI_ROWS_END = 20000   # 구리 근무자명단에서 '근무자 확정' 여부를 볼 끝 행
GURI_COLS_END = "ZZ"    # 구리 근무자명단에서 읽는 끝 열 (월 블록 약 116개월)
OUTN = int(os.environ.get("OUTN", "1000"))   # 명단 기반 출력 목록 표시 한도(명)
GR = int(os.environ.get("GR", "1500"))       # 구리 근무자명단 원본 해석 행 수(4행부터)
GV = int(os.environ.get("GV", "500"))        # 구리 근무자 목록 표시 한도(명)
IVN = int(os.environ.get("IVN", "600"))      # 인턴 현황 표시 한도(명)
CODEN = 40                                   # 코드 목록(병원·진료과 등) 표시 한도

# ---------------------------------------------------------------- 코드 목록
HOSPITALS = ["서울병원", "인천병원", "구리병원", "창원병원", "제주병원", "오산병원"]
DEPTS = ["내과", "소아청소년과", "신경과", "정신건강의학과", "피부과", "외과", "심장혈관흉부외과",
         "정형외과", "신경외과", "성형외과", "산부인과", "안과", "이비인후과", "비뇨의학과",
         "재활의학과", "마취통증의학과", "영상의학과", "진단검사의학과", "병리과", "가정의학과",
         "응급의학과", "핵의학과", "직업환경의학과", "수련교육부(인턴)"]
INTERN_ONLY = ["통합", "핵/방", "직/피", "정/영"]
DEPT_ALIASES = [  # (약칭, 정식 명칭) — 사용자 제공 목록
    ("내과", "내과"), ("소아청소년과", "소아청소년과"), ("신경과", "신경과"), ("정신건강의학과", "정신건강의학과"),
    ("피부과", "피부과"), ("외과", "외과"), ("심장혈관흉부외과", "심장혈관흉부외과"), ("정형외과", "정형외과"),
    ("신경외과", "신경외과"), ("성형외과", "성형외과"), ("산부인과", "산부인과"), ("안과", "안과"),
    ("이비인후과", "이비인후과"), ("비뇨의학과", "비뇨의학과"), ("재활의학과", "재활의학과"),
    ("마취통증의학과", "마취통증의학과"), ("영상의학과", "영상의학과"), ("진단검사의학과", "진단검사의학과"),
    ("병리과", "병리과"), ("가정의학과", "가정의학과"), ("응급의학과", "응급의학과"), ("핵의학과", "핵의학과"),
    ("직업환경의학과", "직업환경의학과"), ("수련교육부(인턴)", "수련교육부(인턴)"),
    ("소아과", "소아청소년과"), ("소아", "소아청소년과"), ("소청", "소아청소년과"), ("소청과", "소아청소년과"),
    ("정신과", "정신건강의학과"), ("정신", "정신건강의학과"), ("흉부외과", "심장혈관흉부외과"),
    ("흉외", "심장혈관흉부외과"), ("심장혈관흉부", "심장혈관흉부외과"), ("정형", "정형외과"), ("신외", "신경외과"),
    ("성형", "성형외과"), ("산부", "산부인과"), ("이비", "이비인후과"), ("이비인후", "이비인후과"),
    ("비뇨", "비뇨의학과"), ("비뇨기과", "비뇨의학과"), ("재활", "재활의학과"), ("마취", "마취통증의학과"),
    ("마통", "마취통증의학과"), ("영상", "영상의학과"), ("진검", "진단검사의학과"), ("진단검사", "진단검사의학과"),
    ("병리", "병리과"), ("가정", "가정의학과"), ("가정의학", "가정의학과"), ("응급", "응급의학과"),
    ("핵의학", "핵의학과"), ("핵", "핵의학과"), ("직업환경", "직업환경의학과"), ("직환", "직업환경의학과"),
    ("수련교육부", "수련교육부(인턴)"), ("인턴", "수련교육부(인턴)"), ("직피", "직/피"), ("직환/피부", "직/피"),
    ("정영", "정/영"), ("정신/영상", "정/영"), ("핵방", "핵/방"), ("핵/방", "핵/방"), ("통합", "통합"),
]
HOSP_ALIASES = [(h, h) for h in HOSPITALS] + [(h[:-2], h) for h in HOSPITALS] + [
    ("한양대학교병원", "서울병원"), ("한양대병원", "서울병원"), ("한양대서울병원", "서울병원"),
    ("한양대구리병원", "구리병원"), ("에스중앙병원", "제주병원"), ("에스중앙", "제주병원"),
]

wb = Workbook()
wb.remove(wb.active)
MD = "DAY(EOMONTH(급여기준월,0))"

# ================================================================ 설정_코드표
ref = wb.create_sheet("설정_코드표")
if DATA:
    HOSPITALS, DEPTS, INTERN_ONLY = DATA["codes"]["병원표"], DATA["codes"]["진료과표"], DATA["codes"]["인턴구분표"]
    DEPT_ALIASES, HOSP_ALIASES = DATA["codes"]["과약칭표"], DATA["codes"]["병원약칭표"]
code_tables = [
    ("병원표", "A", [("병원", 14)], [{"병원": h} for h in HOSPITALS]),
    ("진료과표", "C", [("근무과", 18)], [{"근무과": d} for d in DEPTS]),
    ("인턴구분표", "E", [("인턴 구분", 12)], [{"인턴 구분": d} for d in INTERN_ONLY]),
    ("과약칭표", "I", [("약칭", 16), ("정식 명칭", 18)], [{"약칭": a, "정식 명칭": b} for a, b in DEPT_ALIASES]),
    ("병원약칭표", "L", [("표기", 16), ("정식 병원", 12)], [{"표기": a, "정식 병원": b} for a, b in HOSP_ALIASES]),
]
from openpyxl.utils import column_index_from_string as CI
for tname, col, cdefs, rows in code_tables:
    make_table(ref, tname, 1, CI(col), [(h, w, "in", None, None) for h, w in cdefs], rows, style="TableStyleLight1")
for c in "BDFGHKN":
    ref.column_dimensions[c].width = 3
ref.column_dimensions["O"].width = 26
ref.column_dimensions["P"].width = 14
notes = [
    "※ 모든 목록은 엑셀 '표'입니다. 표 바로 아래 행에 입력하면 자동으로 목록이 늘어나고 드롭다운·판정에 바로 반영됩니다.",
    "※ 병원약칭표(L~M): 인턴근무계획의 지역 단어와 기본자료의 근무지를 정식 병원으로 바꿉니다. 실제 기관명은 여기에 추가하세요.",
    "※ 과약칭표(I~J): 인턴근무계획 월 칸의 마지막 단어(진료과 약칭)를 정식 명칭으로 바꿉니다.",
    "※ 인턴구분표(E): 통합(과 미지정), 핵/방, 직/피, 정/영처럼 인턴 순환에만 쓰는 구분입니다.",
]
for i, t in enumerate(notes):
    ref[f"O{1 + i}"].value, ref[f"O{1 + i}"].font = t, f_base
head_cell(ref["O8"], "설정값 (노란 칸만 수정)")
head_cell(ref["P8"], "값")
for rr, (label, val, fmt, nm) in enumerate([
        ("급여 기준월(1일)", DATA["settings"][0] if DATA else "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "yyyy-mm-dd", "급여기준월"),
        ("서울·구리 15일 기준(일)", DATA["settings"][1] if DATA else 15, "0", "판정기준일")], 9):
    ref[f"O{rr}"].value, ref[f"O{rr}"].font, ref[f"O{rr}"].border = label, f_bold, border
    c = ref[f"P{rr}"]
    c.value, c.number_format, c.font, c.fill, c.border, c.alignment = val, fmt, f_input, fill_key, border, center
    add_name(wb, nm, f"{q('설정_코드표')}!$P${rr}")
ref["O11"].value, ref["O12"].font = "※ 급여 기준월: 다른 달을 보려면 그 달 1일 날짜를 입력(기본 수식 = 이번 달 1일).", f_note
for nm, r in [("병원목록", "병원표[병원]"), ("진료과목록", "진료과표[근무과]"),
              ("병원표기목록", "병원약칭표[표기]"), ("과약칭목록", "과약칭표[약칭]")]:
    add_name(wb, nm, r)

# ================================================================ 입력_기본자료
bs = wb.create_sheet("입력_기본자료")
base_rows = [
    ("재직", "2261212", "김ㅇㅇ", "가정의학과", "레지던트1", "서울", 20260301, "11-1111", 111111, "여", 19999999),
    ("재직", "2261213", "이ㅇㅇ", "내과", "레지던트2", "서울", 20250301, "11-1112", 111112, "남", 19950702),
    ("재직", "2261214", "박ㅇㅇ", "내과", "레지던트3", "서울", 20240301, "21-1111", 111113, "여", 19970125),
    ("재직", "2261215", "최ㅇㅇ", "산부인과", "레지던트3", "서울", 20240301, "11-1114", 111114, "남", 19961109),
    ("재직", "2261216", "정ㅇㅇ", "마취통증의학과", "레지던트2", "구리", 20250301, "11-1115", 111115, "여", 19980530),
    ("재직", "222222", "김ㅇㅇ", "인턴", "인턴", "서울", 20260301, "22-2222", 222222, "여", 19990101),
    ("재직", "222223", "이ㅇㅇ", "인턴", "인턴", "서울", 20260301, "22-2223", 222223, "남", 19990202),
    ("재직", "222224", "박ㅇㅇ", "인턴", "인턴", "구리", 20260301, "22-2224", 222224, "남", 19990303),
]
if DATA:
    base_rows = DATA["base_rows"]
bh = ["현재근무 여부", "사번", "이름", "수련과목", "연차", "근무지", "수련개시일", "전공의등록번호", "의사면허번호", "성별", "생년월일", "비고"]
bw = [11, 11, 10, 14, 11, 9, 12, 14, 12, 6, 12, 16]
bcols = [(h, w, "in", None, "@" if h in ("사번", "전공의등록번호") else None) for h, w in zip(bh, bw)]
bcols += [("사번키", 12, "calc", '=IF(TRIM([@사번]&"")="","","S|"&TRIM([@사번]&""))', None)]
make_table(bs, "기본자료표", 1, 1, bcols, [dict(zip(bh, r)) for r in base_rows])
bs.freeze_panes = "A2"

# ================================================================ 입력_인턴근무계획
ip = wb.create_sheet("입력_인턴근무계획")
ip.column_dimensions["A"].width = 2
ip.merge_cells("B1:K1")
ip["B1"].value, ip["B1"].font, ip["B1"].alignment = (DATA["plan_title"] if DATA else "2026년도 하반기 인턴근무계획표(9~2월)"), f_title, center
plan_rows = [
    (1, "김ㅇㅇ", 222222, "구리 안과", "구리 내과", "서울 핵방", "서울 재활의학과", "제주 산부", "서울 소아과", 2026.03),
    (2, "이ㅇㅇ", 222223, "서울 내과", "서울 외과", "구리 통합", "인천 산부인과", "제주 에스중앙 내과", "서울 가정의학과", 2026.03),
    (3, "박ㅇㅇ", 222224, "제주 외과", "서울 내과", "서울 응급", "구리 직피", "서울 정영", "제주 산부", 2026.03),
]
ph = ["순번", "이름", "사번", "9월", "10월", "11월", "12월", "1월", "2월", "비고"]
if DATA:
    ph, plan_rows = DATA["plan_headers"], DATA["plan_rows"]
pcols = [(h, w, "in", None, None) for h, w in zip(ph, [6, 10, 11, 17, 17, 17, 17, 17, 17, 10])]
pcols.append(("사번키", 11, "calc", '=IF(TRIM([@사번]&"")="","","S|"&TRIM([@사번]&""))', None))
for k in range(1, 7):
    r = f'TRIM(INDEX(인턴계획표[#This Row],1,{3 + k})&"")'
    t1 = f'LEFT({r},FIND(" ",{r})-1)'
    t2 = f'TRIM(RIGHT(SUBSTITUTE({r}," ",REPT(" ",60)),60))'
    pcols.append((f"정규화{k}", 17, "calc",
                  f'=IF({r}="","",IFERROR(INDEX(병원약칭표[정식 병원],MATCH({t1},병원약칭표[표기],0))&" "&'
                  f'INDEX(과약칭표[정식 명칭],MATCH({t2},과약칭표[약칭],0)),"⚠"&{r}))', None))
kk = 'INDEX(인턴계획표[#This Row],1,10)&""'
vv, mm = f"VALUE({kk})", f"ROUND((VALUE({kk})-INT(VALUE({kk})))*100,0)"
pcols.append(("입사일(비고)", 12, "calc",
              f'=IF({kk}="","",IFERROR(IF(AND(INT({vv})>=1900,{mm}>=1,{mm}<=12),DATE(INT({vv}),{mm},1),{kk}),{kk}))', "yyyy-mm-dd"))
pcols.append(("표시순번", 7, "calc",
              '=IF(TRIM([@이름]&"")="","",SUMPRODUCT(--(TRIM(INDEX(인턴계획표[이름],1):[@이름]&"")<>"")))', "0"))
pcols.append(("입력오류", 8, "calc", '=COUNTIF(인턴계획표[[#This Row],[정규화1]:[정규화6]],"⚠*")', "0"))
make_table(ip, "인턴계획표", 2, 2, pcols, [dict(zip(ph, r)) for r in plan_rows])
ip.freeze_panes = "E3"
t = 'TRIM(E3)'
t1 = f'LEFT({t},FIND(" ",{t})-1)'
dv(ip, f"E3:J{MAXR}",
   f'OR({t}="",AND(ISNUMBER(FIND(" ",{t})),COUNTIF(병원표기목록,{t1})>0,'
   f'COUNTIF(과약칭목록,TRIM(RIGHT(SUBSTITUTE({t}," ",REPT(" ",60)),60)))>0))',
   "등록되지 않은 병원/진료과", "'지역 진료과' 형식(예: 서울 내과, 제주 산부)으로 입력하세요.\n"
   "처음 보는 지역이나 진료과(약칭)라면 먼저 '설정_코드표'에 추가한 뒤 다시 입력하세요.", kind="custom")
ip.conditional_formatting.add(f"E3:J{MAXR}", FormulaRule(formula=['COUNTIF($M3:$R3,"⚠"&TRIM(E3))>0'],
                              fill=PatternFill("solid", bgColor="FFC7CE")))
ip.conditional_formatting.add(f"M3:R{MAXR}", FormulaRule(formula=['LEFT(M3,1)="⚠"'], font=Font(name=FONT, color="C00000", bold=True)))

# ================================================================ 입력_구리근무자명단 (양식 그대로 붙여넣기)
gi = wb.create_sheet("입력_구리근무자명단")
gi["A1"].value, gi["A1"].font = "전공의(모.자 파견) 근무자 명단", f_title
gi["A2"].value, gi["A2"].font = "2026년도", f_bold
g_heads = ["소속", "년차"]
for y, m in [(2026, m) for m in range(9, 13)] + [(2027, m) for m in range(1, 9)]:
    g_heads += [f"{y}년{m}월\n근무예정자", f"{m}월\n근무자", "근무 시작일자\n(시작일자 작성)", "근무 종료일자\n(종료일자 작성)",
                f"{m}월 휴가 스케줄\n(휴가일 작성)", "비고"]
if DATA:
    gi["A1"].value = gi["A2"].value = None
    for (r, c), v in DATA["guri_cells"].items():
        cell = gi.cell(row=r, column=c, value=v)
        cell.font = f_input if r >= 4 else (f_title if r == 1 else f_bold)
    g_heads = [DATA["guri_cells"].get((3, c)) for c in range(1, 1 + max(c for (r, c) in DATA["guri_cells"] if r == 3))]
for i, h in enumerate(g_heads, 1):
    head_cell(gi.cell(row=3, column=i), h)
    gi.column_dimensions[CL(i)].width = 11
gi.column_dimensions["A"].width = 14
gi.column_dimensions["B"].width = 6
gi.row_dimensions[3].height = 44
g_sample = [
    ("내과", 3, ["박ㅇㅇ", None, None, None, None, None], ["박ㅇㅇ", "박ㅇㅇ", "2026.10.01", "2026.10.31", None, None]),
    ("가정의학과", 1, [None] * 6, ["김ㅇㅇ", "김ㅇㅇ", "2026.10.21", "2026.10.31", None, None]),
    ("마취통증의학과", 2, ["정ㅇㅇ", None, None, None, None, "구리독자"], ["정ㅇㅇ", "정ㅇㅇ", "2026.10.01", "2026.10.17", None, "구리독자"]),
    ("외과", 1, ["홍ㅇㅇ", None, None, None, None, "구리독자"], ["홍ㅇㅇ", "홍ㅇㅇ", "2026.10.01", "2026.10.31", None, "구리독자"]),
    ("수련교육부", "인턴", [None] * 6, ["김ㅇㅇ", "김ㅇㅇ", "2026.10.01", "2026.10.31", None, None]),
]
for k, (dept, yr, sep, octb) in enumerate([] if DATA else g_sample):
    r = 4 + k
    gi.cell(row=r, column=1, value=dept).font = f_input
    gi.cell(row=r, column=2, value=yr).font = f_input
    for j, v in enumerate(sep + octb):
        if v is not None:
            gi.cell(row=r, column=3 + j, value=v).font = f_input
gi.freeze_panes = "C4"

# ================================================================ 입력_전공의명단
ws = wb.create_sheet("입력_전공의명단")
ROW_ON = 'TRIM([@사번]&"")<>""'


def lk(col):
    return f'=IF([@기본자료행]="","",IF(INDEX(기본자료표[{col}],[@기본자료행])="","",INDEX(기본자료표[{col}],[@기본자료행])))'


def lk_date(col):
    x = f"INDEX(기본자료표[{col}],[@기본자료행])"
    return (f'=IF([@기본자료행]="","",IF({x}="","",IF(NOT(ISNUMBER({x})),{x}&"",IF({x}<1000000,{x},'
            f'IF(AND({x}>=10000101,INT(MOD({x},10000)/100)>=1,INT(MOD({x},10000)/100)<=12,MOD({x},100)>=1,MOD({x},100)<=31),'
            f'DATE(INT({x}/10000),INT(MOD({x},10000)/100),MOD({x},100)),{x}&"")))))')


home_n = 'IF([@근무지]="","",IFERROR(INDEX(병원약칭표[정식 병원],MATCH(TRIM([@근무지]&""),병원약칭표[표기],0))&"",[@근무지]&""))'
ok_now = 'AND([@인턴 이번달]<>"",LEFT([@인턴 이번달],1)<>"⚠")'
ok_mon = 'AND([@인턴 기준월]<>"",LEFT([@인턴 기준월],1)<>"⚠")'
seoul, guri, th = "[@서울 근무일(기준월)]", "[@구리 근무일(기준월)]", "판정기준일"
mcols = [
    ("번호", 6, "calc", '=ROW()-ROW(명단표[#Headers])', None),
    ("사번", 11, "in", None, "@"),
    ("이름", 9, "calc", lk("이름"), None),
    ("성별", 5, "calc", lk("성별"), None),
    ("생년월일", 11, "calc", lk_date("생년월일"), "yyyy-mm-dd"),
    ("수련과목", 13, "calc", lk("수련과목"), None),
    ("연차", 10, "calc", lk("연차"), None),
    ("근무지", 8, "calc", lk("근무지"), None),
    ("수련개시일", 11, "calc", lk_date("수련개시일"), "yyyy-mm-dd"),
    ("전공의등록번호", 12, "calc", lk("전공의등록번호"), None),
    ("의사면허번호", 11, "calc", lk("의사면허번호"), None),
    ("재직여부", 8, "calc", lk("현재근무 여부"), None),
    ("입사장소", 10, "in", None, None),
    ("현재 근무병원", 11, "calc",
     f'=IF(NOT({ROW_ON}),"",IF({ok_now},LEFT([@인턴 이번달],FIND(" ",[@인턴 이번달])-1),'
     f'IF([@구리 근무 중]="Y","구리병원",{home_n})))', None),
    ("현재 진료과", 14, "calc",
     f'=IF(NOT({ROW_ON}),"",IF({ok_now},MID([@인턴 이번달],FIND(" ",[@인턴 이번달])+1,60),'
     f'IF([@수련과목]="","",IF([@수련과목]="인턴","수련교육부(인턴)",[@수련과목]))))', None),
    ("구리 근무 시작일(기준월)", 11, "calc", '=IF([@구리행]="","",INDEX(구리_시작,[@구리행]))', None),
    ("구리 근무 종료일(기준월)", 11, "calc", '=IF([@구리행]="","",INDEX(구리_종료,[@구리행]))', None),
    ("구리독자", 8, "calc", '=IF([@구리행]="","",INDEX(구리_독자,[@구리행]))', None),
    ("서울 근무일(기준월)", 9, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@구리행]<>"",MAX(0,{MD}-N([@구리명단 근무일])),IF([@기준월 근무병원]="서울병원",{MD},0)))', "0"),
    ("구리 근무일(기준월)", 9, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@구리행]<>"",N([@구리명단 근무일]),IF([@기준월 근무병원]="구리병원",{MD},0)))', "0"),
    ("급여지급병원(기준월)", 13, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@기준월 근무병원]="","",IF([@구리행]="",[@기준월 근무병원],'
     f'IF(AND({seoul}>={th},{guri}<{th}),"서울병원",IF(AND({guri}>={th},{seoul}<{th}),"구리병원",'
     f'"⚠15일 판정불가(서울 "&{seoul}&"일/구리 "&{guri}&"일)")))))', None),
    ("급여 판정 근거", 30, "calc",
     f'=IF(OR(NOT({ROW_ON}),[@급여지급병원(기준월)]=""),"",IF([@구리행]<>"",'
     f'"서울 "&{seoul}&"일 / 구리 "&{guri}&"일 (구리 근무자명단) → 15일 이상 근무 병원 지급",'
     f'IF({ok_mon},"인턴 근무계획 → 그 달 근무병원 지급","구리 근무자명단에 없음 → 소속병원 지급")))', None),
    ("점검", 26, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@기본자료행]="","기본자료에 없는 사번",IF(COUNTIF(명단표[사번],[@사번])>1,"중복 사번",'
     f'IF([@소속병원]="","소속(입사장소) 입력",'
     f'IF(LEFT([@급여지급병원(기준월)],1)="⚠","서울·구리 15일 판정 불가(양쪽 모두 이상/미만)",'
     f'IF(IF([@인턴계획행]="",0,N(INDEX(인턴계획표[입력오류],[@인턴계획행])))>0,"인턴계획 입력오류 확인",'
     f'IF(AND([@구리행]="",COUNTIF(구리_이름,[@이름])>0),"구리 근무자명단 연결 확인(동명이인·과 불일치)","정상")))))))', None),
    ("비고", 16, "in", None, None),
    # ---- 보조 열 (수정 금지)
    ("기본자료행", 8, "calc", '=IF(NOT(' + ROW_ON + '),"",IFERROR(MATCH("S|"&TRIM([@사번]&""),기본자료표[사번키],0),""))', None),
    ("이름|과", 14, "calc", '=IF([@이름]="","",[@이름]&"|"&[@수련과목])', None),
    ("인턴계획행", 8, "calc", '=IF(NOT(' + ROW_ON + '),"",IFERROR(MATCH("S|"&TRIM([@사번]&""),인턴계획표[사번키],0),""))', None),
    ("인턴 이번달", 15, "calc",
     '=IF(OR([@인턴계획행]="",인턴위치_이번=""),"",INDEX(인턴계획표[[정규화1]:[정규화6]],[@인턴계획행],인턴위치_이번)&"")', None),
    ("인턴 기준월", 15, "calc",
     '=IF(OR([@인턴계획행]="",인턴위치_기준월=""),"",INDEX(인턴계획표[[정규화1]:[정규화6]],[@인턴계획행],인턴위치_기준월)&"")', None),
    ("소속병원", 10, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@입사장소]<>"",[@입사장소],IF([@구리독자]="구리독자","구리병원",{home_n})))', None),
    ("기준월 근무병원", 10, "calc",
     f'=IF(NOT({ROW_ON}),"",IF({ok_mon},LEFT([@인턴 기준월],FIND(" ",[@인턴 기준월])-1),[@소속병원]))', None),
    ("구리행", 7, "calc",
     '=IF([@이름]="","",IFERROR(MATCH([@이름|과],구리_키,0),'
     'IF(AND(COUNTIF(구리_이름,[@이름])=1,COUNTIF(명단표[이름],[@이름])=1),MATCH([@이름],구리_이름,0),"")))', None),
    ("구리명단 근무일", 8, "calc", '=IF([@구리행]="","",INDEX(구리_근무일,[@구리행]))', "0"),
    ("구리 근무 중", 7, "calc",
     '=IF([@구리행]="","",IFERROR(IF(AND(TODAY()>=DATEVALUE([@구리 근무 시작일(기준월)]),'
     'TODAY()<=DATEVALUE([@구리 근무 종료일(기준월)])),"Y",""),""))', None),
    # ---- 출력 목록용 순번 (수정 금지)
    ("서울순번", 7, "calc",
     '=IF([@급여지급병원(기준월)]="서울병원",COUNTIF(INDEX(명단표[급여지급병원(기준월)],1):[@급여지급병원(기준월)],"서울병원"),"")', "0"),
    ("판정불가순번", 7, "calc",
     '=IF(LEFT([@급여지급병원(기준월)],1)="⚠",COUNTIF(INDEX(명단표[급여지급병원(기준월)],1):[@급여지급병원(기준월)],"⚠*"),"")', "0"),
    ("명단순번", 7, "calc", '=IF(NOT(' + ROW_ON + '),"",SUMPRODUCT(--(TRIM(INDEX(명단표[사번],1):[@사번]&"")<>"")))', "0"),
]
m_rows = DATA["roster_rows"] if DATA else [{"사번": s, "입사장소": h, "비고": n} for s, h, n in [
    ("2261212", "서울병원", "예시 데이터"), ("2261213", "서울병원", "예시 데이터"), ("2261214", "서울병원", "예시 데이터"),
    ("2261215", "서울병원", "예시 데이터"), ("2261216", "구리병원", "예시 데이터"), ("222222", "서울병원", "예시(인턴)"),
    ("222223", "서울병원", "예시(인턴)"), ("222224", "구리병원", "예시(인턴)")]]
make_table(ws, "명단표", 1, 1, mcols, m_rows)
MC = {h: CL(i + 1) for i, (h, *_rest) in enumerate(mcols)}
ws.row_dimensions[1].height = 42
ws.freeze_panes = "D2"
dv(ws, f"{MC['입사장소']}2:{MC['입사장소']}{MAXR}", "=병원목록", "등록되지 않은 값",
   "병원 목록에 없습니다. 새 병원이면 '설정_코드표' 병원표에 먼저 추가하세요.")
cf = ws.conditional_formatting
B, Z, Y, AE = MC["사번"], MC["점검"], MC["급여지급병원(기준월)"], MC["소속병원"]
cf.add(f"{B}2:{B}{MAXR}", FormulaRule(formula=[f'AND(${B}2<>"",${MC["기본자료행"]}2="")'], fill=PatternFill("solid", bgColor="FF0000"), font=Font(name=FONT, color="FFFFFF", bold=True)))
cf.add(f"{Z}2:{Z}{MAXR}", FormulaRule(formula=[f'AND(${Z}2<>"",${Z}2<>"정상")'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
cf.add(f"{Y}2:{Y}{MAXR}", FormulaRule(formula=[f'LEFT(${Y}2,1)="⚠"'], fill=PatternFill("solid", bgColor="FF0000"), font=Font(name=FONT, color="FFFFFF", bold=True)))
cf.add(f"{Y}2:{Y}{MAXR}", FormulaRule(formula=[f'AND(${Y}2<>"",${AE}2<>"",${Y}2<>${AE}2)'], fill=PatternFill("solid", bgColor="FFFF00")))
cf.add(f"{MC['생년월일']}2:{MC['생년월일']}{MAXR}", FormulaRule(formula=[f'ISTEXT(${MC["생년월일"]}2)'], font=Font(name=FONT, color="FF0000", bold=True)))

# ================================================================ 출력_서울급여
R0 = 6  # 출력 목록 첫 행
sp = wb.create_sheet("출력_서울급여")
sp["A1"].value, sp["A1"].font = "서울병원 급여 지급 대상자", f_title
sp["A2"].value, sp["A2"].font = "급여 기준월", f_bold
sp["B2"].value, sp["B2"].number_format, sp["B2"].font = "=급여기준월", 'yyyy"년" m"월"', Font(name=FONT, size=12, bold=True, color="1F3864")
YN = "명단표[급여지급병원(기준월)]"
for lc, lt, vc, vf in [("D2", "서울 지급 인원", "E2", f'=COUNTIF({YN},"서울병원")'),
                       ("F2", "그중 타 병원 소속", "G2", f'=COUNTIFS({YN},"서울병원",명단표[소속병원],"<>서울병원")'),
                       ("H2", "⚠ 판정 불가", "I2", f'=COUNTIF({YN},"⚠*")')]:
    sp[lc].value, sp[lc].font, sp[lc].alignment = lt, f_bold, Alignment(horizontal="right", vertical="center")
    sp[vc].value, sp[vc].font, sp[vc].alignment, sp[vc].border = vf, Font(name=FONT, size=12, bold=True), center, border
sp.conditional_formatting.add("I2", FormulaRule(formula=["I2>0"], fill=PatternFill("solid", bgColor="FF0000"), font=Font(name=FONT, color="FFFFFF", bold=True)))
sp["K2"].value = f'=IF(E2>{OUTN},"⚠ 표시 한도({OUTN}명) 초과 — 생성 파일의 OUTN을 늘려야 합니다","")'
sp["K2"].font = Font(name=FONT, size=10, bold=True, color="C00000")
sp["A3"].value, sp["A3"].font = (f"※ 자동 목록입니다(입력 금지, 최대 {OUTN}명 표시). 노란 행 = 소속은 다른 병원인데 이번 달 서울병원이 지급. "
                                 "급여 기준월은 설정_코드표에서 바꿉니다."), f_note
sp_heads = [("번호", 5), ("이름", 9), ("사번", 10), ("전공의등록번호", 12), ("수련과목", 13), ("연차", 9), ("소속병원", 10),
            ("구분", 24), ("현재 근무병원", 11), ("구리 근무 시작일", 11), ("구리 근무 종료일", 11),
            ("서울 근무일", 8), ("구리 근무일", 8), ("판정 근거", 40), ("점검", 26), ("(명단 행)", 6)]
for i, (h, w) in enumerate(sp_heads, 1):
    head_cell(sp.cell(row=5, column=i), h, fill_aux if h.startswith("(") else None)
    sp.column_dimensions[CL(i)].width = w
sp.row_dimensions[5].height = 30
ix = lambda col, j: f"INDEX(명단표[{col}],{j})"
cols_sp = [ix("이름", "{j}"), ix("사번", "{j}") + '&""', ix("전공의등록번호", "{j}") + '&""', ix("수련과목", "{j}") + '&""',
           ix("연차", "{j}") + '&""', ix("소속병원", "{j}"),
           f'IF({ix("소속병원", "{j}")}="서울병원","서울 소속","타 병원 소속("&{ix("소속병원", "{j}")}&") → 서울 지급")',
           ix("현재 근무병원", "{j}"), ix("구리 근무 시작일(기준월)", "{j}") + '&""', ix("구리 근무 종료일(기준월)", "{j}") + '&""',
           ix("서울 근무일(기준월)", "{j}"), ix("구리 근무일(기준월)", "{j}"), ix("급여 판정 근거", "{j}"), ix("점검", "{j}")]
HC = CL(len(sp_heads))  # 보조: 명단 행
e0 = len(sp_heads) + 2
EHC = CL(e0 + 5)        # 보조: 판정불가 명단 행
e_cols = [ix("이름", "{j}"), ix("사번", "{j}") + '&""', ix("소속병원", "{j}"), ix("서울 근무일(기준월)", "{j}"), ix("구리 근무일(기준월)", "{j}")]
for k in range(1, OUTN + 1):
    r = R0 + k - 1
    sp[f"{HC}{r}"].value = f"=IFERROR(MATCH({k},명단표[서울순번],0),\"\")"
    sp[f"A{r}"].value = f'=IF(${HC}{r}="","",{k})'
    for j, expr in enumerate(cols_sp):
        sp.cell(row=r, column=2 + j).value = f'=IF(${HC}{r}="","",{expr.format(j="$" + HC + str(r))})'
    sp[f"{EHC}{r}"].value = f"=IFERROR(MATCH({k},명단표[판정불가순번],0),\"\")"
    for j, expr in enumerate(e_cols):
        sp.cell(row=r, column=e0 + j).value = f'=IF(${EHC}{r}="","",{expr.format(j="$" + EHC + str(r))})'
    sp[f"{HC}{r}"].font = sp[f"{EHC}{r}"].font = f_note
sp[f"B{R0}"].value = f'=IF(${HC}{R0}="",IF(COUNTIF({YN},"서울병원")=0,"대상 없음",""),{cols_sp[0].format(j="$" + HC + str(R0))})'
sp.cell(row=R0, column=e0).value = f'=IF(${EHC}{R0}="",IF(COUNTIF({YN},"⚠*")=0,"없음",""),{e_cols[0].format(j="$" + EHC + str(R0))})'
RN = R0 + OUTN - 1
last = CL(len(sp_heads) - 1)
sp.conditional_formatting.add(f"A{R0}:{last}{RN}", FormulaRule(formula=[f'LEFT($H{R0},2)="타 "'], fill=PatternFill("solid", bgColor="FFF2CC")))
sp.conditional_formatting.add(f"{last}{R0}:{last}{RN}", FormulaRule(formula=[f'AND(${last}{R0}<>"",${last}{R0}<>"정상")'], fill=PatternFill("solid", bgColor="FFC7CE")))
sp.column_dimensions[CL(e0 - 1)].width = 3
sp.cell(row=4, column=e0, value="⚠ 서울·구리 15일 판정 불가 (양쪽 모두 15일 이상/미만) — 담당자 확인 필요").font = Font(name=FONT, size=10, bold=True, color="C00000")
for i, (h, w) in enumerate([("이름", 9), ("사번", 10), ("소속병원", 10), ("서울 근무일", 8), ("구리 근무일", 8), ("(명단 행)", 6)]):
    head_cell(sp.cell(row=5, column=e0 + i), h, fill_aux if h.startswith("(") else PatternFill("solid", fgColor="C00000"))
    sp.column_dimensions[CL(e0 + i)].width = w
sp.freeze_panes = "C6"

# ================================================================ 출력_기관명단 (시간외 판정 파일로 복사하는 용도)
ex = wb.create_sheet("출력_기관명단")
ex["A1"].value, ex["A1"].font = "기관명단 (급여 기준월 기준) — 시간외근무 판정 파일의 입력_기관명단에 값으로 붙여넣기", f_title
ex["A2"].value, ex["A2"].font = ("※ 아래 '복사할 범위'만 복사 → 시간외 판정 파일 입력_기관명단의 표 아래에 '값 붙여넣기'(빈 행까지 복사하면 표가 쓸데없이 늘어납니다). "
                                 "전월·대상월·익월이 필요하면 설정_코드표의 급여 기준월을 바꿔 가며 반복하세요."), f_note
ex["A3"].value, ex["A3"].font = "복사할 범위", f_bold
ex["B3"].value = f'=IF(COUNT($J$5:$J${4 + OUTN})=0,"없음","A5:I"&(4+COUNT($J$5:$J${4 + OUTN})))'
ex["B3"].font, ex["B3"].fill, ex["B3"].border = f_bold, fill_key, border
ex["D3"].value = f'=IF(SUMPRODUCT(--(TRIM(명단표[사번]&"")<>""))>{OUTN},"⚠ 표시 한도({OUTN}명) 초과","")'
ex["D3"].font = Font(name=FONT, size=10, bold=True, color="C00000")
ex_heads = [("기준월", 9), ("사번", 10), ("이름", 9), ("수련과목", 13), ("연차", 10), ("소속병원", 10),
            ("기준월 근무병원", 11), ("급여지급병원(기준월)", 13), ("급여 판정 근거", 40), ("(명단 행)", 6)]
for i, (h, w) in enumerate(ex_heads, 1):
    head_cell(ex.cell(row=4, column=i), h, fill_aux if h.startswith("(") else None)
    ex.column_dimensions[CL(i)].width = w
ex.row_dimensions[4].height = 30
ex_cols = ['TEXT(급여기준월,"yyyy-mm")', ix("사번", "{j}") + '&""', ix("이름", "{j}"), ix("수련과목", "{j}"), ix("연차", "{j}") + '&""',
           ix("소속병원", "{j}"), ix("기준월 근무병원", "{j}"), ix("급여지급병원(기준월)", "{j}"), ix("급여 판정 근거", "{j}")]
for k in range(1, OUTN + 1):
    r = 4 + k
    ex[f"J{r}"].value, ex[f"J{r}"].font = f"=IFERROR(MATCH({k},명단표[명단순번],0),\"\")", f_note
    for j, expr in enumerate(ex_cols):
        ex.cell(row=r, column=1 + j).value = f'=IF($J{r}="","",{expr.format(j="$J" + str(r))})'
ex["A5"].value = f'=IF($J5="","명단 없음",{ex_cols[0]})'
ex.freeze_panes = "C5"

# ================================================================ 출력_구리근무자
go = wb.create_sheet("출력_구리근무자")
GIN = q("입력_구리근무자명단")
GD = f"{GIN}!$A$4:${GURI_COLS_END}${GURI_ROWS_END}"
H0 = CI("O")             # 보조 계산 시작 열
HE = 5 + GR              # 보조 계산 끝 행 (6행 ↔ 원본 4행)
go["A1"].value, go["A1"].font = "구리병원 근무자 (구리 근무자명단 기준)", f_title
go["A2"].value, go["A2"].font = "급여 기준월", f_bold
go["B2"].value, go["B2"].number_format, go["B2"].font = "=급여기준월", 'yyyy"년" m"월"', Font(name=FONT, size=12, bold=True, color="1F3864")
go["D2"].value, go["D2"].font = "블록 시작 열", f_note
# 3행 머리글 'YYYY년M월\n근무예정자'에서 급여 기준월 블록을 찾는다 (한 셀 배열수식)
hh = f'{GIN}!$A$3:${GURI_COLS_END}$3&""'
p1, p2 = f'IFERROR(FIND("월",{hh}),999)', f'IFERROR(FIND(CHAR(10),{hh}),999)'
pe = f"IF({p1}<{p2},{p1},{p2})"
cse(go, "E2", f'=IFERROR(MATCH(YEAR(급여기준월)*100+MONTH(급여기준월),IF(ISNUMBER(SEARCH("근무예정자",{hh})),'
              f'IFERROR(VALUE(LEFT({hh},4)),0)*100+IFERROR(VALUE(MID({hh},6,{pe}-6)),0),0),0),"")')
go["F2"].value, go["F2"].font = "근무자 확정", f_note
go["G2"].value = f'=IF($E$2="","",IF(COUNTIF(INDEX({GD},0,$E$2+1),"?*")>0,"확정","예정"))'
for c in ("E2", "G2"):
    go[c].font, go[c].alignment = f_bold, center
go["A3"].value, go["A3"].font = ("※ 자동 결과(입력 금지). 그 달 '근무자' 칸이 하나라도 채워져 있으면 근무자, 아니면 근무예정자 기준. 날짜가 없으면 1일~말일. "
                                 "명단 연결: 이름+진료과(인턴은 이름+인턴), 안 되면 이름만. '블록 시작 열'이 비면 3행 머리글(연·월)을 확인하세요. "
                                 f"원본은 4~{3 + GR}행까지 해석, 목록은 최대 {GV}명 표시. O열부터는 보조 계산(+ 단추로 펼치기)."), f_note
HCOL = {n: CL(H0 + i) for i, n in enumerate(["원본행", "진료과", "연차", "이름", "시작값", "종료값", "근무일", "비고", "구리독자",
                                              "이름|과", "명단 연결", "시작일", "종료일", "누계", "순번"])}
H = lambda n: f"${HCOL[n]}$6:${HCOL[n]}${HE}"
for lc, lt, vc, vf in [("I1", "명단 연결", "J1", f'=COUNTIF({H("명단 연결")},"명단 일치*")'),
                       ("I2", "명단에 없음/동명이인", "J2", f'=COUNTIF({H("명단 연결")},"명단에 없음*")+COUNTIF({H("명단 연결")},"동명이인*")'),
                       ("I3", "날짜 오류", "J3", f'=COUNTIF({H("시작일")},"⚠*")+COUNTIF({H("종료일")},"⚠*")')]:
    go[lc].value, go[lc].font, go[lc].alignment = lt, f_bold, Alignment(horizontal="right")
    go[vc].value, go[vc].font, go[vc].border, go[vc].alignment = vf, f_bold, border, center
go["K2"].value = (f'=IF(MAX({H("누계")})>{GV},"⚠ 목록 표시 한도({GV}명) 초과",'
                  f'IF(COUNTA({GIN}!$A${4 + GR}:$B${GURI_ROWS_END})>0,"⚠ 원본 {3 + GR}행 이후 자료는 해석되지 않음",""))')
go["K2"].font = Font(name=FONT, size=10, bold=True, color="C00000")
go.conditional_formatting.add("J2:J3", FormulaRule(formula=["J2>0"], fill=PatternFill("solid", bgColor="FFC7CE")))
go_heads = [("원본행", 6), ("진료과", 14), ("연차", 6), ("이름", 9), ("확정/예정", 8), ("시작일", 11), ("종료일", 11),
            ("기준월 구리 근무일", 9), ("비고", 10), ("구리독자", 8), ("이름|과", 16), ("명단 연결", 24)]
for i, (h, w) in enumerate(go_heads, 1):
    head_cell(go.cell(row=5, column=i), h)
    go.column_dimensions[CL(i)].width = w
go.row_dimensions[5].height = 30
go.cell(row=4, column=H0, value="보조 계산 (수정 금지) — 원본 각 행을 해석").font = f_note
for n, c in HCOL.items():
    head_cell(go[f"{c}5"], n, fill_aux)
    go.column_dimensions[c].width = 10
BOM = "급여기준월"


def gcell(raw, off):
    return f"INDEX({GIN}!$A{raw}:${GURI_COLS_END}{raw},1,$E$2+{off})"


def pdate(raw, off, default):
    x = gcell(raw, off)
    return f'IF({x}&""="",{default},IF(ISNUMBER({x}),{x},IFERROR(DATE(VALUE(LEFT({x},4)),VALUE(MID({x},6,2)),VALUE(MID({x},9,2))),-1)))'


for i in range(GR):
    r, raw = 6 + i, 4 + i
    c = {n: f"${HCOL[n]}{r}" for n in HCOL}
    a_raw = f'TRIM({GIN}!$A{raw}&"")'
    name_cell = gcell(raw, 'IF($G$2="확정",1,0)')
    f = {
        "원본행": f"={raw}",
        "진료과": f"={a_raw}" if i == 0 else f'=IF({a_raw}<>"",{a_raw},{HCOL["진료과"]}{r - 1})',
        "연차": f'={GIN}!$B{raw}&""',
        "이름": f'=IF($E$2="","",TRIM({name_cell}&""))',
        "시작값": f'=IF({c["이름"]}="","",{pdate(raw, 2, BOM)})',
        "종료값": f'=IF({c["이름"]}="","",{pdate(raw, 3, "EOMONTH(" + BOM + ",0)")})',
        "근무일": (f'=IF({c["이름"]}="","",IF(OR({c["시작값"]}<0,{c["종료값"]}<0),0,'
                 f'MAX(0,MIN({c["종료값"]},EOMONTH({BOM},0))-MAX({c["시작값"]},{BOM})+1)))'),
        "비고": f'=IF({c["이름"]}="","",TRIM({gcell(raw, 5)}&""))',
        "구리독자": f'=IF(ISNUMBER(SEARCH("구리독자",{c["비고"]})),"구리독자","")',
        "이름|과": (f'=IF({c["이름"]}="","",{c["이름"]}&"|"&IF(OR({c["연차"]}="인턴",'
                  f'ISNUMBER(SEARCH("수련교육부",{c["진료과"]}))),"인턴",{c["진료과"]}))'),
        "명단 연결": (f'=IF({c["이름"]}="","",IF(COUNTIF(명단표[이름|과],{c["이름|과"]})=1,"명단 일치",'
                  f'IF(AND(COUNTIF(명단표[이름|과],{c["이름|과"]})=0,COUNTIF(명단표[이름],{c["이름"]})=1),"명단 일치(이름만, 과 다름)",'
                  f'IF(COUNTIF(명단표[이름],{c["이름"]})=0,IF(COUNTIF(기본자료표[이름],{c["이름"]})>0,"명단에 없음(기본자료에는 있음)","명단에 없음"),'
                  f'"동명이인 확인"))))'),
        "시작일": f'=IF({c["이름"]}="","",IF({c["시작값"]}<0,"⚠"&{gcell(raw, 2)},TEXT({c["시작값"]},"yyyy-mm-dd")))',
        "종료일": f'=IF({c["이름"]}="","",IF({c["종료값"]}<0,"⚠"&{gcell(raw, 3)},TEXT({c["종료값"]},"yyyy-mm-dd")))',
        "누계": (f'=IF({c["이름"]}="",0,1)' if i == 0 else f'={HCOL["누계"]}{r - 1}+IF({c["이름"]}="",0,1)'),
        "순번": f'=IF({c["이름"]}="","",{c["누계"]})',
    }
    for n, fm in f.items():
        go[f"{HCOL[n]}{r}"].value = fm
go.column_dimensions.group(HCOL["원본행"], HCOL["순번"], hidden=True, outline_level=1)
vis = [("진료과", None), ("연차", None), ("이름", None), (None, "$G$2"), ("시작일", None), ("종료일", None), ("근무일", None),
       ("비고", None), ("구리독자", None), ("이름|과", None), ("명단 연결", None)]
for k in range(1, GV + 1):
    r = 5 + k
    go[f"A{r}"].value = f'=IFERROR(INDEX({H("원본행")},MATCH({k},{H("순번")},0)),"")'
    for j, (n, fixed) in enumerate(vis):
        go.cell(row=r, column=2 + j).value = f'=IF(ISNUMBER($A{r}),{fixed if fixed else "INDEX(" + H(n) + ",$A" + str(r) + "-3)"},"")'
go["A6"].value = f'=IF($E$2="","해당 월 블록 없음",IF(MAX({H("누계")})=0,"해당 월 근무자 없음",INDEX({H("원본행")},MATCH(1,{H("순번")},0))))'
GE = 5 + GV
go.conditional_formatting.add(f"L6:L{GE}", FormulaRule(formula=['AND($L6<>"",LEFT($L6,5)<>"명단 일치")'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
go.conditional_formatting.add(f"F6:G{GE}", FormulaRule(formula=['LEFT(F6,1)="⚠"'], font=Font(name=FONT, color="C00000", bold=True)))
go.conditional_formatting.add(f"A6:K{GE}", FormulaRule(formula=['$J6="구리독자"'], fill=PatternFill("solid", bgColor="E2EFDA")))
go.freeze_panes = "E6"
GO = q("출력_구리근무자")
for nm, n in [("구리_키", "이름|과"), ("구리_이름", "이름"), ("구리_근무일", "근무일"), ("구리_독자", "구리독자"),
              ("구리_시작", "시작일"), ("구리_종료", "종료일")]:
    add_name(wb, nm, f"{GO}!{H(n)}")

# ================================================================ 출력_요약
sm = wb.create_sheet("출력_요약")
sm["A1"].value, sm["A1"].font = "현황 요약", f_title
for r, (lab, fml, fmt) in enumerate([("기준일", "=TODAY()", "yyyy-mm-dd"), ("급여 기준월", "=급여기준월", "yyyy-mm-dd"),
                                     ("15일 기준(일)", "=판정기준일", "0")], 2):
    sm[f"A{r}"].value, sm[f"A{r}"].font = lab, f_bold
    sm[f"B{r}"].value, sm[f"B{r}"].number_format, sm[f"B{r}"].font = fml, fmt, f_base
sm["C3"].value, sm["C3"].font = "← 설정값은 설정_코드표에서 변경", f_note
head_cell(sm["A6"], "인원 (급여 기준월)")
head_cell(sm["B6"], "인원")
people = [
    ("명단 등록 인원", '=SUMPRODUCT(--(TRIM(명단표[사번]&"")<>""))'),
    ("서울병원 지급", '=COUNTIF(명단표[급여지급병원(기준월)],"서울병원")'),
    ("구리병원 지급", '=COUNTIF(명단표[급여지급병원(기준월)],"구리병원")'),
    ("그 외 병원 지급", '=SUMPRODUCT((명단표[급여지급병원(기준월)]<>"")*(명단표[급여지급병원(기준월)]<>"서울병원")*(명단표[급여지급병원(기준월)]<>"구리병원")*(LEFT(명단표[급여지급병원(기준월)],1)<>"⚠"))'),
    ("⚠ 판정 불가", '=COUNTIF(명단표[급여지급병원(기준월)],"⚠*")'),
    ("구리 근무자명단 연결 인원", '=COUNT(명단표[구리행])'),
]
for i, (lab, fml) in enumerate(people):
    r = 7 + i
    sm[f"A{r}"].value, sm[f"B{r}"].value = lab, fml
for r in range(7, 7 + len(people)):
    for c in "AB":
        sm[f"{c}{r}"].border, sm[f"{c}{r}"].font = border, f_base
    sm[f"B{r}"].alignment = center
CK0 = 15
head_cell(sm[f"A{CK0}"], "데이터 점검")
head_cell(sm[f"B{CK0}"], "건수")
checks = [
    ("기본자료 인원", '=SUMPRODUCT(--(TRIM(기본자료표[사번]&"")<>""))'),
    ("기본자료 재직 인원", '=COUNTIF(기본자료표[현재근무 여부],"재직")'),
    ("기본자료에 없는 사번", '=SUMPRODUCT((TRIM(명단표[사번]&"")<>"")*(명단표[기본자료행]=""))'),
    ("명단 중복 사번", '=SUMPRODUCT((TRIM(명단표[사번]&"")<>"")*(COUNTIF(명단표[사번],명단표[사번])>1))'),
    ("날짜 형식 오류(생년월일/수련개시일)", '=SUMPRODUCT(--ISTEXT(명단표[생년월일]))+SUMPRODUCT(--ISTEXT(명단표[수련개시일]))'),
    ("인턴계획 입력 오류(병원/진료과)", "=SUM(인턴계획표[입력오류])"),
    ("서울·구리 15일 판정 불가", '=COUNTIF(명단표[급여지급병원(기준월)],"⚠*")'),
    ("구리 근무자명단: 명단에 없음/동명이인", f"={q('출력_구리근무자')}!$J$2"),
    ("구리 근무자명단: 날짜 오류", f"={q('출력_구리근무자')}!$J$3"),
    ("점검 필요 행(전공의명단 '점검' 열)", '=SUMPRODUCT((TRIM(명단표[사번]&"")<>"")*(명단표[점검]<>"정상"))'),
    ("출력 표시 한도 초과(서울급여·기관명단·구리·인턴)",
     f'=(COUNTIF(명단표[급여지급병원(기준월)],"서울병원")>{OUTN})+(SUMPRODUCT(--(TRIM(명단표[사번]&"")<>""))>{OUTN})'
     f'+({q("출력_구리근무자")}!$K$2<>"")+(SUMPRODUCT(--(TRIM(인턴계획표[이름]&"")<>""))>{IVN})'),
]
for i, (lab, fml) in enumerate(checks):
    r = CK0 + 1 + i
    sm[f"A{r}"].value, sm[f"B{r}"].value = lab, fml
    for c in "AB":
        sm[f"{c}{r}"].border, sm[f"{c}{r}"].font = border, f_base
    sm[f"B{r}"].alignment = center
sm.conditional_formatting.add(f"B{CK0 + 3}:B{CK0 + len(checks)}", FormulaRule(formula=[f"B{CK0 + 3}>0"], fill=PatternFill("solid", bgColor="FFC7CE")))
blocks = [("D", ["병원", "현재 근무", "급여 지급(기준월)"], "병원표[병원]",
           ["COUNTIF(명단표[현재 근무병원],{k})", "COUNTIF(명단표[급여지급병원(기준월)],{k})"]),
          ("H", ["현재 진료과", "인원"], "진료과표[근무과]", ["COUNTIF(명단표[현재 진료과],{k})"]),
          ("K", ["인턴 구분", "인원"], "인턴구분표[인턴 구분]", ["COUNTIF(명단표[현재 진료과],{k})"])]
for col, heads, keyref, counts in blocks:
    c0 = CI(col)
    for j, h in enumerate(heads):
        head_cell(sm.cell(row=6, column=c0 + j), h)
        sm.column_dimensions[CL(c0 + j)].width = 14 if j == 0 else 12
    for k in range(1, CODEN + 1):
        r = 6 + k
        key = f"${col}{r}"
        sm[f"{col}{r}"].value = f'=IFERROR(INDEX({keyref},{k})&"","")'
        for j, cnt in enumerate(counts, 1):
            sm.cell(row=r, column=c0 + j).value = f'=IF({key}="","",{cnt.format(k=key)})'
sm.column_dimensions["A"].width = 40
sm.column_dimensions["B"].width = 12
for c in "CGJ":
    sm.column_dimensions[c].width = 3
sm.row_dimensions[6].height = 30

# ================================================================ 출력_인턴현황
iv = wb.create_sheet("출력_인턴현황")
iv["A1"].value, iv["A1"].font = "인턴 근무 현황", f_title
iv["A2"].value, iv["A2"].font = "기준일", f_bold
iv["B2"].value, iv["B2"].number_format = "=TODAY()", "yyyy-mm-dd"
iv["A3"].value, iv["A3"].font = "계획 월 번호(자동)", f_note
iv["F3"].value, iv["F3"].font = "이번/다음 달 위치→", f_note
iv["G3"].value = '=IFERROR(MATCH(MONTH(TODAY()),$J$3:$O$3,0),"")'
iv["I3"].value = '=IFERROR(MATCH(MONTH(EDATE(TODAY(),1)),$J$3:$O$3,0),"")'
iv["J2"].value, iv["J2"].font = "급여 기준월 위치→", f_note
iv["K2"].value = '=IFERROR(MATCH(MONTH(급여기준월),$J$3:$O$3,0),"")'
for c in ("G3", "I3", "K2"):
    iv[c].font, iv[c].alignment = f_note, center
iv["L2"].value = f'=IF(SUMPRODUCT(--(TRIM(인턴계획표[이름]&"")<>""))>{IVN},"⚠ 표시 한도({IVN}명) 초과","")'
iv["L2"].font = Font(name=FONT, size=10, bold=True, color="C00000")
add_name(wb, "인턴위치_이번", f"{q('출력_인턴현황')}!$G$3")
add_name(wb, "인턴위치_다음", f"{q('출력_인턴현황')}!$I$3")
add_name(wb, "인턴위치_기준월", f"{q('출력_인턴현황')}!$K$2")
iv_heads = ["순번", "이름", "사번", "기본자료 확인", "입사일",
            '="이번 달("&IF($G$3="","계획 외",MONTH(TODAY())&"월")&") 병원"', "이번 달 진료과",
            '="다음 달("&IF($I$3="","계획 외",MONTH(EDATE(TODAY(),1))&"월")&") 병원"', "다음 달 진료과"]
iv_heads += [f'=TRIM(INDEX(인턴계획표[#Headers],1,{3 + k})&"")' for k in range(1, 7)] + ["입력오류", "(계획 행)"]
for i, h in enumerate(iv_heads, 1):
    head_cell(iv.cell(row=5, column=i), h, fill_aux if h.startswith("(") else None)
    iv.column_dimensions[CL(i)].width = 17 if 6 <= i <= 15 else 10
iv.column_dimensions["Q"].width = 6
iv.row_dimensions[5].height = 30
for k in range(6):
    c = CL(10 + k)
    iv[f"{c}3"].value = f'=IFERROR(VALUE(SUBSTITUTE(TRIM({c}5),"월","")),"")'
    iv[f"{c}3"].font, iv[f"{c}3"].alignment = f_note, center
hosp = lambda x: f'IF({x}="","",IF(LEFT({x},1)="⚠",{x},IFERROR(LEFT({x},FIND(" ",{x})-1),{x})))'
dept = lambda x: f'IF({x}="","",IF(LEFT({x},1)="⚠","",IFERROR(MID({x},FIND(" ",{x})+1,60),"")))'
NRM = "인턴계획표[[정규화1]:[정규화6]]"
for k in range(1, IVN + 1):
    r = 5 + k
    j = f"$Q{r}"
    ii = lambda col: f"INDEX(인턴계획표[{col}],{j})"
    cur = f'IF($G$3="","",INDEX({NRM},{j},$G$3)&"")'
    nxt = f'IF($I$3="","",INDEX({NRM},{j},$I$3)&"")'
    vals = [ii("순번") + '&""', f'TRIM({ii("이름")}&"")', ii("사번") + '&""',
            f'IF(COUNTIF(기본자료표[사번키],{ii("사번키")})>0,"확인","미등록")',
            f'IF(ISNUMBER({ii("입사일(비고)")}),TEXT({ii("입사일(비고)")},"yyyy-mm-dd"),{ii("입사일(비고)")}&"")',
            hosp(cur), dept(cur), hosp(nxt), dept(nxt)] + [ii(f"정규화{m}") for m in range(1, 7)] + [ii("입력오류")]
    iv[f"Q{r}"].value, iv[f"Q{r}"].font = f'=IFERROR(MATCH({k},인턴계획표[표시순번],0),"")', f_note
    for c0, v in enumerate(vals, 1):
        iv.cell(row=r, column=c0).value = f'=IF({j}="","",{v})'
iv["B6"].value = f'=IF($Q6="",IF(SUMPRODUCT(--(TRIM(인턴계획표[이름]&"")<>""))=0,"인턴 계획 없음",""),TRIM(INDEX(인턴계획표[이름],$Q6)&""))'
END = 5 + IVN
iv.conditional_formatting.add(f"D6:D{END}", FormulaRule(formula=['$D6="미등록"'], fill=PatternFill("solid", bgColor="FFC7CE")))
iv.conditional_formatting.add(f"J6:O{END}", FormulaRule(formula=['LEFT(J6,1)="⚠"'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=FONT, color="C00000", bold=True)))
iv.conditional_formatting.add(f"J5:O{END}", FormulaRule(formula=['J$3=MONTH(TODAY())'], fill=PatternFill("solid", bgColor="FFF2CC")))
iv.freeze_panes = "D6"
# 월별 병원별 인원
iv["S2"].value, iv["S2"].font = "월별 병원별 인턴 인원", f_bold
head_cell(iv["S3"], "병원")
iv["S4"].value, iv["S5"].value = "합계(배정)", "⚠ 오류"
for c in ("S4", "S5"):
    iv[c].font, iv[c].border = f_bold, border
iv.column_dimensions["R"].width = 3
iv.column_dimensions["S"].width = 14
for k in range(1, CODEN + 1):
    iv[f"S{5 + k}"].value = f'=IFERROR(INDEX(병원표[병원],{k})&"","")'
for k in range(6):
    c = CL(20 + k)
    head_cell(iv[f"{c}3"], f"={CL(10 + k)}5")
    iv[f"{c}4"].value = f'=COUNTIF(인턴계획표[정규화{k + 1}],"?*")'
    iv[f"{c}5"].value = f'=COUNTIF(인턴계획표[정규화{k + 1}],"⚠*")'
    for rr in (4, 5):
        iv[f"{c}{rr}"].font, iv[f"{c}{rr}"].border, iv[f"{c}{rr}"].alignment = f_bold, border, center
    for kk in range(1, CODEN + 1):
        iv[f"{c}{5 + kk}"].value = f'=IF($S{5 + kk}="","",COUNTIF(인턴계획표[정규화{k + 1}],$S{5 + kk}&" *"))'
    iv.column_dimensions[c].width = 10

# ================================================================ 사용안내
gd = wb.create_sheet("사용안내", 0)
lines = [
    ("전공의·인턴 파견 모니터링 사용 안내 (Excel 2019 이상 · Microsoft 365 공통)", f_title),
    ("", f_base),
    ("[시트 구분] — 탭 색: 파랑=입력, 초록=출력(자동, 입력 금지), 회색=설정", f_bold),
    ("• 입력_기본자료 / 입력_인턴근무계획 / 입력_전공의명단: 엑셀 '표'. 표 바로 아래(또는 첫 자료 행)에 붙여넣으면 표가 자동으로 늘어나고 초록 머리글 계산 열도 자동으로 채워집니다.", f_base),
    ("• 입력_구리근무자명단: 구리병원 파일의 시트를 A1에 양식 그대로 붙여넣는 시트(병합 셀이 있어 표가 아님). 4행부터 "+str(GR)+"행("+str(3 + GR)+"행까지)·ZZ열까지 해석합니다.", f_base),
    ("• 출력_서울급여 / 출력_구리근무자 / 출력_요약 / 출력_인턴현황 / 출력_기관명단: 미리 수식이 들어 있는 고정 행 목록. 셀을 지우거나 덮어쓰지 마세요.", f_base),
    (f"• 출력 표시 한도: 서울급여·기관명단 {OUTN}명, 구리 근무자 {GV}명, 인턴현황 {IVN}명, 코드 목록 {CODEN}개. 넘으면 시트 위쪽과 출력_요약에 ⚠ 경고가 뜹니다(입력 표는 제한 없음).", f_base),
    ("• 회색 머리글 '(명단 행)'·'(계획 행)' 열과 출력_구리근무자 O열 이후(+ 단추로 펼침)는 목록을 만드는 보조 계산입니다(수정 금지).", f_base),
    ("• 설정_코드표: 병원·진료과·인턴구분·약칭 목록(표)과 설정값(급여 기준월·15일 기준).", f_base),
    ("• 출력_기관명단: 급여 기준월의 전원 기관(근무병원·급여지급병원) — 시간외근무 판정 파일의 입력_기관명단에 값으로 붙여넣는 용도", f_base),
    ("", f_base),
    ("[붙여넣는 방법]", f_bold),
    ("• 기본자료: A2부터 (머리글 순서: 현재근무 여부·사번·이름·수련과목·연차·근무지·수련개시일·전공의등록번호·의사면허번호·성별·생년월일·비고). 날짜는 20260301 형식도 됩니다.", f_base),
    ("• 인턴근무계획: B3부터 자료 행(순번·이름·사번·월 6칸·비고). 학기가 바뀌면 2행의 월 머리글(9월~2월 등)도 바꿔 적으세요. 비고 2026.03 = 2026-03-01 입사.", f_base),
    ("• 구리 근무자명단: 구리병원 파일의 시트를 A1에 양식 그대로(1~2행 제목, 3행 머리글, 4행부터 자료). 날짜는 'YYYY.MM.DD' 글자도 됩니다.", f_base),
    ("• 전공의명단: 사번·입사장소·비고만 입력. 나머지는 자동.", f_base),
    ("• 예시 데이터(ㅇㅇ)는 모두 지우고 실제 자료를 넣으세요. 표의 행을 지울 때는 행 전체 삭제(표 행 삭제)를 쓰세요.", f_base),
    ("", f_base),
    ("[표기 해석]", f_bold),
    ("• 인턴 월 칸은 '지역 진료과'(예: 서울 산부 → 서울병원 산부인과, 제주 에스중앙 내과 → 제주병원 내과). 첫 단어=지역(병원약칭표), 마지막 단어=진료과(과약칭표).", f_base),
    ("• 통합(과 미지정), 핵/방, 직/피, 정/영 등은 인턴 전용 구분입니다(과약칭표에서 직피·정영 등 약칭도 인정).", f_base),
    ("• 코드표에 없는 값을 직접 입력하면 오류 팝업이 뜹니다. 붙여넣기는 팝업이 뜨지 않으므로 해당 칸이 빨갛게 표시되고 요약에 집계됩니다.", f_base),
    ("", f_base),
    ("[급여지급병원(기준월) 규칙]", f_bold),
    ("• 서울병원·구리병원 사이: 소속과 무관하게 급여 기준월에 15일 이상 근무한 병원이 그 달 전체를 지급(예: 구리 10일·서울 20일 → 서울 지급).", f_base),
    ("• 근무일: 구리 근무자명단에 있으면 그 날짜가 구리 근무일, 그 달의 나머지 날은 모두 서울 근무일입니다.", f_base),
    ("• 양쪽 모두 15일 이상이거나 모두 미만이면 ⚠ 판정 불가로 표시합니다.", f_base),
    ("• 구리 근무자명단에 없는 사람은 그 달 근무병원(인턴은 인턴근무계획의 병원, 그 외는 소속병원)이 지급합니다.", f_base),
    ("• 소속병원 = 입사장소 → (비어 있으면) 구리 근무자명단 비고 '구리독자'면 구리병원 → 기본자료 근무지 순서로 정합니다.", f_base),
    ("", f_base),
    ("[참고]", f_bold),
    ("• Excel 2019·2021·Microsoft 365에서 같은 결과가 나옵니다(FILTER·XLOOKUP 같은 365 전용 함수를 쓰지 않음). 2016 이하는 확인하지 않았습니다.", f_base),
    ("• 목록 순서는 입력_전공의명단(인턴은 인턴근무계획) 행 순서입니다. 다른 순서로 보려면 출력 시트를 복사해 값으로 붙여넣은 뒤 정렬하세요.", f_base),
    ("• 개인정보(생년월일·면허번호 등)가 포함되므로 파일 암호 설정과 접근 권한 관리를 권장합니다.", f_base),
]
for i, (tx, fo) in enumerate(lines, 1):
    gd.cell(row=i, column=1, value=tx).font = fo
gd.column_dimensions["A"].width = 175
gd.sheet_view.showGridLines = False

# ================================================================ 시트 순서·탭 색, 저장
ORDER = ["사용안내", "입력_기본자료", "입력_인턴근무계획", "입력_구리근무자명단", "입력_전공의명단",
         "출력_서울급여", "출력_구리근무자", "출력_요약", "출력_인턴현황", "출력_기관명단", "설정_코드표"]
wb._sheets = [wb[n] for n in ORDER]
wb.active = 0
for sh in wb.worksheets:
    for pre, color in {"입력_": "2F5597", "출력_": "548235", "설정_": "7F7F7F"}.items():
        if sh.title.startswith(pre):
            sh.sheet_properties.tabColor = color
save(wb, OUT)
print("saved", OUT)
