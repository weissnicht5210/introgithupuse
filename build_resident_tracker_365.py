"""전공의·인턴 파견 모니터링 엑셀 — Microsoft 365 · Excel 2021용 (행 수 제한 없음)

- Excel 2019용은 build_resident_tracker.py. 이 판에는 입력_월별근무지(전월·익월 등 근무지 직접 입력)가 더 있다.

- 입력 시트는 엑셀 '표'(Table): 아래에 붙여넣으면 표와 계산 열이 자동으로 늘어난다.
- 출력 시트는 동적 배열(FILTER 등): 대상 수만큼 자동으로 펼쳐진다.
- DATA=<pickle> 로 실행하면 기존 파일에서 옮긴 입력 자료로 채운다(실제 자료는 저장소에 넣지 않음).
- VERIFY=1 로 실행하면 LibreOffice 검증용으로 동적 배열을 고정 크기 배열수식으로 기록한다.
"""
import os
import re
import zipfile
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

VERIFY = os.environ.get("VERIFY") == "1"
DATA = None  # 실제 자료로 채울 때: DATA=<pickle 경로> (입력 시트 값만 담은 사전)
if os.environ.get("DATA"):
    import pickle
    with open(os.environ["DATA"], "rb") as fh:
        DATA = pickle.load(fh)
OUT = os.environ.get("OUT", "전공의_파견_모니터링_365.xlsx")
MAXR = 1048576          # 서식·검증 적용 끝 행
GURI_ROWS_END = 20000   # 구리 근무자명단에서 읽는 끝 행 (사실상 무제한)
GURI_COLS_END = "ZZ"    # 구리 근무자명단에서 읽는 끝 열 (월 블록 약 116개월)
SPILL_END = 20005       # 출력_구리근무자 결과를 참조하는 끝 행
VERIFY_ROWS = int(os.environ.get("VERIFY_ROWS", "40"))  # 검증용 배열수식 높이

FONT = "맑은 고딕"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_input = Font(name=FONT, size=10, color="0000FF")  # 입력값=파랑, 수식=검정
f_title = Font(name=FONT, size=14, bold=True)
f_note = Font(name=FONT, size=9, color="7F7F7F")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_head2 = PatternFill("solid", fgColor="548235")
fill_calc = PatternFill("solid", fgColor="F2F2F2")
fill_key = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center")

# ---------------------------------------------------------------- 코드 목록
HOSPITALS = ["서울병원", "인천병원", "구리병원", "창원병원", "제주병원", "오산병원"]
DEPTS = ["내과", "소아청소년과", "신경과", "정신건강의학과", "피부과", "외과", "심장혈관흉부외과",
         "정형외과", "신경외과", "성형외과", "산부인과", "안과", "이비인후과", "비뇨의학과",
         "재활의학과", "마취통증의학과", "영상의학과", "진단검사의학과", "병리과", "가정의학과",
         "응급의학과", "핵의학과", "직업환경의학과", "수련교육부(인턴)"]
INTERN_ONLY = ["통합", "핵/방", "직/피", "정/영"]
DEPT_ALIASES = [  # (약칭, 정식 명칭) — 사용자 제공 목록
    ("내과", "내과"), ("소아청소년과", "소아청소년과"), ("신경과", "신경과"), ("정신건강의학과", "정신건강의학과"),
    ("피부과", "피부과"), ("외과", "외과"), ("심장혈관흉부외과", "심장혈관흉부외과"), ("정형외과", "정형외과"),
    ("신경외과", "신경외과"), ("성형외과", "성형외과"), ("산부인과", "산부인과"), ("안과", "안과"),
    ("이비인후과", "이비인후과"), ("비뇨의학과", "비뇨의학과"), ("재활의학과", "재활의학과"),
    ("마취통증의학과", "마취통증의학과"), ("영상의학과", "영상의학과"), ("진단검사의학과", "진단검사의학과"),
    ("병리과", "병리과"), ("가정의학과", "가정의학과"), ("응급의학과", "응급의학과"), ("핵의학과", "핵의학과"),
    ("직업환경의학과", "직업환경의학과"), ("수련교육부(인턴)", "수련교육부(인턴)"),
    ("소아과", "소아청소년과"), ("소아", "소아청소년과"), ("소청", "소아청소년과"), ("소청과", "소아청소년과"),
    ("정신과", "정신건강의학과"), ("정신", "정신건강의학과"), ("흉부외과", "심장혈관흉부외과"),
    ("흉외", "심장혈관흉부외과"), ("심장혈관흉부", "심장혈관흉부외과"), ("정형", "정형외과"), ("신외", "신경외과"),
    ("성형", "성형외과"), ("산부", "산부인과"), ("이비", "이비인후과"), ("이비인후", "이비인후과"),
    ("비뇨", "비뇨의학과"), ("비뇨기과", "비뇨의학과"), ("재활", "재활의학과"), ("마취", "마취통증의학과"),
    ("마통", "마취통증의학과"), ("영상", "영상의학과"), ("진검", "진단검사의학과"), ("진단검사", "진단검사의학과"),
    ("병리", "병리과"), ("가정", "가정의학과"), ("가정의학", "가정의학과"), ("응급", "응급의학과"),
    ("핵의학", "핵의학과"), ("핵", "핵의학과"), ("직업환경", "직업환경의학과"), ("직환", "직업환경의학과"),
    ("수련교육부", "수련교육부(인턴)"), ("인턴", "수련교육부(인턴)"), ("직피", "직/피"), ("직환/피부", "직/피"),
    ("정영", "정/영"), ("정신/영상", "정/영"), ("핵방", "핵/방"), ("핵/방", "핵/방"), ("통합", "통합"),
]
HOSP_ALIASES = [(h, h) for h in HOSPITALS] + [(h[:-2], h) for h in HOSPITALS] + [
    ("한양대학교병원", "서울병원"), ("한양대병원", "서울병원"), ("한양대서울병원", "서울병원"),
    ("한양대구리병원", "구리병원"), ("에스중앙병원", "제주병원"), ("에스중앙", "제주병원"),
]

# ---------------------------------------------------------------- 수식 처리 도우미
_FN = [("XLOOKUP", "_xlfn.XLOOKUP"), ("XMATCH", "_xlfn.XMATCH"), ("SEQUENCE", "_xlfn.SEQUENCE"),
       ("LET", "_xlfn.LET"), ("FILTER", "_xlfn._xlws.FILTER"), ("SORT", "_xlfn._xlws.SORT")]


def xl(f):
    """사람이 읽는 수식 → 파일 저장 형식 (365 함수 접두어, LET 변수 접두어)."""
    if not isinstance(f, str):
        return f
    f = re.sub(r"(?<![\w.])(v_[A-Za-z0-9_]+)", r"_xlpm.\1", f)
    for name, full in _FN:
        f = re.sub(rf"(?<![\w.]){name}\(", full + "(", f)
    return f


def this_row(table, f):
    """[@열] → 표[[#This Row],[열]]"""
    return re.sub(r"\[@([^\]]+)\]", lambda m: f"{table}[[#This Row],[{m.group(1)}]]", f)


SPILLS = {}  # 시트 이름 -> 동적 배열 셀 목록


def spill(ws, cell, formula, width=1, height=VERIFY_ROWS):
    """동적 배열 수식. Excel은 cm 메타데이터로 펼치고, 검증 모드에선 고정 크기 배열수식."""
    f = xl(formula)
    if VERIFY:
        col = re.match(r"([A-Z]+)(\d+)", cell)
        c0 = col.group(1)
        r0 = int(col.group(2))
        from openpyxl.utils import column_index_from_string as ci
        end = f"{CL(ci(c0) + width - 1)}{r0 + height - 1}"
        ws[cell] = ArrayFormula(f"{cell}:{end}", f)
    else:
        ws[cell] = ArrayFormula(cell, f)
        SPILLS.setdefault(ws.title, []).append(cell)


def q(sheet):
    return f"'{sheet}'"


def head_cell(c, text, fill=fill_head):
    c.value, c.font, c.fill, c.alignment, c.border = text, f_head, fill, wrap_center, border


def add_name(wb, name, ref):
    dn = DefinedName(name, attr_text=xl(ref))
    try:
        wb.defined_names[name] = dn
    except TypeError:
        wb.defined_names.append(dn)


def make_table(ws, tname, top, left_col, cols, rows, style="TableStyleLight9"):
    """cols: [(머리글, 너비, 'in'|'calc', 수식 또는 None, 숫자형식 또는 None)], rows: 입력 열 값 목록(dict)"""
    n = max(1, len(rows))
    tcols = []
    for j, (h, w, kind, fml, fmt) in enumerate(cols):
        c = ws.cell(row=top, column=left_col + j)
        head_cell(c, h, fill_head if kind == "in" else fill_head2)
        ws.column_dimensions[CL(left_col + j)].width = w
        tc = TableColumn(id=j + 1, name=h)
        if kind == "calc":
            body = xl(this_row(tname, fml))
            tc.calculatedColumnFormula = TableFormula(attr_text=body[1:] if body.startswith("=") else body)
        tcols.append(tc)
        for i in range(n):
            cell = ws.cell(row=top + 1 + i, column=left_col + j)
            if kind == "calc":
                cell.value = xl(this_row(tname, fml))
                cell.font, cell.fill = f_base, fill_calc
            else:
                v = rows[i].get(h) if i < len(rows) else None
                if v is not None:
                    cell.value = v
                cell.font = f_input
            cell.border = border
            if fmt:
                cell.number_format = fmt
    ref = f"{CL(left_col)}{top}:{CL(left_col + len(cols) - 1)}{top + n}"
    t = Table(displayName=tname, ref=ref, tableColumns=tcols)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    t._initialise_columns = lambda: None  # 직접 만든 열 정의 유지
    ws.add_table(t)
    return t


def dv(ws, sqref, formula1, title, msg, kind="list"):
    d = DataValidation(type=kind, formula1=xl(formula1), allow_blank=True, showErrorMessage=True,
                       errorStyle="stop", errorTitle=title, error=msg)
    ws.add_data_validation(d)
    d.add(sqref)


wb = Workbook()
wb.remove(wb.active)
MD = "DAY(EOMONTH(급여기준월,0))"

# ================================================================ 설정_코드표
ref = wb.create_sheet("설정_코드표")
if DATA:
    HOSPITALS, DEPTS, INTERN_ONLY = DATA["codes"]["병원표"], DATA["codes"]["진료과표"], DATA["codes"]["인턴구분표"]
    DEPT_ALIASES, HOSP_ALIASES = DATA["codes"]["과약칭표"], DATA["codes"]["병원약칭표"]
code_tables = [
    ("병원표", "A", [("병원", 14)], [{"병원": h} for h in HOSPITALS]),
    ("진료과표", "C", [("근무과", 18)], [{"근무과": d} for d in DEPTS]),
    ("인턴구분표", "E", [("인턴 구분", 12)], [{"인턴 구분": d} for d in INTERN_ONLY]),
    ("과약칭표", "I", [("약칭", 16), ("정식 명칭", 18)], [{"약칭": a, "정식 명칭": b} for a, b in DEPT_ALIASES]),
    ("병원약칭표", "L", [("표기", 16), ("정식 병원", 12)], [{"표기": a, "정식 병원": b} for a, b in HOSP_ALIASES]),
]
from openpyxl.utils import column_index_from_string as CI
for tname, col, cdefs, rows in code_tables:
    make_table(ref, tname, 1, CI(col), [(h, w, "in", None, None) for h, w in cdefs], rows, style="TableStyleLight1")
for c in "BDFGHKN":
    ref.column_dimensions[c].width = 3
ref.column_dimensions["O"].width = 26
ref.column_dimensions["P"].width = 14
notes = [
    "※ 모든 목록은 엑셀 '표'입니다. 표 바로 아래 행에 입력하면 자동으로 목록이 늘어나고 드롭다운·판정에 바로 반영됩니다.",
    "※ 병원약칭표(L~M): 인턴근무계획의 지역 단어와 기본자료의 근무지를 정식 병원으로 바꿉니다. 실제 기관명은 여기에 추가하세요.",
    "※ 과약칭표(I~J): 인턴근무계획 월 칸의 마지막 단어(진료과 약칭)를 정식 명칭으로 바꿉니다.",
    "※ 인턴구분표(E): 통합(과 미지정), 핵/방, 직/피, 정/영처럼 인턴 순환에만 쓰는 구분입니다.",
]
for i, t in enumerate(notes):
    ref[f"O{1 + i}"].value, ref[f"O{1 + i}"].font = t, f_base
head_cell(ref["O8"], "설정값 (노란 칸만 수정)")
head_cell(ref["P8"], "값")
for rr, (label, val, fmt, nm) in enumerate([
        ("급여 기준월(1일)", DATA["settings"][0] if DATA else "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "yyyy-mm-dd", "급여기준월"),
        ("서울·구리 15일 기준(일)", DATA["settings"][1] if DATA else 15, "0", "판정기준일")], 9):
    ref[f"O{rr}"].value, ref[f"O{rr}"].font, ref[f"O{rr}"].border = label, f_bold, border
    c = ref[f"P{rr}"]
    c.value, c.number_format, c.font, c.fill, c.border, c.alignment = val, fmt, f_input, fill_key, border, center
    add_name(wb, nm, f"{q('설정_코드표')}!$P${rr}")
ref["O11"].value, ref["O12"].font = "※ 급여 기준월: 다른 달을 보려면 그 달 1일 날짜를 입력(기본 수식 = 이번 달 1일).", f_note
for nm, r in [("기준월키", 'TEXT(급여기준월,"yyyy-mm")'), ("전월키", 'TEXT(EDATE(급여기준월,-1),"yyyy-mm")'),
              ("익월키", 'TEXT(EDATE(급여기준월,1),"yyyy-mm")')]:
    add_name(wb, nm, r)
for nm, r in [("병원목록", "병원표[병원]"), ("진료과목록", "진료과표[근무과]"),
              ("병원표기목록", "병원약칭표[표기]"), ("과약칭목록", "과약칭표[약칭]")]:
    add_name(wb, nm, r)

# ================================================================ 입력_기본자료
bs = wb.create_sheet("입력_기본자료")
base_rows = [
    ("재직", "2261212", "김ㅇㅇ", "가정의학과", "레지던트1", "서울", 20260301, "11-1111", 111111, "여", 19999999),
    ("재직", "2261213", "이ㅇㅇ", "내과", "레지던트2", "서울", 20250301, "11-1112", 111112, "남", 19950702),
    ("재직", "2261214", "박ㅇㅇ", "내과", "레지던트3", "서울", 20240301, "21-1111", 111113, "여", 19970125),
    ("재직", "2261215", "최ㅇㅇ", "산부인과", "레지던트3", "서울", 20240301, "11-1114", 111114, "남", 19961109),
    ("재직", "2261216", "정ㅇㅇ", "마취통증의학과", "레지던트2", "구리", 20250301, "11-1115", 111115, "여", 19980530),
    ("재직", "222222", "김ㅇㅇ", "인턴", "인턴", "서울", 20260301, "22-2222", 222222, "여", 19990101),
    ("재직", "222223", "이ㅇㅇ", "인턴", "인턴", "서울", 20260301, "22-2223", 222223, "남", 19990202),
    ("재직", "222224", "박ㅇㅇ", "인턴", "인턴", "구리", 20260301, "22-2224", 222224, "남", 19990303),
]
if DATA:
    base_rows = DATA["base_rows"]
bh = ["현재근무 여부", "사번", "이름", "수련과목", "연차", "근무지", "수련개시일", "전공의등록번호", "의사면허번호", "성별", "생년월일", "비고"]
bw = [11, 11, 10, 14, 11, 9, 12, 14, 12, 6, 12, 16]
bcols = [(h, w, "in", None, "@" if h in ("사번", "전공의등록번호") else None) for h, w in zip(bh, bw)]
bcols += [("사번키", 12, "calc", '=IF(TRIM([@사번]&"")="","","S|"&TRIM([@사번]&""))', None)]
make_table(bs, "기본자료표", 1, 1, bcols, [dict(zip(bh, r)) for r in base_rows])
bs.freeze_panes = "A2"

# ================================================================ 입력_인턴근무계획
ip = wb.create_sheet("입력_인턴근무계획")
ip.column_dimensions["A"].width = 2
ip.merge_cells("B1:K1")
ip["B1"].value, ip["B1"].font, ip["B1"].alignment = (DATA["plan_title"] if DATA else "2026년도 하반기 인턴근무계획표(9~2월)"), f_title, center
plan_rows = [
    (1, "김ㅇㅇ", 222222, "구리 안과", "구리 내과", "서울 핵방", "서울 재활의학과", "제주 산부", "서울 소아과", 2026.03),
    (2, "이ㅇㅇ", 222223, "서울 내과", "서울 외과", "구리 통합", "인천 산부인과", "제주 에스중앙 내과", "서울 가정의학과", 2026.03),
    (3, "박ㅇㅇ", 222224, "제주 외과", "서울 내과", "서울 응급", "구리 직피", "서울 정영", "제주 산부", 2026.03),
]
ph = ["순번", "이름", "사번", "9월", "10월", "11월", "12월", "1월", "2월", "비고"]
if DATA:
    ph, plan_rows = DATA["plan_headers"], DATA["plan_rows"]
pcols = [(h, w, "in", None, None) for h, w in zip(ph, [6, 10, 11, 17, 17, 17, 17, 17, 17, 10])]
pcols.append(("사번키", 11, "calc", '=IF(TRIM([@사번]&"")="","","S|"&TRIM([@사번]&""))', None))
for k in range(1, 7):
    pcols.append((f"정규화{k}", 17, "calc",
                  f'=LET(v_r,TRIM(INDEX(인턴계획표[#This Row],1,{3 + k})&""),'
                  f'v_t1,IFERROR(LEFT(v_r,FIND(" ",v_r)-1),""),'
                  f'v_t2,TRIM(RIGHT(SUBSTITUTE(v_r," ",REPT(" ",60)),60)),'
                  f'v_h,XLOOKUP(v_t1,병원약칭표[표기],병원약칭표[정식 병원],""),'
                  f'v_d,XLOOKUP(v_t2,과약칭표[약칭],과약칭표[정식 명칭],""),'
                  f'IF(v_r="","",IF(OR(v_t1="",v_h="",v_d=""),"⚠"&v_r,v_h&" "&v_d)))', None))
pcols.append(("입사일(비고)", 12, "calc",
              '=LET(v_k,INDEX(인턴계획표[#This Row],1,10)&"",v_v,IFERROR(VALUE(v_k),""),'
              'v_m,IFERROR(ROUND((v_v-INT(v_v))*100,0),0),'
              'IF(v_k="","",IF(AND(ISNUMBER(v_v),v_m>=1,v_m<=12,IFERROR(INT(v_v),0)>=1900),DATE(INT(v_v),v_m,1),v_k)))', "yyyy-mm-dd"))
pcols.append(("입력오류", 8, "calc", '=COUNTIF(인턴계획표[[#This Row],[정규화1]:[정규화6]],"⚠*")', "0"))
make_table(ip, "인턴계획표", 2, 2, pcols, [dict(zip(ph, r)) for r in plan_rows])
ip.freeze_panes = "E3"
t = 'TRIM(E3)'
t1 = f'LEFT({t},FIND(" ",{t})-1)'
dv(ip, f"E3:J{MAXR}",
   f'OR({t}="",AND(ISNUMBER(FIND(" ",{t})),COUNTIF(병원표기목록,{t1})>0,'
   f'COUNTIF(과약칭목록,TRIM(RIGHT(SUBSTITUTE({t}," ",REPT(" ",60)),60)))>0))',
   "등록되지 않은 병원/진료과", "'지역 진료과' 형식(예: 서울 내과, 제주 산부)으로 입력하세요.\n"
   "처음 보는 지역이나 진료과(약칭)라면 먼저 '설정_코드표'에 추가한 뒤 다시 입력하세요.", kind="custom")
ip.conditional_formatting.add(f"E3:J{MAXR}", FormulaRule(formula=['COUNTIF($M3:$R3,"⚠"&TRIM(E3))>0'],
                              fill=PatternFill("solid", bgColor="FFC7CE")))
ip.conditional_formatting.add(f"M3:R{MAXR}", FormulaRule(formula=['LEFT(M3,1)="⚠"'], font=Font(name=FONT, color="C00000", bold=True)))

# ================================================================ 입력_구리근무자명단 (양식 그대로 붙여넣기)
gi = wb.create_sheet("입력_구리근무자명단")
gi["A1"].value, gi["A1"].font = "전공의(모.자 파견) 근무자 명단", f_title
gi["A2"].value, gi["A2"].font = "2026년도", f_bold
g_heads = ["소속", "년차"]
for y, m in [(2026, m) for m in range(9, 13)] + [(2027, m) for m in range(1, 9)]:
    g_heads += [f"{y}년{m}월\n근무예정자", f"{m}월\n근무자", "근무 시작일자\n(시작일자 작성)", "근무 종료일자\n(종료일자 작성)",
                f"{m}월 휴가 스케줄\n(휴가일 작성)", "비고"]
if DATA:
    gi["A1"].value = gi["A2"].value = None
    for (r, c), v in DATA["guri_cells"].items():
        cell = gi.cell(row=r, column=c, value=v)
        cell.font = f_input if r >= 4 else (f_title if r == 1 else f_bold)
    g_heads = [DATA["guri_cells"].get((3, c)) for c in range(1, 1 + max(c for (r, c) in DATA["guri_cells"] if r == 3))]
for i, h in enumerate(g_heads, 1):
    head_cell(gi.cell(row=3, column=i), h)
    gi.column_dimensions[CL(i)].width = 11
gi.column_dimensions["A"].width = 14
gi.column_dimensions["B"].width = 6
gi.row_dimensions[3].height = 44
g_sample = [
    ("내과", 3, ["박ㅇㅇ", None, None, None, None, None], ["박ㅇㅇ", "박ㅇㅇ", "2026.10.01", "2026.10.31", None, None]),
    ("가정의학과", 1, [None] * 6, ["김ㅇㅇ", "김ㅇㅇ", "2026.10.21", "2026.10.31", None, None]),
    ("마취통증의학과", 2, ["정ㅇㅇ", None, None, None, None, "구리독자"], ["정ㅇㅇ", "정ㅇㅇ", "2026.10.01", "2026.10.17", None, "구리독자"]),
    ("외과", 1, ["홍ㅇㅇ", None, None, None, None, "구리독자"], ["홍ㅇㅇ", "홍ㅇㅇ", "2026.10.01", "2026.10.31", None, "구리독자"]),
    ("수련교육부", "인턴", [None] * 6, ["김ㅇㅇ", "김ㅇㅇ", "2026.10.01", "2026.10.31", None, None]),
]
for k, (dept, yr, sep, octb) in enumerate([] if DATA else g_sample):
    r = 4 + k
    gi.cell(row=r, column=1, value=dept).font = f_input
    gi.cell(row=r, column=2, value=yr).font = f_input
    for j, v in enumerate(sep + octb):
        if v is not None:
            gi.cell(row=r, column=3 + j, value=v).font = f_input
gi.freeze_panes = "C4"

# ================================================================ 입력_월별근무지 (전월·익월 등 근무지 직접 입력)
dw = wb.create_sheet("입력_월별근무지")
dw["A1"].value, dw["A1"].font = "월별 근무지 직접 입력 (자동 판정보다 우선)", f_title
dw["A2"].value, dw["A2"].font = ("※ 그 달 근무병원을 직접 정할 사람만 한 줄씩: 월(예: 2026-10)·사번·근무병원. 급여지급병원이 근무병원과 다를 때만 급여지급병원도 입력. "
                                 "급여 기준월·전월·익월 계산과 출력_기관명단에 그대로 쓰입니다. 같은 사번·월이 두 줄이면 위쪽 줄이 적용됩니다."), f_note
dw_rows = DATA.get("direct_rows", []) if DATA else [
    {"월": date(2026, 9, 1), "사번": "2261216", "근무병원": "구리병원", "비고": "예시: 9월 구리 근무"},
    {"월": date(2026, 11, 1), "사번": "2261212", "근무병원": "구리병원", "비고": "예시: 11월 구리 파견 예정"},
    {"월": date(2026, 11, 1), "사번": "2261214", "근무병원": "구리병원", "급여지급병원": "서울병원", "비고": "예시: 구리 근무, 서울 지급"},
]
dw_cols = [("월", 9, "in", None, "yyyy-mm"), ("사번", 11, "in", None, "@"), ("근무병원", 11, "in", None, None),
           ("급여지급병원", 12, "in", None, None), ("비고", 26, "in", None, None),
           ("월키", 9, "calc", '=LET(v_m,[@월],IF(v_m&""="","",IF(ISNUMBER(v_m),IF(v_m<10000,'
            'TEXT(DATE(INT(v_m),ROUND(MOD(v_m,1)*100,0),1),"yyyy-mm"),TEXT(v_m,"yyyy-mm")),'
            'IFERROR(TEXT(DATEVALUE(SUBSTITUTE(SUBSTITUTE(LEFT(TRIM(v_m),7),".","-"),"/","-")&"-01"),"yyyy-mm"),"⚠형식 오류"))))', None),
           ("사번키", 11, "calc", '=IF(TRIM([@사번]&"")="","","S|"&TRIM([@사번]&""))', None),
           ("키", 16, "calc", '=IF(LEFT([@월키],1)="⚠","",IF(OR([@사번키]="",[@월키]=""),"",[@사번키]&"|"&[@월키]))', None),
           ("이름(확인)", 9, "calc", '=IF([@사번키]="","",XLOOKUP([@사번키],기본자료표[사번키],기본자료표[이름]&"","(기본자료에 없음)"))', None),
           ("적용 급여병원", 12, "calc", '=IF([@근무병원]="","",IF([@급여지급병원]="",[@근무병원],[@급여지급병원]))', None),
           ("점검", 26, "calc",
            '=IF(AND(TRIM([@사번]&"")="",[@월]&""=""),"",IF(OR(TRIM([@사번]&"")="",[@월]&""=""),"월·사번 모두 입력",'
            'IF(LEFT([@월키],1)="⚠","월 형식 확인(예: 2026-10)",IF([@근무병원]="","근무병원 입력",'
            'IF(COUNTIF(명단표[사번],TRIM([@사번]&""))=0,"전공의명단에 없는 사번",'
            'IF(COUNTIF(월별근무지표[키],[@키])>1,"같은 사번·월 중복(위쪽 줄 적용)","정상"))))))', None)]
make_table(dw, "월별근무지표", 3, 1, dw_cols, dw_rows)
dw.freeze_panes = "A4"
for c in ("C", "D"):
    dv(dw, f"{c}4:{c}{MAXR}", "=병원목록", "등록되지 않은 병원", "병원 목록에 없습니다. 새 병원이면 '설정_코드표' 병원표에 먼저 추가하세요.")
dw.conditional_formatting.add(f"K4:K{MAXR}", FormulaRule(formula=['AND($K4<>"",$K4<>"정상")'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))

# ================================================================ 입력_전공의명단
ws = wb.create_sheet("입력_전공의명단")
ROW_ON = 'TRIM([@사번]&"")<>""'


def lk(col):
    return f'=IF([@기본자료행]="","",IF(INDEX(기본자료표[{col}],[@기본자료행])="","",INDEX(기본자료표[{col}],[@기본자료행])))'


def lk_date(col):
    return (f'=IF([@기본자료행]="","",LET(v_x,INDEX(기본자료표[{col}],[@기본자료행]),IF(v_x="","",IF(NOT(ISNUMBER(v_x)),v_x&"",'
            f'IF(v_x<1000000,v_x,IF(AND(v_x>=10000101,INT(MOD(v_x,10000)/100)>=1,INT(MOD(v_x,10000)/100)<=12,'
            f'MOD(v_x,100)>=1,MOD(v_x,100)<=31),DATE(INT(v_x/10000),INT(MOD(v_x,10000)/100),MOD(v_x,100)),v_x&""))))))')


home_n = 'IF([@근무지]="","",XLOOKUP(TRIM([@근무지]&""),병원약칭표[표기],병원약칭표[정식 병원],[@근무지]&""))'
ok_now = 'AND([@인턴 이번달]<>"",LEFT([@인턴 이번달],1)<>"⚠")'
ok_mon = 'AND([@인턴 기준월]<>"",LEFT([@인턴 기준월],1)<>"⚠")'
seoul, guri, th = "[@서울 근무일(기준월)]", "[@구리 근무일(기준월)]", "판정기준일"
DIRECT = lambda m: f'INDEX(월별근무지표[근무병원],[@직접행({m})])'
DIRECT_PAY = lambda m: f'INDEX(월별근무지표[적용 급여병원],[@직접행({m})])'
DIRECT_NOTE = lambda m: (f'"직접 입력(입력_월별근무지)"&IF(INDEX(월별근무지표[비고],[@직접행({m})])&""="","",'
                         f'" — "&INDEX(월별근무지표[비고],[@직접행({m})]))')


def side_month(m):
    """전월·익월: 직접 입력 → 인턴 근무계획 → 소속병원 순서."""
    ok = f'AND([@인턴 {m}]<>"",LEFT([@인턴 {m}],1)<>"⚠")'
    return [
        (f"{m} 근무병원", 10, "calc",
         f'=IF(NOT({ROW_ON}),"",IF([@직접행({m})]<>"",{DIRECT(m)},IF({ok},LEFT([@인턴 {m}],FIND(" ",[@인턴 {m}])-1),[@소속병원])))', None),
        (f"{m} 급여지급병원", 10, "calc",
         f'=IF(NOT({ROW_ON}),"",IF([@직접행({m})]<>"",{DIRECT_PAY(m)},[@{m} 근무병원]))', None),
        (f"{m} 근거", 28, "calc",
         f'=IF(NOT({ROW_ON}),"",IF([@직접행({m})]<>"",{DIRECT_NOTE(m)},IF({ok},"인턴 근무계획 → 그 달 근무병원",'
         f'IF([@구리행]<>"","⚠자동(소속병원) — 급여 기준월에 구리 근무, 직접 입력 확인","자동(소속병원) — 구리 근무자명단 미반영"))))', None),
    ]

mcols = [
    ("번호", 6, "calc", '=ROW()-ROW(명단표[#Headers])', None),
    ("사번", 11, "in", None, "@"),
    ("이름", 9, "calc", lk("이름"), None),
    ("성별", 5, "calc", lk("성별"), None),
    ("생년월일", 11, "calc", lk_date("생년월일"), "yyyy-mm-dd"),
    ("수련과목", 13, "calc", lk("수련과목"), None),
    ("연차", 10, "calc", lk("연차"), None),
    ("근무지", 8, "calc", lk("근무지"), None),
    ("수련개시일", 11, "calc", lk_date("수련개시일"), "yyyy-mm-dd"),
    ("전공의등록번호", 12, "calc", lk("전공의등록번호"), None),
    ("의사면허번호", 11, "calc", lk("의사면허번호"), None),
    ("재직여부", 8, "calc", lk("현재근무 여부"), None),
    ("입사장소", 10, "in", None, None),
    ("현재 근무병원", 11, "calc",
     f'=IF(NOT({ROW_ON}),"",IF({ok_now},LEFT([@인턴 이번달],FIND(" ",[@인턴 이번달])-1),'
     f'IF([@구리 근무 중]="Y","구리병원",{home_n})))', None),
    ("현재 진료과", 14, "calc",
     f'=IF(NOT({ROW_ON}),"",IF({ok_now},MID([@인턴 이번달],FIND(" ",[@인턴 이번달])+1,60),'
     f'IF([@수련과목]="","",IF([@수련과목]="인턴","수련교육부(인턴)",[@수련과목]))))', None),
    ("구리 근무 시작일(기준월)", 11, "calc", '=IF([@구리행]="","",INDEX(구리_시작,[@구리행]))', None),
    ("구리 근무 종료일(기준월)", 11, "calc", '=IF([@구리행]="","",INDEX(구리_종료,[@구리행]))', None),
    ("구리독자", 8, "calc", '=IF([@구리행]="","",INDEX(구리_독자,[@구리행]))', None),
    ("서울 근무일(기준월)", 9, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@구리행]<>"",MAX(0,{MD}-N([@구리명단 근무일])),IF([@기준월 근무병원]="서울병원",{MD},0)))', "0"),
    ("구리 근무일(기준월)", 9, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@구리행]<>"",N([@구리명단 근무일]),IF([@기준월 근무병원]="구리병원",{MD},0)))', "0"),
    ("급여지급병원(기준월)", 13, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@직접행(기준월)]<>"",{DIRECT_PAY("기준월")},IF([@기준월 근무병원]="","",IF([@구리행]="",[@기준월 근무병원],'
     f'IF(AND({seoul}>={th},{guri}<{th}),"서울병원",IF(AND({guri}>={th},{seoul}<{th}),"구리병원",'
     f'"⚠15일 판정불가(서울 "&{seoul}&"일/구리 "&{guri}&"일)"))))))', None),
    ("급여 판정 근거", 30, "calc",
     f'=IF(OR(NOT({ROW_ON}),[@급여지급병원(기준월)]=""),"",IF([@직접행(기준월)]<>"",{DIRECT_NOTE("기준월")},IF([@구리행]<>"",'
     f'"서울 "&{seoul}&"일 / 구리 "&{guri}&"일 (구리 근무자명단) → 15일 이상 근무 병원 지급",'
     f'IF({ok_mon},"인턴 근무계획 → 그 달 근무병원 지급","구리 근무자명단에 없음 → 소속병원 지급"))))', None),
    ("점검", 26, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@기본자료행]="","기본자료에 없는 사번",IF(COUNTIF(명단표[사번],[@사번])>1,"중복 사번",'
     f'IF([@소속병원]="","소속(입사장소) 입력",'
     f'IF(LEFT([@급여지급병원(기준월)],1)="⚠","서울·구리 15일 판정 불가(양쪽 모두 이상/미만)",'
     f'IF(IF([@인턴계획행]="",0,N(INDEX(인턴계획표[입력오류],[@인턴계획행])))>0,"인턴계획 입력오류 확인",'
     f'IF(AND([@구리행]="",COUNTIF(구리_이름,[@이름])>0),"구리 근무자명단 연결 확인(동명이인·과 불일치)","정상")))))))', None),
    ("비고", 16, "in", None, None),
] + side_month("전월") + side_month("익월") + [
    # ---- 보조 열 (수정 금지)
    ("기본자료행", 8, "calc", '=IF(NOT(' + ROW_ON + '),"",IFERROR(XMATCH("S|"&TRIM([@사번]&""),기본자료표[사번키]),""))', None),
    ("이름|과", 14, "calc", '=IF([@이름]="","",[@이름]&"|"&[@수련과목])', None),
    ("인턴계획행", 8, "calc", '=IF(NOT(' + ROW_ON + '),"",IFERROR(XMATCH("S|"&TRIM([@사번]&""),인턴계획표[사번키]),""))', None),
    ("인턴 이번달", 15, "calc",
     '=IF(OR([@인턴계획행]="",인턴위치_이번=""),"",INDEX(인턴계획표[[정규화1]:[정규화6]],[@인턴계획행],인턴위치_이번)&"")', None),
    ("인턴 기준월", 15, "calc",
     '=IF(OR([@인턴계획행]="",인턴위치_기준월=""),"",INDEX(인턴계획표[[정규화1]:[정규화6]],[@인턴계획행],인턴위치_기준월)&"")', None),
    ("인턴 전월", 15, "calc",
     '=IF(OR([@인턴계획행]="",인턴위치_전월=""),"",INDEX(인턴계획표[[정규화1]:[정규화6]],[@인턴계획행],인턴위치_전월)&"")', None),
    ("인턴 익월", 15, "calc",
     '=IF(OR([@인턴계획행]="",인턴위치_익월=""),"",INDEX(인턴계획표[[정규화1]:[정규화6]],[@인턴계획행],인턴위치_익월)&"")', None),
] + [(f"직접행({m})", 7, "calc",
      f'=IF(NOT({ROW_ON}),"",IFERROR(XMATCH("S|"&TRIM([@사번]&"")&"|"&{m}키,월별근무지표[키]),""))', None)
     for m in ("기준월", "전월", "익월")] + [
    ("소속병원", 10, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@입사장소]<>"",[@입사장소],IF([@구리독자]="구리독자","구리병원",{home_n})))', None),
    ("기준월 근무병원", 10, "calc",
     f'=IF(NOT({ROW_ON}),"",IF([@직접행(기준월)]<>"",{DIRECT("기준월")},IF({ok_mon},LEFT([@인턴 기준월],FIND(" ",[@인턴 기준월])-1),[@소속병원])))', None),
    ("구리행", 7, "calc",
     '=IF([@이름]="","",LET(v_a,XMATCH([@이름|과],구리_키),IF(ISNUMBER(v_a),v_a,'
     'IF(AND(COUNTIF(구리_이름,[@이름])=1,COUNTIF(명단표[이름],[@이름])=1),XMATCH([@이름],구리_이름),""))))', None),
    ("구리명단 근무일", 8, "calc", '=IF([@구리행]="","",INDEX(구리_근무일,[@구리행]))', "0"),
    ("구리 근무 중", 7, "calc",
     '=IF([@구리행]="","",LET(v_s,IFERROR(DATEVALUE([@구리 근무 시작일(기준월)]),""),v_e,IFERROR(DATEVALUE([@구리 근무 종료일(기준월)]),""),'
     'IF(AND(ISNUMBER(v_s),ISNUMBER(v_e)),IF(AND(TODAY()>=v_s,TODAY()<=v_e),"Y",""),"")))', None),
]
m_rows = DATA["roster_rows"] if DATA else [{"사번": s, "입사장소": h, "비고": n} for s, h, n in [
    ("2261212", "서울병원", "예시 데이터"), ("2261213", "서울병원", "예시 데이터"), ("2261214", "서울병원", "예시 데이터"),
    ("2261215", "서울병원", "예시 데이터"), ("2261216", "구리병원", "예시 데이터"), ("222222", "서울병원", "예시(인턴)"),
    ("222223", "서울병원", "예시(인턴)"), ("222224", "구리병원", "예시(인턴)")]]
make_table(ws, "명단표", 1, 1, mcols, m_rows)
MC = {h: CL(i + 1) for i, (h, *_rest) in enumerate(mcols)}
ws.row_dimensions[1].height = 42
ws.freeze_panes = "D2"
dv(ws, f"{MC['입사장소']}2:{MC['입사장소']}{MAXR}", "=병원목록", "등록되지 않은 값",
   "병원 목록에 없습니다. 새 병원이면 '설정_코드표' 병원표에 먼저 추가하세요.")
cf = ws.conditional_formatting
B, Z, Y, AE = MC["사번"], MC["점검"], MC["급여지급병원(기준월)"], MC["소속병원"]
cf.add(f"{B}2:{B}{MAXR}", FormulaRule(formula=[f'AND(${B}2<>"",${MC["기본자료행"]}2="")'], fill=PatternFill("solid", bgColor="FF0000"), font=Font(name=FONT, color="FFFFFF", bold=True)))
cf.add(f"{Z}2:{Z}{MAXR}", FormulaRule(formula=[f'AND(${Z}2<>"",${Z}2<>"정상")'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
cf.add(f"{Y}2:{Y}{MAXR}", FormulaRule(formula=[f'LEFT(${Y}2,1)="⚠"'], fill=PatternFill("solid", bgColor="FF0000"), font=Font(name=FONT, color="FFFFFF", bold=True)))
cf.add(f"{Y}2:{Y}{MAXR}", FormulaRule(formula=[f'AND(${Y}2<>"",${AE}2<>"",${Y}2<>${AE}2)'], fill=PatternFill("solid", bgColor="FFFF00")))
cf.add(f"{MC['생년월일']}2:{MC['생년월일']}{MAXR}", FormulaRule(formula=[f'ISTEXT(${MC["생년월일"]}2)'], font=Font(name=FONT, color="FF0000", bold=True)))

# ================================================================ 출력_서울급여
sp = wb.create_sheet("출력_서울급여")
sp["A1"].value, sp["A1"].font = "서울병원 급여 지급 대상자", f_title
sp["A2"].value, sp["A2"].font = "급여 기준월", f_bold
sp["B2"].value, sp["B2"].number_format, sp["B2"].font = "=급여기준월", 'yyyy"년" m"월"', Font(name=FONT, size=12, bold=True, color="1F3864")
YN = "명단표[급여지급병원(기준월)]"
for lc, lt, vc, vf in [("D2", "서울 지급 인원", "E2", f'=COUNTIF({YN},"서울병원")'),
                       ("F2", "그중 타 병원 소속", "G2", f'=COUNTIFS({YN},"서울병원",명단표[소속병원],"<>서울병원")'),
                       ("H2", "⚠ 판정 불가", "I2", f'=COUNTIF({YN},"⚠*")')]:
    sp[lc].value, sp[lc].font, sp[lc].alignment = lt, f_bold, Alignment(horizontal="right", vertical="center")
    sp[vc].value, sp[vc].font, sp[vc].alignment, sp[vc].border = vf, Font(name=FONT, size=12, bold=True), center, border
sp.conditional_formatting.add("I2", FormulaRule(formula=["I2>0"], fill=PatternFill("solid", bgColor="FF0000"), font=Font(name=FONT, color="FFFFFF", bold=True)))
sp["A3"].value, sp["A3"].font = ("※ 자동으로 펼쳐지는 목록입니다(입력 금지). 노란 행 = 소속은 다른 병원인데 이번 달 서울병원이 지급. "
                                 "급여 기준월은 설정_코드표에서 바꿉니다."), f_note
sp_heads = [("번호", 5), ("이름", 9), ("사번", 10), ("전공의등록번호", 12), ("수련과목", 13), ("연차", 9), ("소속병원", 10),
            ("구분", 24), ("현재 근무병원", 11), ("구리 근무 시작일", 11), ("구리 근무 종료일", 11),
            ("서울 근무일", 8), ("구리 근무일", 8), ("판정 근거", 40), ("점검", 26)]
for i, (h, w) in enumerate(sp_heads, 1):
    head_cell(sp.cell(row=5, column=i), h)
    sp.column_dimensions[CL(i)].width = w
sp.row_dimensions[5].height = 30
cols_sp = ["명단표[이름]", '명단표[사번]&""', '명단표[전공의등록번호]&""', '명단표[수련과목]&""', '명단표[연차]&""',
           "명단표[소속병원]",
           'IF(명단표[소속병원]="서울병원","서울 소속","타 병원 소속("&명단표[소속병원]&") → 서울 지급")',
           "명단표[현재 근무병원]", '명단표[구리 근무 시작일(기준월)]&""', '명단표[구리 근무 종료일(기준월)]&""',
           "명단표[서울 근무일(기준월)]", "명단표[구리 근무일(기준월)]", "명단표[급여 판정 근거]", "명단표[점검]"]
idx = ",".join(str(i) for i in range(1, len(cols_sp) + 1))
spill(sp, "A6", f'=IF(COUNTIF({YN},"서울병원")=0,"",SEQUENCE(COUNTIF({YN},"서울병원")))')
spill(sp, "B6", f'=IF(COUNTIF({YN},"서울병원")=0,"대상 없음",FILTER(CHOOSE({{{idx}}},{",".join(cols_sp)}),{YN}="서울병원"))', width=len(cols_sp))
last = CL(len(sp_heads))
sp.conditional_formatting.add(f"A6:{last}{SPILL_END}", FormulaRule(formula=['LEFT($H6,2)="타 "'], fill=PatternFill("solid", bgColor="FFF2CC")))
sp.conditional_formatting.add(f"{last}6:{last}{SPILL_END}", FormulaRule(formula=[f'AND(${last}6<>"",${last}6<>"정상")'], fill=PatternFill("solid", bgColor="FFC7CE")))
e0 = len(sp_heads) + 2
sp.column_dimensions[CL(e0 - 1)].width = 3
sp.cell(row=4, column=e0, value="⚠ 서울·구리 15일 판정 불가 (양쪽 모두 15일 이상/미만) — 담당자 확인 필요").font = Font(name=FONT, size=10, bold=True, color="C00000")
for i, (h, w) in enumerate([("이름", 9), ("사번", 10), ("소속병원", 10), ("서울 근무일", 8), ("구리 근무일", 8)]):
    head_cell(sp.cell(row=5, column=e0 + i), h, PatternFill("solid", fgColor="C00000"))
    sp.column_dimensions[CL(e0 + i)].width = w
spill(sp, f"{CL(e0)}6", f'=IF(COUNTIF({YN},"⚠*")=0,"없음",FILTER(CHOOSE({{1,2,3,4,5}},명단표[이름],명단표[사번]&"",명단표[소속병원],'
                        f'명단표[서울 근무일(기준월)],명단표[구리 근무일(기준월)]),LEFT({YN},1)="⚠"))', width=5)
sp.freeze_panes = "C6"

# ================================================================ 출력_기관명단 (시간외 판정 파일로 복사하는 용도)
ex = wb.create_sheet("출력_기관명단")
ex["A1"].value, ex["A1"].font = "기관명단 (전월·급여 기준월·익월 3개월) — 시간외근무 판정 파일의 입력_기관명단에 값으로 붙여넣기", f_title
ex["A2"].value, ex["A2"].font = ("※ A5부터 아래로 펼쳐진 전체를 한 번에 복사 → 시간외 판정 파일 입력_기관명단의 표 아래에 '값 붙여넣기'. "
                                 "급여 기준월은 자동 판정(구리 근무자명단·인턴계획·15일), 전월·익월은 입력_월별근무지의 직접 입력 → 인턴계획 → 소속병원 순서입니다. "
                                 "근거가 '⚠자동'인 줄은 직접 입력이 필요한지 확인하세요."), f_note
ex["A3"].value, ex["A3"].font = "자동 추정 확인 필요(⚠)", f_bold
ex["C3"].value = '=COUNTIF($I$5:$I$60000,"⚠*")'
ex["C3"].font, ex["C3"].border, ex["C3"].alignment = f_bold, border, center
ex.conditional_formatting.add("C3", FormulaRule(formula=["C3>0"], fill=PatternFill("solid", bgColor="FFC7CE")))
ex_heads = [("기준월", 9), ("사번", 10), ("이름", 9), ("수련과목", 13), ("연차", 10), ("소속병원", 10),
            ("기준월 근무병원", 11), ("급여지급병원(기준월)", 13), ("급여 판정 근거", 44)]
for i, (h, w) in enumerate(ex_heads, 1):
    head_cell(ex.cell(row=4, column=i), h)
    ex.column_dimensions[CL(i)].width = w
ex.row_dimensions[4].height = 30
by_m = lambda a, b, c: f"CHOOSE(v_m,INDEX(명단표[{a}],v_r),INDEX(명단표[{b}],v_r),INDEX(명단표[{c}],v_r))"
spill(ex, "A5", '=LET(v_k,TRIM(명단표[사번]&"")<>"",v_rows,FILTER(SEQUENCE(ROWS(명단표[사번])),v_k,0),v_n,ROWS(v_rows),'
                'v_i,SEQUENCE(3*v_n),v_m,INT((v_i-1)/v_n)+1,v_r,INDEX(v_rows,MOD(v_i-1,v_n)+1),'
                'IF(SUM(--v_k)=0,"명단 없음",CHOOSE({1,2,3,4,5,6,7,8,9},CHOOSE(v_m,전월키,기준월키,익월키),'
                'INDEX(명단표[사번],v_r)&"",INDEX(명단표[이름],v_r),INDEX(명단표[수련과목],v_r),INDEX(명단표[연차],v_r)&"",'
                'INDEX(명단표[소속병원],v_r),'
                + by_m("전월 근무병원", "기준월 근무병원", "익월 근무병원") + ","
                + by_m("전월 급여지급병원", "급여지급병원(기준월)", "익월 급여지급병원") + ","
                + by_m("전월 근거", "급여 판정 근거", "익월 근거") + ")))", width=9, height=3 * VERIFY_ROWS)
ex.conditional_formatting.add("A5:I60000", FormulaRule(formula=['LEFT($I5,1)="⚠"'], fill=PatternFill("solid", bgColor="FFF2CC")))
ex.conditional_formatting.add("A5:I60000", FormulaRule(formula=['AND($A5<>"",$A5=기준월키)'], font=Font(name=FONT, bold=True)))
ex.freeze_panes = "C5"

# ================================================================ 출력_구리근무자
go = wb.create_sheet("출력_구리근무자")
GIN = q("입력_구리근무자명단")
GD = f"{GIN}!$A$4:${GURI_COLS_END}${GURI_ROWS_END}"
go["A1"].value, go["A1"].font = "구리병원 근무자 (구리 근무자명단 기준)", f_title
go["A2"].value, go["A2"].font = "급여 기준월", f_bold
go["B2"].value, go["B2"].number_format, go["B2"].font = "=급여기준월", 'yyyy"년" m"월"', Font(name=FONT, size=12, bold=True, color="1F3864")
go["D2"].value, go["D2"].font = "블록 시작 열", f_note
spill(go, "E2",
      f'=LET(v_h,{GIN}!$A$3:${GURI_COLS_END}$3&"",v_p1,IFERROR(FIND("월",v_h),999),v_p2,IFERROR(FIND(CHAR(10),v_h),999),'
      f'v_e,IF(v_p1<v_p2,v_p1,v_p2),v_y,IFERROR(VALUE(LEFT(v_h,4)),0),v_m,IFERROR(VALUE(MID(v_h,6,v_e-6)),0),'
      f'v_k,IF(ISNUMBER(SEARCH("근무예정자",v_h)),v_y*100+v_m,0),IFERROR(XMATCH(YEAR(급여기준월)*100+MONTH(급여기준월),v_k),""))', height=1)
go["F2"].value, go["F2"].font = "근무자 확정", f_note
go["G2"].value = f'=IF($E$2="","",IF(COUNTIF(INDEX({GD},0,$E$2+1),"?*")>0,"확정","예정"))'
for c in ("E2", "G2"):
    go[c].font, go[c].alignment = f_bold, center
go["A3"].value, go["A3"].font = ("※ 자동 결과(입력 금지). 그 달 '근무자' 칸이 하나라도 채워져 있으면 근무자, 아니면 근무예정자 기준. 날짜가 없으면 1일~말일. "
                                 "명단 연결: 이름+진료과(인턴은 이름+인턴), 안 되면 이름만. '블록 시작 열'이 비면 3행 머리글(연·월)을 확인하세요."), f_note
for lc, lt, vc, vf in [("I1", "명단 연결", "J1", f'=COUNTIF($L$6:$L${SPILL_END},"명단 일치*")'),
                       ("I2", "명단에 없음/동명이인", "J2", f'=COUNTIF($L$6:$L${SPILL_END},"명단에 없음*")+COUNTIF($L$6:$L${SPILL_END},"동명이인*")'),
                       ("I3", "날짜 오류", "J3", f'=COUNTIF($F$6:$G${SPILL_END},"⚠*")')]:
    go[lc].value, go[lc].font, go[lc].alignment = lt, f_bold, Alignment(horizontal="right")
    go[vc].value, go[vc].font, go[vc].border, go[vc].alignment = vf, f_bold, border, center
go.conditional_formatting.add("J2:J3", FormulaRule(formula=["J2>0"], fill=PatternFill("solid", bgColor="FFC7CE")))
go_heads = [("원본행", 6), ("진료과", 14), ("연차", 6), ("이름", 9), ("확정/예정", 8), ("시작일", 11), ("종료일", 11),
            ("기준월 구리 근무일", 9), ("비고", 10), ("구리독자", 8), ("이름|과", 16), ("명단 연결", 24)]
for i, (h, w) in enumerate(go_heads, 1):
    head_cell(go.cell(row=5, column=i), h)
    go.column_dimensions[CL(i)].width = w
go.row_dimensions[5].height = 30
BOM = "급여기준월"
pd = lambda v: (f'IF({v}&""="",{{d}},IF(ISNUMBER({v}),{v},IFERROR(DATE(VALUE(LEFT({v},4)),VALUE(MID({v},6,2)),VALUE(MID({v},9,2))),-1)))')
spill(go, "A6",
      f'=IF($E$2="","해당 월 블록 없음",LET(v_g,{GD},v_lr,MAX(IFERROR(XMATCH("?*",INDEX(v_g,0,1)&"",2,-1),0),IFERROR(XMATCH("?*",INDEX(v_g,0,2)&"",2,-1),0),1),'
      f'v_d,INDEX(v_g,1,1):INDEX(v_g,v_lr,COLUMNS(v_g)),v_r,SEQUENCE(ROWS(v_d)),v_a,TRIM(INDEX(v_d,0,1)&""),v_b,INDEX(v_d,0,2)&"",'
      f'v_n,TRIM(INDEX(v_d,0,$E$2+IF($G$2="확정",1,0))&""),'
      f'v_dept,XLOOKUP(v_r,IF(v_a<>"",v_r,0),v_a,"",-1),'
      f'v_s0,INDEX(v_d,0,$E$2+2),v_e0,INDEX(v_d,0,$E$2+3),'
      f'v_s,{pd("v_s0").replace("{d}", BOM)},v_e,{pd("v_e0").replace("{d}", "EOMONTH(" + BOM + ",0)")},'
      f'v_lo,IF(v_s>{BOM},v_s,{BOM}),v_hi,IF(v_e<EOMONTH({BOM},0),v_e,EOMONTH({BOM},0)),'
      f'v_days,IF((v_s<0)+(v_e<0),0,IF(v_hi-v_lo+1>0,v_hi-v_lo+1,0)),'
      f'v_note,TRIM(INDEX(v_d,0,$E$2+5)&""),v_gd,IF(ISNUMBER(SEARCH("구리독자",v_note)),"구리독자",""),'
      f'v_dk,IF((v_b="인턴")+ISNUMBER(SEARCH("수련교육부",v_dept)),"인턴",v_dept),v_key,v_n&"|"&v_dk,'
      f'v_c1,COUNTIF(명단표[이름|과],v_key),v_c2,COUNTIF(명단표[이름],v_n),'
      f'v_st,IF(v_c1=1,"명단 일치",IF((v_c1=0)*(v_c2=1),"명단 일치(이름만, 과 다름)",IF(v_c2=0,'
      f'IF(COUNTIF(기본자료표[이름],v_n)>0,"명단에 없음(기본자료에는 있음)","명단에 없음"),"동명이인 확인"))),'
      f'v_st2,IF(v_s<0,"⚠"&v_s0,TEXT(v_s,"yyyy-mm-dd")),v_et2,IF(v_e<0,"⚠"&v_e0,TEXT(v_e,"yyyy-mm-dd")),'
      f'v_x,v_n<>"",IF(SUM(--v_x)=0,"해당 월 근무자 없음",'
      f'FILTER(CHOOSE({{1,2,3,4,5,6,7,8,9,10,11,12}},v_r+3,v_dept,v_b,v_n,$G$2,v_st2,v_et2,v_days,v_note,v_gd,v_key,v_st),v_x))))', width=12)
go.conditional_formatting.add(f"L6:L{SPILL_END}", FormulaRule(formula=['AND($L6<>"",LEFT($L6,5)<>"명단 일치")'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
go.conditional_formatting.add(f"F6:G{SPILL_END}", FormulaRule(formula=['LEFT(F6,1)="⚠"'], font=Font(name=FONT, color="C00000", bold=True)))
go.conditional_formatting.add(f"A6:K{SPILL_END}", FormulaRule(formula=['$J6="구리독자"'], fill=PatternFill("solid", bgColor="E2EFDA")))
go.freeze_panes = "E6"
GO = q("출력_구리근무자")
for nm, c in [("구리_키", "K"), ("구리_이름", "D"), ("구리_근무일", "H"), ("구리_독자", "J"), ("구리_시작", "F"), ("구리_종료", "G")]:
    add_name(wb, nm, f"{GO}!${c}$6:${c}${SPILL_END}")

# ================================================================ 출력_요약
sm = wb.create_sheet("출력_요약")
sm["A1"].value, sm["A1"].font = "현황 요약", f_title
for r, (lab, fml, fmt) in enumerate([("기준일", "=TODAY()", "yyyy-mm-dd"), ("급여 기준월", "=급여기준월", "yyyy-mm-dd"),
                                     ("15일 기준(일)", "=판정기준일", "0")], 2):
    sm[f"A{r}"].value, sm[f"A{r}"].font = lab, f_bold
    sm[f"B{r}"].value, sm[f"B{r}"].number_format, sm[f"B{r}"].font = fml, fmt, f_base
sm["C3"].value, sm["C3"].font = "← 설정값은 설정_코드표에서 변경", f_note
head_cell(sm["A6"], "인원 (급여 기준월)")
head_cell(sm["B6"], "인원")
people = [
    ("명단 등록 인원", '=SUMPRODUCT(--(TRIM(명단표[사번]&"")<>""))'),
    ("서울병원 지급", '=COUNTIF(명단표[급여지급병원(기준월)],"서울병원")'),
    ("구리병원 지급", '=COUNTIF(명단표[급여지급병원(기준월)],"구리병원")'),
    ("그 외 병원 지급", '=SUMPRODUCT((명단표[급여지급병원(기준월)]<>"")*(명단표[급여지급병원(기준월)]<>"서울병원")*(명단표[급여지급병원(기준월)]<>"구리병원")*(LEFT(명단표[급여지급병원(기준월)],1)<>"⚠"))'),
    ("⚠ 판정 불가", '=COUNTIF(명단표[급여지급병원(기준월)],"⚠*")'),
    ("구리 근무자명단 연결 인원", '=COUNT(명단표[구리행])'),
    ("급여 기준월 직접 입력 적용", '=COUNT(명단표[직접행(기준월)])'),
]
for i, (lab, fml) in enumerate(people):
    r = 7 + i
    sm[f"A{r}"].value, sm[f"B{r}"].value = lab, fml
for r in range(7, 7 + len(people)):
    for c in "AB":
        sm[f"{c}{r}"].border, sm[f"{c}{r}"].font = border, f_base
    sm[f"B{r}"].alignment = center
CK0 = 15
head_cell(sm[f"A{CK0}"], "데이터 점검")
head_cell(sm[f"B{CK0}"], "건수")
checks = [
    ("기본자료 인원", '=SUMPRODUCT(--(TRIM(기본자료표[사번]&"")<>""))'),
    ("기본자료 재직 인원", '=COUNTIF(기본자료표[현재근무 여부],"재직")'),
    ("기본자료에 없는 사번", '=SUMPRODUCT((TRIM(명단표[사번]&"")<>"")*(명단표[기본자료행]=""))'),
    ("명단 중복 사번", '=SUMPRODUCT((TRIM(명단표[사번]&"")<>"")*(COUNTIF(명단표[사번],명단표[사번])>1))'),
    ("날짜 형식 오류(생년월일/수련개시일)", '=SUMPRODUCT(--ISTEXT(명단표[생년월일]))+SUMPRODUCT(--ISTEXT(명단표[수련개시일]))'),
    ("인턴계획 입력 오류(병원/진료과)", "=SUM(인턴계획표[입력오류])"),
    ("서울·구리 15일 판정 불가", '=COUNTIF(명단표[급여지급병원(기준월)],"⚠*")'),
    ("구리 근무자명단: 명단에 없음/동명이인", f"={q('출력_구리근무자')}!$J$2"),
    ("구리 근무자명단: 날짜 오류", f"={q('출력_구리근무자')}!$J$3"),
    ("월별근무지 직접 입력 오류", '=SUMPRODUCT((월별근무지표[점검]<>"")*(월별근무지표[점검]<>"정상"))'),
    ("전월·익월 자동 추정 확인 필요(⚠, 구리 근무자)", '=COUNTIF(명단표[전월 근거],"⚠*")+COUNTIF(명단표[익월 근거],"⚠*")'),
    ("점검 필요 행(전공의명단 '점검' 열)", '=SUMPRODUCT((TRIM(명단표[사번]&"")<>"")*(명단표[점검]<>"정상"))'),
]
for i, (lab, fml) in enumerate(checks):
    r = CK0 + 1 + i
    sm[f"A{r}"].value, sm[f"B{r}"].value = lab, xl(fml)
    for c in "AB":
        sm[f"{c}{r}"].border, sm[f"{c}{r}"].font = border, f_base
    sm[f"B{r}"].alignment = center
sm.conditional_formatting.add(f"B{CK0 + 3}:B{CK0 + len(checks)}", FormulaRule(formula=[f"B{CK0 + 3}>0"], fill=PatternFill("solid", bgColor="FFC7CE")))
blocks = [("D", ["병원", "현재 근무", "급여 지급(기준월)"], "병원표[병원]",
           ["COUNTIF(명단표[현재 근무병원],{k})", "COUNTIF(명단표[급여지급병원(기준월)],{k})"]),
          ("H", ["현재 진료과", "인원"], "진료과표[근무과]", ["COUNTIF(명단표[현재 진료과],{k})"]),
          ("K", ["인턴 구분", "인원"], "인턴구분표[인턴 구분]", ["COUNTIF(명단표[현재 진료과],{k})"])]
for col, heads, keyref, counts in blocks:
    c0 = CI(col)
    for j, h in enumerate(heads):
        head_cell(sm.cell(row=6, column=c0 + j), h)
        sm.column_dimensions[CL(c0 + j)].width = 14 if j == 0 else 12
    spill(sm, f"{col}7", f"={keyref}")
    for j, cnt in enumerate(counts, 1):
        spill(sm, f"{CL(c0 + j)}7", "=" + cnt.format(k=keyref))
sm.column_dimensions["A"].width = 36
sm.column_dimensions["B"].width = 12
for c in "CGJ":
    sm.column_dimensions[c].width = 3
sm.row_dimensions[6].height = 30

# ================================================================ 출력_인턴현황
iv = wb.create_sheet("출력_인턴현황")
iv["A1"].value, iv["A1"].font = "인턴 근무 현황", f_title
iv["A2"].value, iv["A2"].font = "기준일", f_bold
iv["B2"].value, iv["B2"].number_format = "=TODAY()", "yyyy-mm-dd"
iv["A3"].value, iv["A3"].font = "계획 월 번호(자동)", f_note
iv["F3"].value, iv["F3"].font = "이번/다음 달 위치→", f_note
iv["G3"].value = '=IFERROR(MATCH(MONTH(TODAY()),$J$3:$O$3,0),"")'
iv["I3"].value = '=IFERROR(MATCH(MONTH(EDATE(TODAY(),1)),$J$3:$O$3,0),"")'
iv["J2"].value, iv["J2"].font = "급여 기준월 위치→", f_note
iv["K2"].value = '=IFERROR(MATCH(MONTH(급여기준월),$J$3:$O$3,0),"")'
for c in ("G3", "I3", "K2"):
    iv[c].font, iv[c].alignment = f_note, center
add_name(wb, "인턴위치_이번", f"{q('출력_인턴현황')}!$G$3")
add_name(wb, "인턴위치_다음", f"{q('출력_인턴현황')}!$I$3")
add_name(wb, "인턴위치_기준월", f"{q('출력_인턴현황')}!$K$2")
iv["L2"].value, iv["L2"].font = "전월/익월 위치→", f_note
iv["M2"].value = '=IFERROR(MATCH(MONTH(EDATE(급여기준월,-1)),$J$3:$O$3,0),"")'
iv["N2"].value = '=IFERROR(MATCH(MONTH(EDATE(급여기준월,1)),$J$3:$O$3,0),"")'
for c in ("M2", "N2"):
    iv[c].font, iv[c].alignment = f_note, center
add_name(wb, "인턴위치_전월", f"{q('출력_인턴현황')}!$M$2")
add_name(wb, "인턴위치_익월", f"{q('출력_인턴현황')}!$N$2")
iv_heads = ["순번", "이름", "사번", "기본자료 확인", "입사일",
            '="이번 달("&IF($G$3="","계획 외",MONTH(TODAY())&"월")&") 병원"', "이번 달 진료과",
            '="다음 달("&IF($I$3="","계획 외",MONTH(EDATE(TODAY(),1))&"월")&") 병원"', "다음 달 진료과"]
iv_heads += [f'=TRIM(INDEX(인턴계획표[#Headers],1,{3 + k})&"")' for k in range(1, 7)] + ["입력오류"]
for i, h in enumerate(iv_heads, 1):
    head_cell(iv.cell(row=5, column=i), h)
    iv.column_dimensions[CL(i)].width = 17 if 6 <= i <= 15 else 10
iv.row_dimensions[5].height = 30
for k in range(6):
    c = CL(10 + k)
    iv[f"{c}3"].value = f'=IFERROR(VALUE(SUBSTITUTE(TRIM({c}5),"월","")),"")'
    iv[f"{c}3"].font, iv[f"{c}3"].alignment = f_note, center
hosp = lambda x: f'IF({x}="","",IF(LEFT({x},1)="⚠",{x},IFERROR(LEFT({x},FIND(" ",{x})-1),{x})))'
dept = lambda x: f'IF({x}="","",IF(LEFT({x},1)="⚠","",IFERROR(MID({x},FIND(" ",{x})+1,60),"")))'
NRM = "인턴계획표[[정규화1]:[정규화6]]"
spill(iv, "A6",
      f'=LET(v_n,TRIM(인턴계획표[이름]&""),v_x,v_n<>"",'
      f'v_c,IF($G$3="","",INDEX({NRM},0,$G$3)&""),v_t,IF($I$3="","",INDEX({NRM},0,$I$3)&""),'
      f'v_ok,IF(ISNUMBER(XMATCH(인턴계획표[사번키],기본자료표[사번키])),"확인","미등록"),'
      f'v_in,IF(ISNUMBER(인턴계획표[입사일(비고)]),TEXT(인턴계획표[입사일(비고)],"yyyy-mm-dd"),인턴계획표[입사일(비고)]&""),'
      f'IF(SUM(--v_x)=0,"인턴 계획 없음",FILTER(CHOOSE({{1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,16}},'
      f'인턴계획표[순번]&"",v_n,인턴계획표[사번]&"",v_ok,v_in,{hosp("v_c")},{dept("v_c")},{hosp("v_t")},{dept("v_t")},'
      f'인턴계획표[정규화1],인턴계획표[정규화2],인턴계획표[정규화3],인턴계획표[정규화4],인턴계획표[정규화5],인턴계획표[정규화6],'
      f'인턴계획표[입력오류]),v_x)))', width=16)
END = 20005
iv.conditional_formatting.add(f"D6:D{END}", FormulaRule(formula=['$D6="미등록"'], fill=PatternFill("solid", bgColor="FFC7CE")))
iv.conditional_formatting.add(f"J6:O{END}", FormulaRule(formula=['LEFT(J6,1)="⚠"'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=FONT, color="C00000", bold=True)))
iv.conditional_formatting.add(f"J5:O{END}", FormulaRule(formula=['J$3=MONTH(TODAY())'], fill=PatternFill("solid", bgColor="FFF2CC")))
iv.freeze_panes = "D6"
# 월별 병원별 인원
iv.column_dimensions["Q"].width = 3
iv["R2"].value, iv["R2"].font = "월별 병원별 인턴 인원", f_bold
head_cell(iv["R3"], "병원")
iv["R4"].value, iv["R5"].value = "합계(배정)", "⚠ 오류"
for c in ("R4", "R5"):
    iv[c].font, iv[c].border = f_bold, border
iv.column_dimensions["R"].width = 14
spill(iv, "R6", "=병원표[병원]")
for k in range(6):
    c = CL(19 + k)
    head_cell(iv[f"{c}3"], f"={CL(10 + k)}5")
    iv[f"{c}4"].value = f'=COUNTIF(인턴계획표[정규화{k + 1}],"?*")'
    iv[f"{c}5"].value = f'=COUNTIF(인턴계획표[정규화{k + 1}],"⚠*")'
    for rr in (4, 5):
        iv[f"{c}{rr}"].font, iv[f"{c}{rr}"].border, iv[f"{c}{rr}"].alignment = f_bold, border, center
    spill(iv, f"{c}6", f'=COUNTIF(인턴계획표[정규화{k + 1}],병원표[병원]&" *")')
    iv.column_dimensions[c].width = 10

# ================================================================ 사용안내
gd = wb.create_sheet("사용안내", 0)
lines = [
    ("전공의·인턴 파견 모니터링 사용 안내 (Microsoft 365 · Excel 2021용, 인원 제한 없음)", f_title),
    ("", f_base),
    ("[시트 구분] — 탭 색: 파랑=입력, 초록=출력(자동, 입력 금지), 회색=설정", f_bold),
    ("• 입력_기본자료 / 입력_인턴근무계획 / 입력_전공의명단: 엑셀 '표'. 표 바로 아래(또는 첫 자료 행)에 붙여넣으면 표가 자동으로 늘어나고 초록 머리글 계산 열도 자동으로 채워집니다.", f_base),
    ("• 입력_구리근무자명단: 구리병원 파일의 시트를 A1에 양식 그대로 붙여넣는 시트(병합 셀이 있어 표가 아님). 2만 행·ZZ열까지 읽습니다.", f_base),
    ("• 출력_서울급여 / 출력_구리근무자 / 출력_요약 / 출력_인턴현황: 대상 수만큼 자동으로 펼쳐지는 목록(동적 배열). 셀을 지우거나 아래에 입력하지 마세요(#SPILL! 오류).", f_base),
    ("• 설정_코드표: 병원·진료과·인턴구분·약칭 목록(표)과 설정값(급여 기준월·15일 기준).", f_base),
    ("• 입력_월별근무지: 전월·익월(또는 어느 달이든) 근무병원을 사람별로 직접 입력하는 표. 자동 판정보다 우선합니다.", f_base),
    ("• 출력_기관명단: 전월·급여 기준월·익월 3개월의 전원 기관(근무병원·급여지급병원) — 시간외근무 판정 파일의 입력_기관명단에 값으로 붙여넣는 용도", f_base),
    ("", f_base),
    ("[붙여넣는 방법]", f_bold),
    ("• 기본자료: A2부터 (머리글 순서: 현재근무 여부·사번·이름·수련과목·연차·근무지·수련개시일·전공의등록번호·의사면허번호·성별·생년월일·비고). 날짜는 20260301 형식도 됩니다.", f_base),
    ("• 인턴근무계획: B3부터 자료 행(순번·이름·사번·월 6칸·비고). 학기가 바뀌면 2행의 월 머리글(9월~2월 등)도 바꿔 적으세요. 비고 2026.03 = 2026-03-01 입사.", f_base),
    ("• 구리 근무자명단: 구리병원 파일의 시트를 A1에 양식 그대로(1~2행 제목, 3행 머리글, 4행부터 자료). 날짜는 'YYYY.MM.DD' 글자도 됩니다.", f_base),
    ("• 전공의명단: 사번·입사장소·비고만 입력. 나머지는 자동.", f_base),
    ("• 예시 데이터(ㅇㅇ)는 모두 지우고 실제 자료를 넣으세요. 표의 행을 지울 때는 행 전체 삭제(표 행 삭제)를 쓰세요.", f_base),
    ("", f_base),
    ("[표기 해석]", f_bold),
    ("• 인턴 월 칸은 '지역 진료과'(예: 서울 산부 → 서울병원 산부인과, 제주 에스중앙 내과 → 제주병원 내과). 첫 단어=지역(병원약칭표), 마지막 단어=진료과(과약칭표).", f_base),
    ("• 통합(과 미지정), 핵/방, 직/피, 정/영 등은 인턴 전용 구분입니다(과약칭표에서 직피·정영 등 약칭도 인정).", f_base),
    ("• 코드표에 없는 값을 직접 입력하면 오류 팝업이 뜹니다. 붙여넣기는 팝업이 뜨지 않으므로 해당 칸이 빨갛게 표시되고 요약에 집계됩니다.", f_base),
    ("", f_base),
    ("[급여지급병원(기준월) 규칙]", f_bold),
    ("• 서울병원·구리병원 사이: 소속과 무관하게 급여 기준월에 15일 이상 근무한 병원이 그 달 전체를 지급(예: 구리 10일·서울 20일 → 서울 지급).", f_base),
    ("• 근무일: 구리 근무자명단에 있으면 그 날짜가 구리 근무일, 그 달의 나머지 날은 모두 서울 근무일입니다.", f_base),
    ("• 양쪽 모두 15일 이상이거나 모두 미만이면 ⚠ 판정 불가로 표시합니다.", f_base),
    ("• 구리 근무자명단에 없는 사람은 그 달 근무병원(인턴은 인턴근무계획의 병원, 그 외는 소속병원)이 지급합니다.", f_base),
    ("• 소속병원 = 입사장소 → (비어 있으면) 구리 근무자명단 비고 '구리독자'면 구리병원 → 기본자료 근무지 순서로 정합니다.", f_base),
    ("", f_base),
    ("[월별 근무지 직접 입력]", f_bold),
    ("• 입력_월별근무지에 월(예: 2026-10)·사번·근무병원을 넣으면 그 달은 그 값을 씁니다. 급여지급병원이 근무병원과 다를 때만 급여지급병원 칸도 입력하세요.", f_base),
    ("• 급여 기준월: 직접 입력 → 구리 근무자명단 15일 판정 → 인턴 근무계획 → 소속병원 순서. 직접 입력하면 ⚠15일 판정불가도 그 값으로 정리됩니다.", f_base),
    ("• 전월·익월: 직접 입력 → 인턴 근무계획 → 소속병원 순서(구리 근무자명단은 급여 기준월 블록만 읽음). 입력_전공의명단의 '전월/익월 근무병원·급여지급병원·근거' 열과 출력_기관명단에 나옵니다.", f_base),
    ("• 급여 기준월에 구리 근무자명단에 있는 사람의 전월·익월이 자동 추정이면 근거에 ⚠ — 실제 근무지를 직접 입력하세요.", f_base),
    ("• 출력_기관명단은 전월·기준월·익월 3개월을 한 번에 펼치므로, 시간외 판정 파일에 한 번만 복사하면 됩니다.", f_base),
    ("", f_base),
    ("[참고]", f_bold),
    ("• Microsoft 365(또는 Excel 2021 이상)에서만 정상 동작합니다. 구버전에서는 출력 시트가 오류로 보입니다.", f_base),
    ("• 개인정보(생년월일·면허번호 등)가 포함되므로 파일 암호 설정과 접근 권한 관리를 권장합니다.", f_base),
]
for i, (tx, fo) in enumerate(lines, 1):
    gd.cell(row=i, column=1, value=tx).font = fo
gd.column_dimensions["A"].width = 175
gd.sheet_view.showGridLines = False

# ================================================================ 시트 순서·탭 색, 저장
ORDER = ["사용안내", "입력_기본자료", "입력_인턴근무계획", "입력_구리근무자명단", "입력_전공의명단", "입력_월별근무지",
         "출력_서울급여", "출력_구리근무자", "출력_요약", "출력_인턴현황", "출력_기관명단", "설정_코드표"]
wb._sheets = [wb[n] for n in ORDER]
wb.active = 0
for sh in wb.worksheets:
    for pre, color in {"입력_": "2F5597", "출력_": "548235", "설정_": "7F7F7F"}.items():
        if sh.title.startswith(pre):
            sh.sheet_properties.tabColor = color
    for row in sh.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and "_xlfn" not in c.value:
                c.value = xl(c.value)
    for cfr in sh.conditional_formatting:
        for rule in cfr.rules:
            rule.formula = [xl(f) for f in rule.formula]
wb.save(OUT)

# ---------------------------------------------------------------- 동적 배열 메타데이터 (Excel이 결과를 펼치도록)
if SPILLS:
    META = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
            '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" pasteAll="1" '
            'pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" '
            'coerce="1" cellMeta="1"/></metadataTypes><futureMetadata name="XLDAPR" count="1"><bk><extLst>'
            '<ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}"><xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/>'
            '</ext></extLst></bk></futureMetadata><cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>')
    zin = zipfile.ZipFile(OUT)
    files = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    wbxml = files["xl/workbook.xml"].decode("utf-8")
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rid_target = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    rid_target.update({a: b for b, a in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wbxml):
        if name not in SPILLS:
            continue
        path = "xl/" + rid_target[rid].lstrip("/").replace("xl/", "")
        x = files[path].decode("utf-8")
        for cell in SPILLS[name]:
            x, n = re.subn(rf'<c r="{cell}"(?![^>]*cm=)', f'<c r="{cell}" cm="1"', x, count=1)
            assert n == 1, (name, cell)
        files[path] = x.encode("utf-8")
    files["xl/metadata.xml"] = META.encode("utf-8")
    ct = files["[Content_Types].xml"].decode("utf-8")
    ct = ct.replace("</Types>", '<Override PartName="/xl/metadata.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/></Types>')
    files["[Content_Types].xml"] = ct.encode("utf-8")
    rels = rels.replace("</Relationships>", '<Relationship Id="rIdMeta1" '
                        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata" '
                        'Target="metadata.xml"/></Relationships>')
    files["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in files.items():
            zout.writestr(n, data)
print("saved", OUT, "verify" if VERIFY else "", sum(len(v) for v in SPILLS.values()), "dynamic cells")
