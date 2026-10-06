"""Microsoft 365용 엑셀 생성 도우미 (엑셀 표, 동적 배열, 365 함수 접두어).

VERIFY=1 이면 동적 배열을 고정 크기 배열수식으로 기록해 LibreOffice로 계산을 검증할 수 있게 한다.
"""
import os
import re
import zipfile

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

VERIFY = os.environ.get("VERIFY") == "1"
VERIFY_ROWS = int(os.environ.get("VERIFY_ROWS", "40"))

FONT = "맑은 고딕"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_input = Font(name=FONT, size=10, color="0000FF")  # 입력값=파랑, 수식=검정
f_title = Font(name=FONT, size=14, bold=True)
f_note = Font(name=FONT, size=9, color="7F7F7F")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_head2 = PatternFill("solid", fgColor="548235")
fill_calc = PatternFill("solid", fgColor="F2F2F2")
fill_key = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

_FN = [("XLOOKUP", "_xlfn.XLOOKUP"), ("XMATCH", "_xlfn.XMATCH"), ("SEQUENCE", "_xlfn.SEQUENCE"),
       ("LET", "_xlfn.LET"), ("MINIFS", "_xlfn.MINIFS"), ("MAXIFS", "_xlfn.MAXIFS"),
       ("FILTER", "_xlfn._xlws.FILTER"), ("SORT", "_xlfn._xlws.SORT")]


def xl(f):
    """사람이 읽는 수식 → 파일 저장 형식 (365 함수 접두어, LET 변수 접두어)."""
    if not isinstance(f, str):
        return f
    f = re.sub(r"(?<![\w.])(v_[A-Za-z0-9_]+)", r"_xlpm.\1", f)
    for name, full in _FN:
        f = re.sub(rf"(?<![\w.]){name}\(", full + "(", f)
    return f


def this_row(table, f):
    """[@열] → 표[[#This Row],[열]]"""
    return re.sub(r"\[@([^\]]+)\]", lambda m: f"{table}[[#This Row],[{m.group(1)}]]", f)


class Book:
    """동적 배열 셀을 기억했다가 저장할 때 Excel 메타데이터를 붙인다."""

    def __init__(self, wb):
        self.wb = wb
        self.spills = {}

    def spill(self, ws, cell, formula, width=1, height=None):
        f = xl(formula)
        if VERIFY:
            m = re.match(r"([A-Z]+)(\d+)", cell)
            end = f"{CL(CI(m.group(1)) + width - 1)}{int(m.group(2)) + (height or VERIFY_ROWS) - 1}"
            ws[cell] = ArrayFormula(f"{cell}:{end}", f)
        else:
            ws[cell] = ArrayFormula(cell, f)
            self.spills.setdefault(ws.title, []).append(cell)

    def add_name(self, name, ref):
        dn = DefinedName(name, attr_text=xl(ref))
        try:
            self.wb.defined_names[name] = dn
        except TypeError:
            self.wb.defined_names.append(dn)

    def save(self, path):
        for sh in self.wb.worksheets:
            for row in sh.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and "_xlfn" not in c.value:
                        c.value = xl(c.value)
            for cfr in sh.conditional_formatting:
                for rule in cfr.rules:
                    rule.formula = [xl(f) for f in rule.formula]
        self.wb.save(path)
        if self.spills:
            add_dynamic_metadata(path, self.spills)


def q(sheet):
    return f"'{sheet}'"


def head_cell(c, text, fill=fill_head):
    c.value, c.font, c.fill, c.alignment, c.border = text, f_head, fill, wrap_center, border


def make_table(ws, tname, top, left_col, cols, rows, style="TableStyleLight9"):
    """cols: [(머리글, 너비, 'in'|'calc', 수식 또는 None, 숫자형식 또는 None)], rows: 입력 열 값(dict) 목록"""
    n = max(1, len(rows))
    tcols = []
    for j, (h, w, kind, fml, fmt) in enumerate(cols):
        c = ws.cell(row=top, column=left_col + j)
        head_cell(c, h, fill_head if kind == "in" else fill_head2)
        ws.column_dimensions[CL(left_col + j)].width = w
        tc = TableColumn(id=j + 1, name=h)
        if kind == "calc":
            body = xl(this_row(tname, fml))
            tc.calculatedColumnFormula = TableFormula(attr_text=body[1:] if body.startswith("=") else body)
        tcols.append(tc)
        for i in range(n):
            cell = ws.cell(row=top + 1 + i, column=left_col + j)
            if kind == "calc":
                cell.value = xl(this_row(tname, fml))
                cell.font, cell.fill = f_base, fill_calc
            else:
                v = rows[i].get(h) if i < len(rows) else None
                if v is not None:
                    cell.value = v
                cell.font = f_input
            cell.border = border
            if fmt:
                cell.number_format = fmt
    ref = f"{CL(left_col)}{top}:{CL(left_col + len(cols) - 1)}{top + n}"
    t = Table(displayName=tname, ref=ref, tableColumns=tcols)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)
    return t


def dv(ws, sqref, formula1, title, msg, kind="list"):
    d = DataValidation(type=kind, formula1=xl(formula1), allow_blank=True, showErrorMessage=True,
                       errorStyle="stop", errorTitle=title, error=msg)
    ws.add_data_validation(d)
    d.add(sqref)


_META = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
         '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
         'xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
         '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" pasteAll="1" '
         'pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" '
         'coerce="1" cellMeta="1"/></metadataTypes><futureMetadata name="XLDAPR" count="1"><bk><extLst>'
         '<ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}"><xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/>'
         '</ext></extLst></bk></futureMetadata><cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>')


def add_dynamic_metadata(path, spills):
    """동적 배열 셀에 cm="1"을 달고 metadata.xml을 추가해 Excel이 결과를 펼치게 한다."""
    zin = zipfile.ZipFile(path)
    files = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    wbxml = files["xl/workbook.xml"].decode("utf-8")
    rels = files["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rid_target = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    rid_target.update({a: b for b, a in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wbxml):
        if name not in spills:
            continue
        sheet_path = "xl/" + rid_target[rid].lstrip("/").replace("xl/", "")
        x = files[sheet_path].decode("utf-8")
        for cell in spills[name]:
            x, n = re.subn(rf'<c r="{cell}"(?![^>]*cm=)', f'<c r="{cell}" cm="1"', x, count=1)
            assert n == 1, (name, cell)
        files[sheet_path] = x.encode("utf-8")
    files["xl/metadata.xml"] = _META.encode("utf-8")
    ct = files["[Content_Types].xml"].decode("utf-8")
    files["[Content_Types].xml"] = ct.replace("</Types>", '<Override PartName="/xl/metadata.xml" ContentType='
                                              '"application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/></Types>').encode("utf-8")
    files["xl/_rels/workbook.xml.rels"] = rels.replace("</Relationships>", '<Relationship Id="rIdMeta1" Type='
                                                      '"http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata" '
                                                      'Target="metadata.xml"/></Relationships>').encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in files.items():
            zout.writestr(n, data)
