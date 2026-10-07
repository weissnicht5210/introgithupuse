"""전공의 시간외 근무 → 연속 근무 묶기 + 서울급여 해당 판정 (Excel 2019 이상 호환 — Microsoft 365에서도 동작)

입력: 시간외 근무 상세내역(전월 말일 기록 포함 가능), 시간외 근무(월 합계), 기관명단(전공의 파견 모니터링의 출력_기관명단, 전월·당월·익월),
      월별근무지(사람·월별 근무병원 직접 입력 — 기관명단보다 우선)
출력: 출력_연속근무, 출력_월경계당직(월을 넘는 당직의 반영 급여월·기관), 출력_서울급여판정(사람별 서울 반영 시간), 출력_타병원발송
입력은 엑셀 표(행 수 제한 없음), 출력은 정해진 행 수만큼 미리 수식을 넣은 목록(INDEX/MATCH, 순번 보조 열).
DATA=<pickle> → 실제 자료로 채우기(저장소에 넣지 않음)
"""
import os
import pickle

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill

from xlsx2019 import (CL, FONT, add_name, border, center, dv, f_base, f_bold, f_input, f_note, f_title, fill_aux, fill_key,
                      head_cell, make_table, q, save)

OUT = os.environ.get("OUT", "시간외근무_서울급여_판정.xlsx")
DATA = None
if os.environ.get("DATA"):
    with open(os.environ["DATA"], "rb") as fh:
        DATA = pickle.load(fh)
MAXR = 1048576
SEC = "1/86400"  # 시각 비교 여유 1초 (소수점 오차 방지)
OCN = int(os.environ.get("OCN", "3000"))   # 출력_연속근무 표시 한도(건)
OBN = int(os.environ.get("OBN", "1000"))   # 출력_월경계당직 표시 한도(건)
OPN = int(os.environ.get("OPN", "1000"))   # 출력_서울급여판정 표시 한도(명)
OSN = 300                                  # 판정 시트 옆 '합계표에 없는 사번' 표시 한도(명)
OTP = 300                                  # 출력_타병원발송 사람 표시 한도(명)
OTR = int(os.environ.get("OTR", "2000"))   # 출력_타병원발송 기록 표시 한도(건)


def rank(flag):
    """flag=1인 기록의 순번: 사번 → 시작 시각 → 입력 행 순서 (SORT 대신 COUNTIFS)."""
    f = f"상세표[{flag}],1"
    return (f'=IF([@{flag}]<>1,"",COUNTIFS({f},상세표[사번키],"<"&[@사번키])'
            f'+COUNTIFS({f},상세표[사번키],[@사번키],상세표[시작],"<"&([@시작]-{SEC}))'
            f'+COUNTIFS({f},상세표[사번키],[@사번키],상세표[시작],">="&([@시작]-{SEC}),상세표[시작],"<="&([@시작]+{SEC}),'
            f'상세표[행],"<"&[@행])+1)')


def first_flag(flag):
    """flag=1인 기록 중 사번별 첫 행이면 1 (UNIQUE 대신)."""
    return f'=IF([@{flag}]<>1,0,IF(COUNTIFS(상세표[{flag}],1,상세표[사번키],[@사번키],상세표[행],"<"&[@행])>0,0,1))'


def first_rank(first):
    """사번별 첫 행의 순번(입력 순서)."""
    return f'=IF([@{first}]<>1,"",COUNTIFS(상세표[{first}],1,상세표[행],"<="&[@행]))'


wb = Workbook()
wb.remove(wb.active)

# ================================================================ 설정
st = wb.create_sheet("설정")
head_cell(st["A1"], "설정값 (노란 칸만 수정)")
head_cell(st["B1"], "값")
st["A2"].value, st["A2"].font, st["A2"].border = "연속으로 볼 최대 공백(분)", f_bold, border
c = st["B2"]
c.value, c.font, c.fill, c.border, c.alignment = (DATA or {}).get("gap_min", 0), f_input, fill_key, border, center
add_name(wb, "연속허용분", f"{q('설정')}!$B$2")
st["A3"].value, st["A3"].font = "※ 0이면 앞 근무의 종료 시각과 다음 근무의 시작 시각이 같을 때만 이어진 근무로 봅니다. 예: 10 → 10분 이내 공백도 연속.", f_note
st["A4"].value, st["A4"].font, st["A4"].border = "대상월(급여 처리할 달, 1일)", f_bold, border
c = st["B4"]
c.value, c.number_format, c.font, c.fill, c.border, c.alignment = ((DATA or {}).get("target") or "=DATE(YEAR(TODAY()),MONTH(TODAY())-1,1)"), "yyyy-mm", f_input, fill_key, border, center
add_name(wb, "대상월", f"{q('설정')}!$B$4")
add_name(wb, "대상월키", 'TEXT(대상월,"yyyy-mm")')
st.column_dimensions["A"].width = 30
st.column_dimensions["B"].width = 12
HOLIDAYS_2026 = [
    ("2026-01-01", "신정"), ("2026-02-16", "설날 연휴"), ("2026-02-17", "설날"), ("2026-02-18", "설날 연휴"),
    ("2026-03-01", "삼일절"), ("2026-03-02", "대체공휴일(삼일절)"), ("2026-05-05", "어린이날"),
    ("2026-05-24", "부처님오신날"), ("2026-05-25", "대체공휴일(부처님오신날)"), ("2026-06-03", "전국동시지방선거"),
    ("2026-06-06", "현충일"), ("2026-08-15", "광복절"), ("2026-08-17", "대체공휴일(광복절)"),
    ("2026-09-24", "추석 연휴"), ("2026-09-25", "추석"), ("2026-09-26", "추석 연휴"), ("2026-10-03", "개천절"),
    ("2026-10-05", "대체공휴일(개천절)"), ("2026-10-09", "한글날"), ("2026-12-25", "성탄절"),
]
from datetime import date as _d
hol_rows = [{"날짜": _d.fromisoformat(d), "이름": n} for d, n in ((DATA or {}).get("holidays") or HOLIDAYS_2026)]
make_table(st, "공휴일표", 6, 1, [("날짜", 12, "in", None, "yyyy-mm-dd"), ("이름", 22, "in", None, None)], hol_rows, style="TableStyleLight1")
st["A5"].value, st["A5"].font = "공휴일 (주말·휴일 포함 여부 판단에 사용 — 표 아래에 계속 추가, 2027년은 확정 후 추가)", f_bold
add_name(wb, "공휴일목록", "공휴일표[날짜]")
st["D5"].value, st["D5"].font = "구리와 오가면 '익월초 급여 반영'으로 표시할 기관", f_bold
make_table(st, "익월초기관표", 6, 4, [("기관", 14, "in", None, None)], [{"기관": h} for h in ["인천병원", "오산병원", "창원병원", "제주병원"]], style="TableStyleLight1")
st.column_dimensions["C"].width = 3
add_name(wb, "익월초기관목록", "익월초기관표[기관]")
st["F5"].value, st["F5"].font = "병원 목록 (입력_월별근무지 드롭다운)", f_bold
make_table(st, "병원표", 6, 6, [("병원", 14, "in", None, None)],
           [{"병원": h} for h in ["서울병원", "인천병원", "구리병원", "창원병원", "제주병원", "오산병원"]], style="TableStyleLight1")
st.column_dimensions["E"].width = 3
add_name(wb, "병원목록", "병원표[병원]")

# ================================================================ 입력_시간외상세
ds = wb.create_sheet("입력_시간외상세")
ds["A1"].value, ds["A1"].font = "월별 전공의 시간외 근무 상세내역", f_title
ds["A2"].value, ds["A2"].font = "※ 원본 파일의 5행(자료 첫 행)부터 마지막 행까지 복사해 A5에 붙여넣으세요. 표가 자동으로 늘어납니다. 초록 머리글 열은 자동 계산.", f_note
ds["A3"].value, ds["A3"].font = "※ 여러 달을 이어 붙여도 됩니다(월이 바뀌어도 이어진 근무를 찾음).", f_note
d_in = ["부서", "사번", "성명", "직위", "근무일자", "시작시간", "종료시간", "06~07시", "07~18시", "18~22시", "22~06시"]
d_w = [10, 10, 8, 10, 11, 16, 16, 7, 7, 7, 7]
BANDS = ["06~07시", "07~18시", "18~22시", "22~06시"]
d_rows = (DATA or {}).get("detail_rows") or [dict(zip(d_in, r)) for r in [
    ("ㅇㅇ과", 2222222, "김ㅇㅇ", "레지던트2", "2026-08-31", "2026-08-31 20:00", "2026-09-01 08:00", 1, 1, 2, 8),
    ("ㅇㅇ과", 2222222, "김ㅇㅇ", "레지던트2", "2026-09-01", "2026-09-01 08:00", "2026-09-01 20:00", 0, 10, 2, 0),
    ("ㅇㅇ과", 2222222, "김ㅇㅇ", "레지던트2", "2026-09-04", "2026-09-04 06:00", "2026-09-04 20:00", 1, 11, 2, 0),
    ("ㅇㅇ과", 2222222, "김ㅇㅇ", "레지던트2", "2026-09-04", "2026-09-04 20:00", "2026-09-05 08:00", 1, 1, 2, 8),
    ("ㅇㅇ과", 2222222, "김ㅇㅇ", "전임의", "2026-09-24", "2026-09-24 08:00", "2026-09-24 20:00", 0, 10, 2, 0),
    ("ㅇㅇ과", 2222222, "김ㅇㅇ", "레지던트2", "2026-09-30", "2026-09-30 20:00", "2026-10-01 08:00", 1, 1, 2, 8),
    ("수련교육부", 3333333, "이ㅇㅇ", "인턴", "2026-08-31", "2026-08-31 20:00", "2026-09-01 08:00", 1, 1, 2, 8),
    ("수련교육부", 3333333, "이ㅇㅇ", "인턴", "2026-09-07", "2026-09-07 18:00", "2026-09-07 22:00", 0, 0, 4, 0),
    ("수련교육부", 3333333, "이ㅇㅇ", "인턴", "2026-09-07", "2026-09-07 07:00", "2026-09-07 18:00", 0, 11, 0, 0),
    ("수련교육부", 4444444, "박ㅇㅇ", "인턴", "2026-09-30", "2026-09-30 20:00", "2026-10-01 08:00", 1, 1, 2, 8),
    ("△△과", 5555555, "정ㅇㅇ", "레지던트3", "2026-09-10", "2026-09-10 08:00", "2026-09-10 20:00", 0, 10, 2, 0),
    ("△△과", 5555555, "정ㅇㅇ", "레지던트3", "2026-09-10", "2026-09-10 20:00", "2026-09-11 08:00", 1, 1, 2, 8),
]]


def ptime(col):
    x = f"[@{col}]"
    return (f'=IF(ISNUMBER({x}),{x},IF(TRIM({x}&"")="","",IFERROR(VALUE(TRIM({x})),'
            f'IFERROR(DATEVALUE(LEFT(TRIM({x}),10))+VALUE(MID(TRIM({x}),12,2))/24+VALUE(MID(TRIM({x}),15,2))/1440,"⚠"&{x}))))')


NEXT = 'IF([@다음블록시작]="",1E+9,[@다음블록시작])'
KEYCOND = '상세표[사번키],[@사번키]'
NEXTMIN = f'MINIFS(상세표[시작],{KEYCOND},상세표[블록시작],1,상세표[시작],">"&([@시작]+{SEC}))'
RANGE = f'상세표[시작],">="&([@시작]-{SEC}),상세표[시작],"<"&({NEXT}-{SEC})'
d_cols = [(h, w, "in", None, "0.00" if h in BANDS else None) for h, w in zip(d_in, d_w)]
d_cols += [
    ("행", 6, "calc", '=ROW()-ROW(상세표[#Headers])', "0"),
    ("사번키", 10, "calc", '=IF(TRIM([@사번]&"")="","","S|"&TRIM([@사번]&""))', None),
    ("시작", 15, "calc", ptime("시작시간"), "yyyy-mm-dd hh:mm"),
    ("종료", 15, "calc", ptime("종료시간"), "yyyy-mm-dd hh:mm"),
    ("이어짐", 7, "calc",
     f'=IF(NOT(ISNUMBER([@시작])),"",IF(COUNTIFS({KEYCOND},상세표[시작],"<"&([@시작]-{SEC}),'
     f'상세표[종료],">="&([@시작]-연속허용분/1440-{SEC}))>0,"이어짐",""))', None),
    ("블록시작", 7, "calc", '=IF(AND(ISNUMBER([@시작]),ISNUMBER([@종료]),[@이어짐]=""),1,0)', "0"),
    ("다음블록시작", 15, "calc",
     f'=IF([@블록시작]<>1,"",IF({NEXTMIN}=0,"",{NEXTMIN}))',
     "yyyy-mm-dd hh:mm"),
    ("블록종료", 15, "calc", f'=IF([@블록시작]<>1,"",MAXIFS(상세표[종료],{KEYCOND},{RANGE}))', "yyyy-mm-dd hh:mm"),
    ("연속 시간", 8, "calc", '=IF([@블록시작]<>1,"",ROUND(([@블록종료]-[@시작])*24,2))', "0.00"),
    ("기록 수", 6, "calc", f'=IF([@블록시작]<>1,"",COUNTIFS({KEYCOND},{RANGE}))', "0"),
]
for b in BANDS:
    d_cols.append((f"블록 {b}", 8, "calc", f'=IF([@블록시작]<>1,"",SUMIFS(상세표[{b}],{KEYCOND},{RANGE}))', "0.00"))
d_cols += [
    ("주말·휴일 포함", 9, "calc",
     f'=IF([@블록시작]<>1,"",IF((INT([@블록종료]-{SEC})-INT([@시작])+1)'
     f'-NETWORKDAYS(INT([@시작]),INT([@블록종료]-{SEC}),공휴일목록)>0,"포함",""))', None),
    ("시작월", 8, "calc", '=IF([@블록시작]<>1,"",TEXT([@시작],"yyyy-mm"))', None),
    ("블록종료월", 8, "calc", f'=IF([@블록시작]<>1,"",TEXT([@블록종료]-{SEC},"yyyy-mm"))', None),
    ("월 넘김 포함", 8, "calc", f'=IF([@블록시작]<>1,"",IF(COUNTIFS({KEYCOND},{RANGE},상세표[월 넘김],"월 넘김")>0,"포함",""))', None),
    # ---- 기록(당직 1건) 단위: 월을 넘는 기록의 반영 급여월·기관
    ("기록 시간외", 8, "calc", '=N([@06~07시])+N([@07~18시])+N([@18~22시])+N([@22~06시])', "0.00"),
    ("기록월", 8, "calc", '=IF(ISNUMBER([@시작]),TEXT([@시작],"yyyy-mm"),"")', None),
    ("종료월", 8, "calc", f'=IF(ISNUMBER([@종료]),TEXT([@종료]-{SEC},"yyyy-mm"),"")', None),
    ("월 넘김", 7, "calc", '=IF(AND([@기록월]<>"",[@종료월]<>"",[@기록월]<>[@종료월]),"월 넘김","")', None),
    ("기관명단", 8, "calc",
     '=IF([@사번키]="","",IF(COUNTIF(기관표[사번키],[@사번키])>0,"있음",IF(COUNTIF(월별근무지표[사번키],[@사번키])>0,"직접 입력만","없음")))', None),
    ("시작월 직접행", 7, "calc", '=IF([@기록월]="","",IFERROR(MATCH([@사번키]&"|"&[@기록월],월별근무지표[키],0),""))', "0"),
    ("종료월 직접행", 7, "calc", '=IF([@종료월]="","",IFERROR(MATCH([@사번키]&"|"&[@종료월],월별근무지표[키],0),""))', "0"),
    ("시작월 기관", 10, "calc",
     '=IF([@기록월]="","",IF([@시작월 직접행]<>"",INDEX(월별근무지표[적용 급여병원],[@시작월 직접행])&"",'
     'IFERROR(INDEX(기관표[급여지급병원(기준월)],MATCH([@사번키]&"|"&[@기록월],기관표[키],0))&"","")))', None),
    ("종료월 기관", 10, "calc",
     '=IF([@월 넘김]="",[@시작월 기관],IF([@종료월 직접행]<>"",INDEX(월별근무지표[적용 급여병원],[@종료월 직접행])&"",'
     'IFERROR(INDEX(기관표[급여지급병원(기준월)],MATCH([@사번키]&"|"&[@종료월],기관표[키],0))&"","")))', None),
    ("기관 출처", 9, "calc",
     '=IF([@기록월]="","",IF(OR([@시작월 직접행]<>"",AND([@월 넘김]<>"",[@종료월 직접행]<>"")),"직접 입력","기관명단"))', None),
    ("월경계 판정", 24, "calc",
     '=IF([@월 넘김]="","",IF([@기관명단]="없음","기관명단에 없음(다른 병원으로 발송)",'
     'IF(OR([@시작월 기관]="",[@종료월 기관]=""),"기관 정보 없음(해당 월 기관명단·직접 입력 필요)",'
     'IF(OR(LEFT([@시작월 기관],1)="⚠",LEFT([@종료월 기관],1)="⚠"),"기관 판정 불가(15일 판정불가)",'
     'IF([@시작월 기관]=[@종료월 기관],"같은 기관",'
     'IF(AND(OR([@시작월 기관]="서울병원",[@시작월 기관]="구리병원"),OR([@종료월 기관]="서울병원",[@종료월 기관]="구리병원")),"서울↔구리: 15일 이상 근무 기관",'
     'IF(OR(AND([@시작월 기관]="구리병원",COUNTIF(익월초기관목록,[@종료월 기관])>0),AND([@종료월 기관]="구리병원",COUNTIF(익월초기관목록,[@시작월 기관])>0)),"익월초 급여 반영",'
     '"규정 확인 필요")))))))', None),
    ("반영월", 8, "calc", '=IF([@기록월]="","",IF([@월경계 판정]="익월초 급여 반영",[@종료월],[@기록월]))', None),
    ("반영 기관", 22, "calc",
     '=IF([@기록월]="","",IF([@기관명단]="없음","기관명단에 없음(다른 병원으로 발송)",'
     'IF([@월경계 판정]="익월초 급여 반영","확인 필요("&[@시작월 기관]&"→"&[@종료월 기관]&")",[@시작월 기관])))', None),
    # ---- 출력 목록용 표시·순번 (수정 금지)
    ("연속표시", 6, "calc", '=IF(AND([@블록시작]=1,OR([@시작월]=대상월키,[@블록종료월]=대상월키)),1,0)', "0"),
    ("연속순번", 6, "calc", rank("연속표시"), "0"),
    ("경계표시", 6, "calc", '=IF(AND([@월 넘김]="월 넘김",OR([@기록월]=대상월키,[@종료월]=대상월키)),1,0)', "0"),
    ("경계순번", 6, "calc", rank("경계표시"), "0"),
    ("발송표시", 6, "calc", '=IF(AND([@기관명단]="없음",OR([@기록월]=대상월키,[@반영월]=대상월키)),1,0)', "0"),
    ("발송순번", 6, "calc", rank("발송표시"), "0"),
    ("발송첫행", 6, "calc", first_flag("발송표시"), "0"),
    ("발송인원순번", 6, "calc", first_rank("발송첫행"), "0"),
    ("합계없음표시", 6, "calc", '=IF(AND([@반영월]=대상월키,[@사번키]<>"",COUNTIF(합계표[사번키],[@사번키])=0),1,0)', "0"),
    ("합계없음첫행", 6, "calc", first_flag("합계없음표시"), "0"),
    ("합계없음순번", 6, "calc", first_rank("합계없음첫행"), "0"),
]
make_table(ds, "상세표", 4, 1, d_cols, d_rows)
ds.freeze_panes = "D5"

# ================================================================ 입력_시간외합계
ss = wb.create_sheet("입력_시간외합계")
ss["A1"].value, ss["A1"].font = "월별 전공의 시간외 근무 (합계)", f_title
ss["A2"].value, ss["A2"].font = "※ 원본 파일의 5행부터 마지막 행까지 복사해 A5에 붙여넣으세요. 시간 칸이 '20.00' 같은 글자여도 숫자로 읽습니다.", f_note
s_in = ["부서", "사번", "성명", "직위", "직위(변환)"] + BANDS
s_rows = (DATA or {}).get("summary_rows") or [dict(zip(s_in, r)) for r in [
    ("ㅇㅇ과", 2222222, "김ㅇㅇ", "레지던트2", "레지던트2", "3.00", "23.00", "8.00", "16.00"),
    ("ㅇㅇ과", 2222222, "김ㅇㅇ", "전임의", "전임의", "0.00", "10.00", "2.00", "0.00"),
    ("수련교육부", 3333333, "이ㅇㅇ", "인턴", "인턴", "0.00", "11.00", "4.00", "0.00"),
    ("수련교육부", 4444444, "박ㅇㅇ", "인턴", "인턴", "1.00", "1.00", "2.00", "8.00"),
    ("△△과", 5555555, "정ㅇㅇ", "레지던트3", "레지던트3", "1.00", "11.00", "4.00", "8.00"),
]]
s_cols = [(h, w, "in", None, None) for h, w in zip(s_in, [10, 10, 8, 10, 10, 8, 8, 8, 8])]
s_cols.append(("사번키", 10, "calc", '=IF(TRIM([@사번]&"")="","","S|"&TRIM([@사번]&""))', None))
s_cols.append(("시스템 합계", 9, "calc", "=" + "+".join(f'IFERROR(VALUE(TRIM([@{b}]&"")),0)' for b in BANDS), "0.00"))
# 같은 사번이 직위만 달리해 여러 줄이면 한 사람으로: 첫 직위와 마지막 직위를 한 칸에 표시(배열 함수 없이)
s_cols.append(("사번내순서", 7, "calc", '=IF([@사번키]="","",COUNTIF(INDEX(합계표[사번키],1):[@사번키],[@사번키]))', "0"))
s_cols.append(("순서키", 12, "calc", '=IF([@사번키]="","",[@사번키]&"#"&[@사번내순서])', None))
_first = 'INDEX(합계표[직위(변환)],MATCH([@사번키]&"#1",합계표[순서키],0))&""'
_last = 'INDEX(합계표[직위(변환)],MATCH([@사번키]&"#"&COUNTIF(합계표[사번키],[@사번키]),합계표[순서키],0))&""'
s_cols.append(("직위(통합)", 18, "calc",
               f'=IF([@사번키]="","",IF({_first}={_last},{_first},{_first}&", "&{_last}))', None))
s_cols.append(("판정순번", 7, "calc", '=IF([@사번내순서]=1,COUNTIF(INDEX(합계표[사번내순서],1):[@사번내순서],1),"")', "0"))
make_table(ss, "합계표", 4, 1, s_cols, s_rows)
ss.freeze_panes = "D5"

# ================================================================ 입력_기관명단 (월별 근무·급여 기관)
sl = wb.create_sheet("입력_기관명단")
l_in = ["기준월", "사번", "이름", "수련과목", "연차", "소속병원", "기준월 근무병원", "급여지급병원(기준월)", "급여 판정 근거"]
l_rows = (DATA or {}).get("org_rows") or [dict(zip(l_in, r)) for r in [
    ("2026-08", "2222222", "김ㅇㅇ", "ㅇㅇ과", "레지던트2", "서울병원", "서울병원", "서울병원", "예시"),
    ("2026-09", "2222222", "김ㅇㅇ", "ㅇㅇ과", "레지던트2", "서울병원", "서울병원", "서울병원", "예시"),
    ("2026-10", "2222222", "김ㅇㅇ", "ㅇㅇ과", "레지던트2", "서울병원", "서울병원", "구리병원", "예시"),
    ("2026-09", "3333333", "이ㅇㅇ", "인턴", "인턴", "서울병원", "구리병원", "구리병원", "예시"),
    ("2026-10", "3333333", "이ㅇㅇ", "인턴", "인턴", "서울병원", "구리병원", "구리병원", "예시"),
    ("2026-08", "4444444", "박ㅇㅇ", "인턴", "인턴", "구리병원", "구리병원", "구리병원", "예시"),
    ("2026-09", "4444444", "박ㅇㅇ", "인턴", "인턴", "구리병원", "구리병원", "구리병원", "예시"),
    ("2026-10", "4444444", "박ㅇㅇ", "인턴", "인턴", "구리병원", "인천병원", "인천병원", "예시"),
]]
l_cols = [(h, w, "in", None, "@" if h in ("사번", "기준월") else None) for h, w in zip(l_in, [9, 10, 8, 12, 9, 10, 11, 13, 34])]
l_cols += [
    ("월키", 8, "calc", '=IF(ISNUMBER([@기준월]),TEXT([@기준월],"yyyy-mm"),LEFT(TRIM([@기준월]&""),7))', None),
    ("사번키", 10, "calc", '=IF(TRIM([@사번]&"")="","","S|"&TRIM([@사번]&""))', None),
    ("키", 16, "calc", '=IF([@사번키]="","",[@사번키]&"|"&[@월키])', None),
]
make_table(sl, "기관표", 1, 1, l_cols, l_rows)
sl.freeze_panes = "C2"

# ================================================================ 입력_월별근무지 (근무지 직접 입력 — 기관명단보다 우선)
dw = wb.create_sheet("입력_월별근무지")
dw["A1"].value, dw["A1"].font = "월별 근무지 직접 입력 (입력_기관명단보다 우선)", f_title
dw["A2"].value, dw["A2"].font = ("※ 그 달 기관을 직접 정할 사람만 한 줄씩: 월(예: 2026-10)·사번·근무병원. 급여지급병원이 근무병원과 다를 때만 급여지급병원도 입력. "
                                 "판정에는 '적용 급여병원'이 쓰입니다. 모니터링 파일(365판) 입력_월별근무지와 열 순서가 같아 A~E열을 그대로 복사해 와도 됩니다. "
                                 "같은 사번·월이 두 줄이면 위쪽 줄이 적용됩니다."), f_note
dw_rows = DATA.get("direct_rows", []) if DATA else [
    {"월": _d(2026, 8, 1), "사번": "3333333", "근무병원": "제주병원", "비고": "예시: 8월 기관명단 대신 직접 입력"},
]
dw_cols = [("월", 9, "in", None, "yyyy-mm"), ("사번", 11, "in", None, "@"), ("근무병원", 11, "in", None, None),
           ("급여지급병원", 12, "in", None, None), ("비고", 28, "in", None, None),
           ("월키", 9, "calc",
            '=IF([@월]&""="","",IFERROR(IF(ISNUMBER([@월]),IF([@월]<10000,TEXT(DATE(INT([@월]),ROUND(MOD([@월],1)*100,0),1),"yyyy-mm"),'
            'TEXT([@월],"yyyy-mm")),TEXT(DATEVALUE(SUBSTITUTE(SUBSTITUTE(LEFT(TRIM([@월]),7),".","-"),"/","-")&"-01"),"yyyy-mm")),"⚠형식 오류"))', None),
           ("사번키", 11, "calc", '=IF(TRIM([@사번]&"")="","","S|"&TRIM([@사번]&""))', None),
           ("키", 16, "calc", '=IF(LEFT([@월키],1)="⚠","",IF(OR([@사번키]="",[@월키]=""),"",[@사번키]&"|"&[@월키]))', None),
           ("이름(확인)", 9, "calc",
            '=IF([@사번키]="","",IFERROR(INDEX(기관표[이름],MATCH([@사번키],기관표[사번키],0))&"",'
            'IFERROR(INDEX(상세표[성명],MATCH([@사번키],상세표[사번키],0))&"","(자료에 없음)")))', None),
           ("적용 급여병원", 12, "calc", '=IF([@근무병원]="","",IF([@급여지급병원]="",[@근무병원],[@급여지급병원]))', None),
           ("점검", 26, "calc",
            '=IF(AND(TRIM([@사번]&"")="",[@월]&""=""),"",IF(OR(TRIM([@사번]&"")="",[@월]&""=""),"월·사번 모두 입력",'
            'IF(LEFT([@월키],1)="⚠","월 형식 확인(예: 2026-10)",IF([@근무병원]="","근무병원 입력",'
            'IF(COUNTIF(월별근무지표[키],[@키])>1,"같은 사번·월 중복(위쪽 줄 적용)","정상")))))', None)]
make_table(dw, "월별근무지표", 3, 1, dw_cols, dw_rows)
dw.freeze_panes = "A4"
for c in ("C", "D"):
    dv(dw, f"{c}4:{c}{MAXR}", "=병원목록", "등록되지 않은 병원", "병원 목록에 없습니다. 새 병원이면 '설정' 시트의 병원 목록 표에 먼저 추가하세요.")
dw.conditional_formatting.add(f"K4:K{MAXR}", FormulaRule(formula=['AND($K4<>"",$K4<>"정상")'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))

# ================================================================ 출력 목록 공통
RED = Font(name=FONT, size=10, bold=True, color="C00000")


def fill_list(ws, r0, n, c0, hcol, rank_ref, exprs, empty_msg=None, empty_cond=None):
    """r0행부터 n행: hcol(보조 열)에 순번 k의 행 번호(MATCH), c0열부터 exprs({j}=행 번호 셀)."""
    for k in range(1, n + 1):
        r = r0 + k - 1
        j = f"${hcol}{r}"
        ws[f"{hcol}{r}"].value, ws[f"{hcol}{r}"].font = f'=IFERROR(MATCH({k},{rank_ref},0),"")', f_note
        for i, e in enumerate(exprs):
            ws.cell(row=r, column=c0 + i).value = f'=IF({j}="","",{e.format(j=j)})'
    if empty_msg:
        j = f"${hcol}{r0}"
        ws.cell(row=r0, column=c0).value = f'=IF({j}="",IF({empty_cond},"{empty_msg}",""),{exprs[0].format(j=j)})'


def aux_head(ws, cell, text="(행)"):
    head_cell(ws[cell], text, fill_aux)


D = lambda col: f"INDEX(상세표[{col}],{{j}})"
fmtdt = lambda c: f'IF(ISNUMBER({c}),TEXT({c},"yyyy-mm-dd hh:mm"),{c}&"")'

# ================================================================ 출력_연속근무
oc = wb.create_sheet("출력_연속근무")
oc["A1"].value, oc["A1"].font = "연속 근무 (이어진 시간외 근무를 한 줄로)", f_title
oc["A2"].value, oc["A2"].font = (f"※ 자동 목록(입력 금지, 최대 {OCN}건). 대상월에 걸친 연속 근무만 사번·시작 시각 순으로 빈 줄 없이 보여줍니다. 전월 말일에 시작한 당직도 포함됩니다. "
                                 "월 넘김 표시가 있으면 출력_월경계당직에서 반영 급여월·기관을 확인하세요."), f_note
for lc, lt, vc, vf in [("A3", "연속 근무 건수", "B3", '=SUM(상세표[연속표시])'),
                       ("C3", "2개 이상 기록을 묶은 건수", "D3", '=COUNTIFS(상세표[연속표시],1,상세표[기록 수],">1")'),
                       ("E3", "주말·휴일 포함", "F3", '=COUNTIFS(상세표[연속표시],1,상세표[주말·휴일 포함],"포함")'),
                       ("G3", "36시간 초과", "H3", '=COUNTIFS(상세표[연속표시],1,상세표[연속 시간],">36")'),
                       ("I3", "시각 형식 오류", "J3", '=COUNTIF(상세표[시작],"⚠*")+COUNTIF(상세표[종료],"⚠*")')]:
    oc[lc].value, oc[lc].font, oc[lc].alignment = lt, f_bold, Alignment(horizontal="right")
    oc[vc].value, oc[vc].font, oc[vc].border, oc[vc].alignment = vf, f_bold, border, center
oc["K3"].value, oc["K3"].font = f'=IF(B3>{OCN},"⚠ 표시 한도({OCN}건) 초과","")', RED
oc.conditional_formatting.add("H3", FormulaRule(formula=["H3>0"], fill=PatternFill("solid", bgColor="FFC7CE")))
oc.conditional_formatting.add("J3", FormulaRule(formula=["J3>0"], fill=PatternFill("solid", bgColor="FFC7CE")))
oc_heads = [("사번", 10), ("성명", 8), ("부서", 10), ("직위", 10), ("시작", 16), ("종료", 16), ("시작 요일", 6),
            ("연속 시간", 8), ("기록 수", 6)] + [(b, 8) for b in BANDS] + [("주말·휴일 포함", 9), ("월 넘김 포함", 9)]
for i, (h, w) in enumerate(oc_heads, 1):
    head_cell(oc.cell(row=5, column=i), h)
    oc.column_dimensions[CL(i)].width = w
oc.row_dimensions[5].height = 30
oc_cols = [D("사번") + '&""', D("성명") + '&""', D("부서") + '&""', D("직위") + '&""', fmtdt(D("시작")), fmtdt(D("블록종료")),
           f'CHOOSE(WEEKDAY({D("시작")}),"일","월","화","수","목","금","토")',
           D("연속 시간"), D("기록 수")] + [D(f"블록 {b}") for b in BANDS] + [D("주말·휴일 포함") + '&""', D("월 넘김 포함") + '&""']
OCH = CL(len(oc_cols) + 1)
aux_head(oc, f"{OCH}5")
fill_list(oc, 6, OCN, 1, OCH, "상세표[연속순번]", oc_cols, "연속 근무 없음", "SUM(상세표[연속표시])=0")
OCE = 5 + OCN
oc.conditional_formatting.add(f"A6:{CL(len(oc_cols))}{OCE}", FormulaRule(formula=['AND(ISNUMBER($H6),$H6>36)'], fill=PatternFill("solid", bgColor="FFC7CE")))
oc.conditional_formatting.add(f"A6:{CL(len(oc_cols))}{OCE}", FormulaRule(formula=['$O6="포함"'], fill=PatternFill("solid", bgColor="DDEBF7")))
oc.conditional_formatting.add(f"A6:{CL(len(oc_cols))}{OCE}", FormulaRule(formula=['$N6="포함"'], fill=PatternFill("solid", bgColor="FCE4D6")))
oc.freeze_panes = "C6"

# ================================================================ 출력_월경계당직
ob = wb.create_sheet("출력_월경계당직")
ob["A1"].value, ob["A1"].font = "월을 넘는 당직 (전월 말일→이번 달, 이번 달 말일→익월)", f_title
ob["A2"].value, ob["A2"].font = ("※ 같은 기관이면 시작한 달에 반영. 서울↔구리는 시작한 달의 15일 이상 근무 기관(그 달 급여지급병원)에 반영. "
                                 "구리↔인천·오산·창원·제주는 '익월초 급여 반영'(종료한 달 급여). 그 밖의 조합은 '규정 확인 필요'."), f_note
for lc, lt, vc, vf in [("A3", "월 넘김 당직", "B3", '=SUM(상세표[경계표시])'),
                       ("C3", "전월→이번 달", "D3", '=COUNTIFS(상세표[경계표시],1,상세표[종료월],대상월키)'),
                       ("E3", "이번 달→익월", "F3", '=COUNTIFS(상세표[경계표시],1,상세표[기록월],대상월키)'),
                       ("G3", "익월초 급여 반영", "H3", '=COUNTIFS(상세표[경계표시],1,상세표[월경계 판정],"익월초 급여 반영")'),
                       ("I3", "기관 정보 없음·확인 필요", "J3",
                        '=COUNTIFS(상세표[경계표시],1,상세표[월경계 판정],"기관*")+COUNTIFS(상세표[경계표시],1,상세표[월경계 판정],"규정 확인 필요")')]:
    ob[lc].value, ob[lc].font, ob[lc].alignment = lt, f_bold, Alignment(horizontal="right")
    ob[vc].value, ob[vc].font, ob[vc].border, ob[vc].alignment = vf, f_bold, border, center
ob["K3"].value, ob["K3"].font = f'=IF(B3>{OBN},"⚠ 표시 한도({OBN}건) 초과","")', RED
ob.conditional_formatting.add("J3", FormulaRule(formula=["J3>0"], fill=PatternFill("solid", bgColor="FFC7CE")))
ob_heads = [("사번", 10), ("성명", 8), ("구분", 12), ("시작", 16), ("종료", 16), ("시간외 시간", 8),
            ("시작월 기관", 11), ("종료월 기관", 11), ("판정", 26), ("반영 급여월", 9), ("반영 기관", 26), ("기관 출처", 9)]
for i, (h, w) in enumerate(ob_heads, 1):
    head_cell(ob.cell(row=5, column=i), h)
    ob.column_dimensions[CL(i)].width = w
ob.row_dimensions[5].height = 30
ob_cols = [D("사번") + '&""', D("성명") + '&""', f'IF({D("종료월")}=대상월키,"전월→이번 달","이번 달→익월")',
           fmtdt(D("시작")), fmtdt(D("종료")), D("기록 시간외"), D("시작월 기관") + '&""', D("종료월 기관") + '&""',
           D("월경계 판정") + '&""', D("반영월") + '&""', D("반영 기관") + '&""',
           D("기관 출처") + '&""']
aux_head(ob, "M5")
fill_list(ob, 6, OBN, 1, "M", "상세표[경계순번]", ob_cols, "월을 넘는 당직 없음", "SUM(상세표[경계표시])=0")
OBE = 5 + OBN
ob.conditional_formatting.add(f"A6:L{OBE}", FormulaRule(formula=['$I6="익월초 급여 반영"'], fill=PatternFill("solid", bgColor="FFF2CC")))
ob.conditional_formatting.add(f"A6:L{OBE}", FormulaRule(formula=['OR(LEFT($I6,2)="기관",$I6="규정 확인 필요")'], fill=PatternFill("solid", bgColor="FFC7CE")))
ob.conditional_formatting.add(f"L6:L{OBE}", FormulaRule(formula=['$L6="직접 입력"'], font=Font(name=FONT, bold=True, color="2F5597")))
ob.freeze_panes = "C6"

# ================================================================ 출력_서울급여판정
NOORG = "기관명단에 없음(다른 병원으로 발송)"
NOMON = "대상월 기관명단 없음(추가 필요)"
op = wb.create_sheet("출력_서울급여판정")
op["A1"].value, op["A1"].font = "시간외 근무 서울급여 판정 (대상월)", f_title
op["A2"].value, op["A2"].font = (f"※ 자동 결과(입력 금지, 최대 {OPN}명). 같은 사번은 직위가 달라도 한 사람으로 합칩니다. 서울급여 = 입력_기관명단에서 대상월 급여지급병원이 서울병원. "
                                 "기관명단에 사번이 아예 없으면 '기관명단에 없음(다른 병원으로 발송)' — 출력_타병원발송에 따로 모았습니다."), f_note
OPE = 5 + OPN
for lc, lt, vc, vf in [("A3", "대상월", "B3", "=대상월키"),
                       ("C3", "시간외 인원(사번 기준)", "D3", '=COUNT(합계표[판정순번])'),
                       ("E3", "서울급여 해당", "F3", f'=COUNTIF($L$6:$L${OPE},"해당")'),
                       ("G3", "상세·합계 불일치", "H3", f'=COUNTIF($G$6:$G${OPE},"불일치*")'),
                       ("I3", "기관명단에 없음(타 병원 발송)", "J3", f'=COUNTIF($L$6:$L${OPE},"기관명단에 없음*")'),
                       ("K3", "대상월 기관명단 없음", "L3", f'=COUNTIF($L$6:$L${OPE},"대상월 기관명단*")')]:
    op[lc].value, op[lc].font, op[lc].alignment = lt, f_bold, Alignment(horizontal="right")
    op[vc].value, op[vc].font, op[vc].border, op[vc].alignment = vf, f_bold, border, center
op["M3"].value, op["M3"].font = "직접 입력 적용", f_bold
op["M3"].alignment = Alignment(horizontal="right")
op["N3"].value, op["N3"].font, op["N3"].border, op["N3"].alignment = f'=COUNTIF($R$6:$R${OPE},"직접 입력")', f_bold, border, center
op["O3"].value, op["O3"].font = f'=IF(D3>{OPN},"⚠ 표시 한도({OPN}명) 초과","")', RED
for c in ("H3", "L3"):
    op.conditional_formatting.add(c, FormulaRule(formula=[f"{c}>0"], fill=PatternFill("solid", bgColor="FFC7CE")))
op.conditional_formatting.add("J3", FormulaRule(formula=["J3>0"], fill=PatternFill("solid", bgColor="F8CBAD")))
op_heads = [("사번", 10), ("성명", 8), ("부서", 10), ("직위", 16), ("시스템 합계", 8), ("상세 합계(대상월 기록)", 9),
            ("상세 대조", 13), ("이월(−) 익월초", 8), ("전입(+) 전월 말일", 8), ("대상월 반영 시간", 9),
            ("대상월 급여지급병원", 11), ("서울급여", 20)] + [(f"서울 {b}", 8) for b in BANDS] + [("서울 반영 합계", 9), ("기관 출처", 9)]
for i, (h, w) in enumerate(op_heads, 1):
    head_cell(op.cell(row=5, column=i), h)
    op.column_dimensions[CL(i)].width = w
op.row_dimensions[5].height = 42
U = "INDEX(합계표[사번키],{j})"
sm = lambda col, *crit: f'SUMIFS(상세표[{col}],상세표[사번키],{U}' + "".join("," + c for c in crit) + ")"
C_REC_M, C_REF_M = "상세표[기록월],대상월키", "상세표[반영월],대상월키"
C_REC_NOT, C_REF_NOT = '상세표[기록월],"<>"&대상월키', '상세표[반영월],"<>"&대상월키'
C_SEOUL = '상세표[반영 기관],"서울병원"'
HAS_ORG = f"COUNTIF(기관표[사번키],{U})+COUNTIF(월별근무지표[사번키],{U})>0"
DROW = f'MATCH({U}&"|"&대상월키,월별근무지표[키],0)'
r_ = "{r}"
op_cols = [f'INDEX(합계표[{c}],{{j}})&""' for c in ("사번", "성명", "부서", "직위(통합)")] + [
    f"SUMIFS(합계표[시스템 합계],합계표[사번키],{U})",
    sm("기록 시간외", C_REC_M),
    'IF(ABS($E{r}-$F{r})<0.01,"일치","불일치(차이 "&TEXT($F{r}-$E{r},"0.00")&")")',
    sm("기록 시간외", C_REC_M, C_REF_NOT), sm("기록 시간외", C_REC_NOT, C_REF_M), sm("기록 시간외", C_REF_M),
    f'IFERROR(INDEX(월별근무지표[적용 급여병원],{DROW})&"",IF({HAS_ORG},IFERROR(INDEX(기관표[급여지급병원(기준월)],'
    f'MATCH({U}&"|"&대상월키,기관표[키],0))&"","—"),"—"))',
    f'IF(NOT({HAS_ORG}),"{NOORG}",IF($K{{r}}="—","{NOMON}",IF($K{{r}}="서울병원","해당","비해당")))',
] + [sm(b, C_REF_M, C_SEOUL) for b in BANDS] + ["SUM($M{r}:$P{r})",
       f'IF(ISNUMBER({DROW}),"직접 입력",IF($K{{r}}="—","—","기관명단"))']
aux_head(op, "S5", "(합계 행)")
for k in range(1, OPN + 1):
    r = 5 + k
    j = f"$S{r}"
    op[f"S{r}"].value, op[f"S{r}"].font = f'=IFERROR(MATCH({k},합계표[판정순번],0),"")', f_note
    for i, e in enumerate(op_cols, 1):
        op.cell(row=r, column=i).value = f'=IF({j}="","",{e.replace("{j}", j).replace("{r}", str(r))})'
op["A6"].value = f'=IF($S6="",IF(COUNT(합계표[판정순번])=0,"시간외 합계 없음",""),INDEX(합계표[사번],$S6)&"")'
op.conditional_formatting.add(f"A6:R{OPE}", FormulaRule(formula=['$L6="비해당"'], fill=PatternFill("solid", bgColor="EDEDED"), font=Font(name=FONT, color="7F7F7F")))
op.conditional_formatting.add(f"A6:R{OPE}", FormulaRule(formula=['LEFT($L6,7)="기관명단에 없"'], fill=PatternFill("solid", bgColor="F8CBAD"), font=Font(name=FONT, bold=True, color="833C0B")))
op.conditional_formatting.add(f"L6:L{OPE}", FormulaRule(formula=['LEFT($L6,6)="대상월 기관"'], fill=PatternFill("solid", bgColor="FFC7CE")))
op.conditional_formatting.add(f"G6:G{OPE}", FormulaRule(formula=['LEFT($G6,3)="불일치"'], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
op.conditional_formatting.add(f"H6:I{OPE}", FormulaRule(formula=['AND(ISNUMBER(H6),H6<>0)'], fill=PatternFill("solid", bgColor="FFF2CC"), font=Font(name=FONT, bold=True)))
op.conditional_formatting.add(f"L6:L{OPE}", FormulaRule(formula=['$L6="해당"'], fill=PatternFill("solid", bgColor="C6E0B4"), font=Font(name=FONT, bold=True)))
op.conditional_formatting.add(f"R6:R{OPE}", FormulaRule(formula=['$R6="직접 입력"'], font=Font(name=FONT, bold=True, color="2F5597")))
op.freeze_panes = "C6"
op.column_dimensions["S"].width = 6
op.column_dimensions["T"].width = 3
op["U4"].value, op["U4"].font = "대상월 반영 기록이 있는데 합계표에 없는 사번 (전월 말일 당직만 있는 사람 등)", f_bold
for i, (h, w) in enumerate([("사번", 10), ("성명", 8), ("대상월 반영 시간", 10)]):
    head_cell(op.cell(row=5, column=21 + i), h)
    op.column_dimensions[CL(21 + i)].width = w
aux_head(op, "X5", "(상세 행)")
fill_list(op, 6, OSN, 21, "X", "상세표[합계없음순번]",
          [D("사번") + '&""', D("성명") + '&""', f'SUMIFS(상세표[기록 시간외],상세표[사번키],{D("사번키")},상세표[반영월],대상월키)'],
          "없음", "SUM(상세표[합계없음첫행])=0")

# ================================================================ 출력_타병원발송
ot = wb.create_sheet("출력_타병원발송")
ot["A1"].value, ot["A1"].font = "기관명단에 없는 시간외 근무자 — 다른 병원으로 발송", f_title
ot["A2"].value, ot["A2"].font = ("※ 자동 목록(입력 금지). 대상월에 기록(또는 반영)이 있는데 입력_기관명단에 사번이 한 번도 없는 사람입니다(같은 사번은 직위가 달라도 한 사람). "
                                 f"왼쪽은 사람별 합계(최대 {OTP}명), 오른쪽은 발송할 상세 기록(최대 {OTR}건)입니다. 우리 기관 사람인데 여기 나오면 기관명단을 확인하세요."), f_note
for lc, lt, vc, vf in [("A3", "발송 대상 인원", "B3", '=SUM(상세표[발송첫행])'), ("C3", "발송 기록 수", "D3", '=SUM(상세표[발송표시])')]:
    ot[lc].value, ot[lc].font, ot[lc].alignment = lt, f_bold, Alignment(horizontal="right")
    ot[vc].value, ot[vc].font, ot[vc].border, ot[vc].alignment = vf, f_bold, border, center
ot["E3"].value, ot["E3"].font = f'=IF(OR(B3>{OTP},D3>{OTR}),"⚠ 표시 한도 초과","")', RED
ot_heads = [("사번", 10), ("성명", 8), ("부서", 10), ("직위", 16), ("대상월 기록 수", 8), ("대상월 상세 시간", 9), ("시스템 합계", 9)]
for i, (h, w) in enumerate(ot_heads, 1):
    head_cell(ot.cell(row=5, column=i), h, PatternFill("solid", fgColor="833C0B"))
    ot.column_dimensions[CL(i)].width = w
ot.row_dimensions[5].height = 30
UK = D("사번키")
POS = f'IFERROR(INDEX(합계표[직위(통합)],MATCH({UK},합계표[사번키],0))&"","")'
ot_cols = [D("사번") + '&""', D("성명") + '&""', D("부서") + '&""', f'IF({POS}="",{D("직위")}&"",{POS})',
           f"COUNTIFS(상세표[사번키],{UK},상세표[기록월],대상월키)",
           f"SUMIFS(상세표[기록 시간외],상세표[사번키],{UK},상세표[기록월],대상월키)",
           f"SUMIFS(합계표[시스템 합계],합계표[사번키],{UK})"]
aux_head(ot, "H5")
fill_list(ot, 6, OTP, 1, "H", "상세표[발송인원순번]", ot_cols, "발송 대상 없음", "SUM(상세표[발송첫행])=0")
ot.column_dimensions["H"].width = 6
ot.column_dimensions["I"].width = 3
rec_heads = [("사번", 10), ("성명", 8), ("직위", 10), ("근무일자", 11), ("시작", 16), ("종료", 16)] + [(b, 7) for b in BANDS] + [("시간외 합계", 8)]
for i, (h, w) in enumerate(rec_heads):
    head_cell(ot.cell(row=5, column=10 + i), h, PatternFill("solid", fgColor="833C0B"))
    ot.column_dimensions[CL(10 + i)].width = w
rec_cols = [D("사번") + '&""', D("성명") + '&""', D("직위") + '&""', f'IF(ISNUMBER({D("근무일자")}),TEXT({D("근무일자")},"yyyy-mm-dd"),{D("근무일자")}&"")',
            fmtdt(D("시작")), fmtdt(D("종료"))] + [D(b) for b in BANDS] + [D("기록 시간외")]
RH = CL(10 + len(rec_heads))
aux_head(ot, f"{RH}5")
fill_list(ot, 6, OTR, 10, RH, "상세표[발송순번]", rec_cols)
ot.freeze_panes = "A6"

# ================================================================ 사용안내
gd = wb.create_sheet("사용안내", 0)
lines = [
    ("시간외 근무 — 연속 근무 · 월 넘김 당직 · 서울급여 판정 (Excel 2019 이상 · Microsoft 365 공통)", f_title),
    ("", f_base),
    ("[매달 하는 일]", f_bold),
    ("1) 설정 B4에 대상월(급여 처리할 달) 입력 — 기본값은 지난달", f_base),
    ("2) 입력_시간외상세: 대상월 상세내역 + 전월 상세내역 중 말일 기록(또는 전월 전체)을 A5부터 붙여넣기. 익월 상세가 이미 있으면 함께 넣어도 됩니다.", f_base),
    ("3) 입력_시간외합계: 대상월 '월별 전공의 시간외 근무'의 5행부터 끝까지를 A5에 붙여넣기", f_base),
    ("4) 입력_기관명단: 전공의 파견 모니터링 파일의 출력_기관명단을 값으로 복사해 붙여넣기 — 전월·대상월·익월 3개월치 (365판은 한 번에 3개월, 2019판은 급여 기준월을 바꿔 가며 복사)", f_base),
    ("5) 입력_월별근무지(필요할 때만): 기관명단에 없거나 다르게 처리할 사람의 그 달 근무병원(·급여지급병원)을 직접 입력 — 기관명단보다 우선합니다.", f_base),
    ("", f_base),
    ("[월별 근무지 직접 입력]", f_bold),
    ("• 월(예: 2026-10)·사번·근무병원을 넣고, 급여지급병원이 근무병원과 다를 때만 급여지급병원도 넣습니다. 판정에는 '적용 급여병원'(급여지급병원, 비면 근무병원)이 쓰입니다.", f_base),
    ("• 적용되는 곳: 월을 넘는 당직의 시작월·종료월 기관, 대상월 급여지급병원(서울급여 해당 여부). 출력_월경계당직·출력_서울급여판정의 '기관 출처' 열에 '직접 입력'으로 표시됩니다.", f_base),
    ("• 기관명단에 사번이 없어도 직접 입력이 있으면 '기관명단에 없음(다른 병원으로 발송)'으로 보내지 않습니다.", f_base),
    ("• 모니터링 파일(365판)의 입력_월별근무지와 열 순서(월·사번·근무병원·급여지급병원·비고)가 같아 그대로 복사해 붙여넣을 수 있습니다.", f_base),
    ("", f_base),
    ("[출력 시트]", f_bold),
    ("• 출력_연속근무: 대상월에 걸친 이어진 근무를 한 줄로(전월 말일에 시작한 당직 포함), 사번·시작 순, 빈 줄 없음. 주말·휴일, 36시간 초과, 월 넘김 표시", f_base),
    ("• 출력_월경계당직: 전월 말일→대상월, 대상월 말일→익월로 넘어가는 당직마다 시작월·종료월 기관, 판정, 반영 급여월·기관", f_base),
    ("• 출력_서울급여판정: 사번별(직위가 달라도 같은 사번이면 한 사람) 시스템 합계·상세 대조, 이월(−)/전입(+), 대상월 반영 시간, 대상월 급여지급병원, 서울 반영 시간(시간대별)", f_base),
    ("• 출력_타병원발송: 기관명단에 사번이 없는 시간외 근무자(다른 병원으로 발송할 대상)와 그 사람들의 대상월 상세 기록", f_base),
    ("", f_base),
    ("[월을 넘는 당직 규칙] — 기록(당직 1건) 단위로 판단", f_bold),
    ("• 시작한 달과 끝난 달의 기관(입력_월별근무지 직접 입력 → 입력_기관명단의 급여지급병원)이 같으면: 시작한 달 급여에 반영", f_base),
    ("• 서울↔구리: 시작한 달의 15일 이상 근무 기관(그 달 급여지급병원)에 반영", f_base),
    ("• 구리↔인천·오산·창원·제주(설정의 '익월초 기관' 표): '익월초 급여 반영' — 끝난 달(익월) 급여로 넘기고 반영 기관은 '확인 필요'로 표시", f_base),
    ("• 그 밖의 조합(예: 서울↔인천)은 '규정 확인 필요', 기관명단·직접 입력에 해당 월이 없으면 '기관 정보 없음', 사번이 기관명단·직접 입력에 아예 없으면 '기관명단에 없음(다른 병원으로 발송)'", f_base),
    ("• 예: 08-31 20:00~09-01 08:00 제주→구리 → 9월(익월초) 급여 반영 / 09-30 20:00~10-01 08:00 서울→구리 → 9월 15일 이상 근무 기관(9월 급여지급병원)", f_base),
    ("", f_base),
    ("[연속 근무 판단]", f_bold),
    ("• 같은 사번에서 시작 시각이 다른 근무의 종료 시각과 같거나 겹치면 이어진 근무입니다(설정의 허용 공백(분) 이내 포함). 행 순서와 관계없이 시작 시각으로 찾습니다.", f_base),
    ("• 주말·휴일 포함: 토·일 또는 설정의 공휴일에 하루라도 걸치면 '포함'. 36시간 초과 연속 근무는 빨간색(전공의법 연속수련 한도 참고).", f_base),
    ("", f_base),
    ("[참고]", f_bold),
    ("• 시스템 합계와 비교하는 '상세 합계'는 시작 시각이 대상월인 기록의 합입니다(전월 말일 기록은 제외).", f_base),
    ("• Excel 2019·2021·Microsoft 365에서 같은 결과가 나옵니다(FILTER·XLOOKUP 같은 365 전용 함수를 쓰지 않음). 2016 이하는 확인하지 않았습니다.", f_base),
    ("• 입력 표는 행 수 제한이 없고, 출력 목록은 미리 수식을 넣어 둔 행까지만 보입니다: "
     f"연속근무 {OCN}건, 월경계당직 {OBN}건, 서울급여판정 {OPN}명, 타병원발송 {OTP}명·{OTR}건. 넘으면 시트 위쪽에 ⚠ 경고가 뜹니다.", f_base),
    ("• 출력 시트의 셀은 지우거나 덮어쓰지 마세요. 회색 머리글 '(행)' 열은 목록을 만드는 보조 계산입니다. 입력 표의 초록 머리글 열(…순번·…표시 포함)도 수정 금지.", f_base),
    ("• 연속근무·월경계·발송 기록은 사번 → 시작 시각 순, 사람별 목록은 입력 순서입니다.", f_base),
    ("• 예시 데이터(ㅇㅇ)는 지우고 실제 자료를 넣으세요. 개인정보가 포함되므로 파일 암호와 접근 권한 관리를 권장합니다.", f_base),
]
for i, (tx, fo) in enumerate(lines, 1):
    gd.cell(row=i, column=1, value=tx).font = fo
gd.column_dimensions["A"].width = 170
gd.sheet_view.showGridLines = False

ORDER = ["사용안내", "입력_시간외상세", "입력_시간외합계", "입력_기관명단", "입력_월별근무지", "출력_연속근무", "출력_월경계당직", "출력_서울급여판정", "출력_타병원발송", "설정"]
wb._sheets = [wb[n] for n in ORDER]
wb.active = 0
for sh in wb.worksheets:
    for pre, color in {"입력_": "2F5597", "출력_": "548235", "설정": "7F7F7F"}.items():
        if sh.title.startswith(pre):
            sh.sheet_properties.tabColor = color
save(wb, OUT)
print("saved", OUT)
