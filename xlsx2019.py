"""Excel 2019 호환 엑셀 생성 도우미 (엑셀 표, 구조적 참조, 고전 함수만 사용).

Excel 2019에 없는 함수(FILTER·SORT·UNIQUE·SEQUENCE·XLOOKUP·XMATCH·LET)와 동적 배열(펼침)은 쓰지 않는다.
여러 값을 한 번에 다루는 계산은 기존 배열수식(Ctrl+Shift+Enter 방식, 한 셀)으로만 기록한다.
"""
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo

FONT = "맑은 고딕"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_input = Font(name=FONT, size=10, color="0000FF")  # 입력값=파랑, 수식=검정
f_title = Font(name=FONT, size=14, bold=True)
f_note = Font(name=FONT, size=9, color="7F7F7F")
fill_head = PatternFill("solid", fgColor="1F3864")
fill_head2 = PatternFill("solid", fgColor="548235")
fill_calc = PatternFill("solid", fgColor="F2F2F2")
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_aux = PatternFill("solid", fgColor="808080")  # 보조 계산 열 머리글
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Excel 2019에 있지만 파일에는 _xlfn. 접두어로 저장해야 하는 함수
_FN = ["MINIFS", "MAXIFS", "TEXTJOIN", "IFS", "CONCAT", "SWITCH"]
# Excel 2019에 없는 함수 — 생성 결과에 남아 있으면 저장을 막는다
FORBIDDEN = re.compile(r"(?<![\w.])(XLOOKUP|XMATCH|FILTER|SORT|SORTBY|UNIQUE|SEQUENCE|RANDARRAY|LET|LAMBDA|ANCHORARRAY)\(|_xlpm|_xlws")


def xl(f):
    """사람이 읽는 수식 → 파일 저장 형식 (2019 함수 중 접두어가 필요한 것만)."""
    if not isinstance(f, str):
        return f
    for name in _FN:
        f = re.sub(rf"(?<![\w.]){name}\(", f"_xlfn.{name}(", f)
    return f


def this_row(table, f):
    """[@열] → 표[[#This Row],[열]]"""
    return re.sub(r"\[@([^\]]+)\]", lambda m: f"{table}[[#This Row],[{m.group(1)}]]", f)


def cse(ws, cell, formula):
    """한 셀짜리 기존 배열수식(Ctrl+Shift+Enter) — Excel 2019에서 그대로 계산된다."""
    ws[cell] = ArrayFormula(cell, xl(formula))


def add_name(wb, name, ref):
    dn = DefinedName(name, attr_text=xl(ref))
    try:
        wb.defined_names[name] = dn
    except TypeError:
        wb.defined_names.append(dn)


def check_2019(wb):
    """Excel 2019에 없는 함수가 남아 있으면 오류."""
    bad = []
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                v = c.value.text if isinstance(c.value, ArrayFormula) else c.value
                if isinstance(v, str) and v.startswith("=") and FORBIDDEN.search(v):
                    bad.append(f"{sh.title}!{c.coordinate}")
        for t in sh.tables.values():
            for tc in t.tableColumns:
                if tc.calculatedColumnFormula is not None and FORBIDDEN.search(tc.calculatedColumnFormula.attr_text or ""):
                    bad.append(f"{t.displayName}[{tc.name}]")
        for cfr in sh.conditional_formatting:
            for rule in cfr.rules:
                if any(FORBIDDEN.search(f) for f in rule.formula):
                    bad.append(f"{sh.title} 조건부서식 {cfr.sqref}")
    for dn in (wb.defined_names.values() if hasattr(wb.defined_names, "values") else wb.defined_names):
        if FORBIDDEN.search(dn.attr_text or ""):
            bad.append(f"이름 {dn.name}")
    assert not bad, "Excel 2019에 없는 함수: " + ", ".join(bad[:20])


def save(wb, path):
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and "_xlfn" not in c.value:
                    c.value = xl(c.value)
        for cfr in sh.conditional_formatting:
            for rule in cfr.rules:
                rule.formula = [xl(f) for f in rule.formula]
    check_2019(wb)
    wb.save(path)


def q(sheet):
    return f"'{sheet}'"


def head_cell(c, text, fill=fill_head):
    fill = fill or fill_head
    c.value, c.font, c.fill, c.alignment, c.border = text, f_head, fill, wrap_center, border


def make_table(ws, tname, top, left_col, cols, rows, style="TableStyleLight9"):
    """cols: [(머리글, 너비, 'in'|'calc', 수식 또는 None, 숫자형식 또는 None)], rows: 입력 열 값(dict) 목록"""
    n = max(1, len(rows))
    tcols = []
    for j, (h, w, kind, fml, fmt) in enumerate(cols):
        c = ws.cell(row=top, column=left_col + j)
        head_cell(c, h, fill_head if kind == "in" else fill_head2)
        ws.column_dimensions[CL(left_col + j)].width = w
        tc = TableColumn(id=j + 1, name=h)
        if kind == "calc":
            body = xl(this_row(tname, fml))
            tc.calculatedColumnFormula = TableFormula(attr_text=body[1:] if body.startswith("=") else body)
        tcols.append(tc)
        for i in range(n):
            cell = ws.cell(row=top + 1 + i, column=left_col + j)
            if kind == "calc":
                cell.value = xl(this_row(tname, fml))
                cell.font, cell.fill = f_base, fill_calc
            else:
                v = rows[i].get(h) if i < len(rows) else None
                if v is not None:
                    cell.value = v
                cell.font = f_input
            cell.border = border
            if fmt:
                cell.number_format = fmt
    ref = f"{CL(left_col)}{top}:{CL(left_col + len(cols) - 1)}{top + n}"
    t = Table(displayName=tname, ref=ref, tableColumns=tcols)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    t._initialise_columns = lambda: None  # 직접 만든 열 정의 유지
    ws.add_table(t)
    return t


def dv(ws, sqref, formula1, title, msg, kind="list"):
    d = DataValidation(type=kind, formula1=xl(formula1), allow_blank=True, showErrorMessage=True,
                       errorStyle="stop", errorTitle=title, error=msg)
    ws.add_data_validation(d)
    d.add(sqref)
