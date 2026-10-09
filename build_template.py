"""필수교육 이수확인 엑셀 템플릿 생성 (붙여넣기 → 자동 계산)."""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

EMP_SRC, LRN_SRC, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
EMP_N = 2000   # 직원 데이터 행수 (2~2001행)
LRN_N = 5000   # 학습현황 데이터 행수 (3~5002행)
OUT_N = 500    # 추출 목록 행수
E1, E2 = 2, EMP_N + 1
S1, S2 = 3, LRN_N + 2

F = "맑은 고딕"
base = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
white = Font(name=F, size=10, bold=True, color="FFFFFF")
title = Font(name=F, size=14, bold=True)
blue = Font(name=F, size=10, color="0000FF")
hdr_fill = PatternFill("solid", fgColor="1F4E78")
paste_fill = PatternFill("solid", fgColor="FFF2CC")
calc_fill = PatternFill("solid", fgColor="E2EFDA")
yellow = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
DATE = "yyyy-mm-dd"
DATE0 = "yyyy-mm-dd;;"

wb = openpyxl.Workbook()
ws_set = wb.active
ws_set.title = "설정"
ws_emp = wb.create_sheet("직원내역")
ws_lrn = wb.create_sheet("학습현황")
ws_calc = wb.create_sheet("계산")
ws_sms = wb.create_sheet("문자발송대상")
ws_out = wb.create_sheet("온라인제외")


def header(ws, row, names, fill=hdr_fill, font=white):
    for i, n in enumerate(names, 1):
        c = ws.cell(row, i, n)
        c.font, c.fill, c.border = font, fill, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ---------- 원본 양식 복사 (헤더 + 예시 데이터) ----------
src_e = openpyxl.load_workbook(EMP_SRC).active
src_l = openpyxl.load_workbook(LRN_SRC).active
emp_rows = [r for r in src_e.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
lrn_rows = [r for r in src_l.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]

header(ws_emp, 1, emp_rows[0], fill=paste_fill, font=bold)
for r, row in enumerate(emp_rows[1:], 2):
    for c, v in enumerate(row, 1):
        ws_emp.cell(r, c, v).font = base
ws_emp.freeze_panes = "A2"

ws_lrn["A1"] = lrn_rows[0][0]
ws_lrn["A1"].font = bold
header(ws_lrn, 2, lrn_rows[1], fill=paste_fill, font=bold)
for r, row in enumerate(lrn_rows[2:], 3):
    for c, v in enumerate(row, 1):
        ws_lrn.cell(r, c, v).font = base
ws_lrn.freeze_panes = "A3"

# 학습현황 보조열 (X~AC): 붙여넣기 영역(A~V) 오른쪽
LH = ["사번키", "연결사번", "과정구분", "시작(날짜)", "종료(날짜)", "수료(날짜)"]
for i, n in enumerate(LH):
    c = ws_lrn.cell(2, 24 + i, n)
    c.font, c.fill, c.border = white, hdr_fill, box
ws_lrn["X1"] = "※ X~AC열은 자동계산 (수정 금지)"
ws_lrn["X1"].font = Font(name=F, size=9, italic=True, color="808080")


def dconv(ref):
    return f'IF(ISNUMBER({ref}),{ref},IFERROR(DATEVALUE({ref}),""))'


EMP_ID = f"직원내역!$C${E1}:$C${E2}"
EMP_PH = f"직원내역!$I${E1}:$I${E2}"
for r in range(S1, S2 + 1):
    ws_lrn[f"X{r}"] = f'=IF(G{r}="","",SUBSTITUTE(LOWER(TRIM(G{r})),"hy",""))'
    ws_lrn[f"Y{r}"] = (
        f'=IF(X{r}="","",IF(COUNTIF({EMP_ID},X{r})>0,X{r},'
        f'IFERROR(INDEX({EMP_ID},MATCH(SUBSTITUTE(H{r},"-",""),INDEX(SUBSTITUTE({EMP_PH},"-",""),0),0))&"",X{r})))'
    )
    ws_lrn[f"Z{r}"] = (
        f'=IF(O{r}="","",IF(ISNUMBER(SEARCH("필수과정1",O{r})),1,'
        f'IF(ISNUMBER(SEARCH("필수과정2",O{r})),2,0)))'
    )
    # 날짜 보조열은 숫자만 갖도록(없으면 0, 0은 서식으로 숨김)
    for col, srccol in (("AA", "C"), ("AB", "D"), ("AC", "E")):
        ws_lrn[f"{col}{r}"] = f'=IF(ISNUMBER({srccol}{r}),{srccol}{r},IFERROR(DATEVALUE({srccol}{r}),0))'
        ws_lrn[f"{col}{r}"].number_format = DATE0
    for col in ("X", "Y", "Z", "AA", "AB", "AC"):
        ws_lrn[f"{col}{r}"].font = base
        ws_lrn[f"{col}{r}"].fill = calc_fill

# ---------- 설정 ----------
ws_set.column_dimensions["A"].width = 26
ws_set.column_dimensions["B"].width = 26
ws_set.column_dimensions["D"].width = 90
ws_set["A1"] = "신규입사자 필수교육 이수 확인"
ws_set["A1"].font = title
rows = [
    ("기준일(산출일)", "=TODAY()", "노란 칸: 매월 초 산출일. 비워두면 안 되며, 특정 날짜로 고정하려면 2026-11-02 처럼 직접 입력"),
    ("대상 입사 시작일", "=DATE(YEAR(B3),MONTH(B3)-3,1)", "기준월 1일 기준 3개월 전 1일"),
    ("대상 입사 종료일", "=DATE(YEAR(B3),MONTH(B3),0)", "기준월 직전 달 말일"),
    ("대상 입사기간", '=TEXT(B4,"yyyy-mm-dd")&" ~ "&TEXT(B5,"yyyy-mm-dd")', ""),
]
for i, (k, v, note) in enumerate(rows, 3):
    ws_set[f"A{i}"], ws_set[f"B{i}"], ws_set[f"D{i}"] = k, v, note
    ws_set[f"A{i}"].font = bold
    ws_set[f"B{i}"].font = base
    ws_set[f"D{i}"].font = Font(name=F, size=9, color="595959")
    if i in (3, 4, 5):
        ws_set[f"B{i}"].number_format = DATE
ws_set["B3"].fill = yellow
ws_set["B3"].font = blue

C = f"계산!$A${E1}:$A${E2}"
summary = [
    ("요약", None),
    ("대상자(재직·모자발령, 대상기간 입사)", f'=COUNTIF(계산!$J${E1}:$J${E2},"대상")'),
    ("  이수", f'=COUNTIFS(계산!$J${E1}:$J${E2},"대상",계산!$Q${E1}:$Q${E2},"이수")'),
    ("  미이수(문자발송)", f'=COUNTIF(계산!$U${E1}:$U${E2},"Y")'),
    ("온라인 제외 대상(퇴직)", f'=COUNTIF(계산!$J${E1}:$J${E2},"온라인제외(퇴직)")'),
    ("상태 확인필요", f'=COUNTIF(계산!$J${E1}:$J${E2},"확인필요")'),
    ("대상기간 외 입사(재직·모자발령)", f'=COUNTIF(계산!$J${E1}:$J${E2},"대상기간 외")'),
]
for i, (k, v) in enumerate(summary, 8):
    ws_set[f"A{i}"] = k
    ws_set[f"A{i}"].font = bold
    if v:
        ws_set[f"B{i}"] = v
        ws_set[f"B{i}"].font = base

guide = [
    "사용 방법 (매월 초)",
    "1. [직원내역] 시트: 2행부터 기존 데이터를 지우고(행 삭제 말고 내용 지우기), 다운받은 직원내역 A2부터 붙여넣기 (1행 헤더는 그대로)",
    "2. [학습현황] 시트: A3:V 범위 기존 데이터를 지우고, 다운받은 학습현황의 3행부터 A3에 붙여넣기 (X~AC열 수식은 건드리지 않기)",
    "3. [설정] B3 기준일 확인 → [문자발송대상], [온라인제외] 시트 결과 확인",
    "",
    "판정 규칙",
    "· 상태 '재직', '모자발령' = 이수 대상 / '퇴직' = 온라인 사이트 제외 대상 / 그 외 = 확인필요",
    "· 입사일은 '최종입사일' 기준, 이수기한 = 입사일 + 3개월",
    "· 과정명은 연도와 무관하게 '필수과정1', '필수과정2' 글자로 구분",
    "· 두 과정 모두 있는 경우: 둘 다 수료일이 있어야 이수 / 시작일=빠른 날짜, 종료일·수료일=늦은 날짜",
    "· 한 과정만 있는 경우(예: 필수과정2만): 그 과정만으로 이수·날짜 판단",
    "· 학습현황 매칭: 로그인아이디에서 'hy'를 뺀 값 = 사번, 일치하지 않으면 연락처로 매칭",
    f"· 처리 한도: 직원 {EMP_N}명, 학습현황 {LRN_N}행 (초과 시 수식 행 복사해 늘리기)",
]
for i, t in enumerate(guide, 17):
    ws_set[f"A{i}"] = t
    ws_set[f"A{i}"].font = bold if t in ("사용 방법 (매월 초)", "판정 규칙") else base

# ---------- 계산 ----------
CH = ["사번", "성명", "소속부서", "근무부서", "직종", "전화번호", "상태", "최종입사일", "이수기한",
      "구분", "과정1 건수", "과정1 수료", "과정2 건수", "과정2 수료", "필수과정1", "필수과정2",
      "이수여부", "시작일", "종료일", "수료일", "문자대상", "문자순번", "퇴직순번", "학습현황 등록"]
header(ws_calc, 1, CH)
L_ID = f"학습현황!$Y${S1}:$Y${S2}"
L_CO = f"학습현황!$Z${S1}:$Z${S2}"
L_ST = f"학습현황!$AA${S1}:$AA${S2}"
L_EN = f"학습현황!$AB${S1}:$AB${S2}"
L_CE = f"학습현황!$AC${S1}:$AC${S2}"
for r in range(E1, E2 + 1):
    e = lambda col: f"직원내역!{col}{r}"
    f = {
        "A": f'=IF({e("C")}="","",{e("C")}&"")',
        "B": f'=IF($A{r}="","",{e("D")})',
        "C": f'=IF($A{r}="","",{e("E")})',
        "D": f'=IF($A{r}="","",{e("F")})',
        "E": f'=IF($A{r}="","",{e("L")})',
        "F": f'=IF($A{r}="","",{e("I")})',
        "G": f'=IF($A{r}="","",TRIM({e("R")}))',
        "H": f'=IF($A{r}="","",{dconv(e("V"))})',
        "I": f'=IF(OR($A{r}="",H{r}=""),"",EDATE(H{r},3))',
        "J": (f'=IF($A{r}="","",IF(G{r}="퇴직","온라인제외(퇴직)",IF(OR(G{r}="재직",G{r}="모자발령"),'
              f'IF(AND(H{r}<>"",H{r}>=설정!$B$4,H{r}<=설정!$B$5),"대상","대상기간 외"),"확인필요")))'),
        "K": f'=IF($A{r}="","",COUNTIFS({L_ID},$A{r},{L_CO},1))',
        "L": f'=IF($A{r}="","",COUNTIFS({L_ID},$A{r},{L_CO},1,{L_CE},">0"))',
        "M": f'=IF($A{r}="","",COUNTIFS({L_ID},$A{r},{L_CO},2))',
        "N": f'=IF($A{r}="","",COUNTIFS({L_ID},$A{r},{L_CO},2,{L_CE},">0"))',
        "O": f'=IF($A{r}="","",IF(K{r}=0,IF(M{r}=0,"미신청","해당없음"),IF(L{r}=K{r},"수료","수강중")))',
        "P": f'=IF($A{r}="","",IF(M{r}=0,IF(K{r}=0,"미신청","해당없음"),IF(N{r}=M{r},"수료","수강중")))',
        "Q": f'=IF($A{r}="","",IF(AND(K{r}+M{r}>0,L{r}=K{r},N{r}=M{r}),"이수","미이수"))',
        "R": (f'=IF(Q{r}<>"이수","",IF(SUMPRODUCT(({L_ID}=$A{r})*({L_CO}>0)*({L_ST}>0))=0,"",'
              f'SUMPRODUCT(MIN({L_ST}+(1-({L_ID}=$A{r})*({L_CO}>0)*({L_ST}>0))*99999))))'),
        "S": f'=IF(Q{r}<>"이수","",SUMPRODUCT(MAX(({L_ID}=$A{r})*({L_CO}>0)*{L_EN})))',
        "T": f'=IF(Q{r}<>"이수","",SUMPRODUCT(MAX(({L_ID}=$A{r})*({L_CO}>0)*{L_CE})))',
        "U": f'=IF(AND(J{r}="대상",Q{r}="미이수"),"Y","")',
        "V": f'=IF(U{r}="Y",COUNTIF(U${E1}:U{r},"Y"),"")',
        "W": f'=IF(J{r}="온라인제외(퇴직)",COUNTIF(J${E1}:J{r},"온라인제외(퇴직)"),"")',
        "X": f'=IF($A{r}="","",IF(K{r}+M{r}>0,"등록","미등록"))',
    }
    for col, v in f.items():
        c = ws_calc[f"{col}{r}"]
        c.value, c.font = v, base
        if col in ("H", "I", "R", "S", "T"):
            c.number_format = DATE0
ws_calc.freeze_panes = "C2"

# ---------- 추출 목록 ----------
def extract(ws, title_text, key_col, cols, note):
    ws["A1"] = title_text
    ws["A1"].font = title
    ws["A2"] = '="대상 입사기간: "&설정!B6&"   (기준일 "&TEXT(설정!B3,"yyyy-mm-dd")&")"'
    ws["A2"].font = bold
    ws["A3"] = note
    ws["A3"].font = Font(name=F, size=9, color="595959")
    header(ws, 4, ["No."] + [c[0] for c in cols])
    for i in range(1, OUT_N + 1):
        r = i + 4
        ws[f"A{r}"] = f'=IF(COUNTIF(계산!${key_col}${E1}:${key_col}${E2},{i})=0,"",{i})'
        ws[f"A{r}"].font = base
        for j, (_, src) in enumerate(cols, 2):
            c = ws.cell(r, j, f'=IF($A{r}="","",INDEX(계산!${src}${E1}:${src}${E2},MATCH($A{r},계산!${key_col}${E1}:${key_col}${E2},0)))')
            c.font = base
            if src in ("H", "I", "R", "S", "T"):
                c.number_format = DATE
    ws.freeze_panes = "B5"


extract(ws_sms, "필수교육 미이수자 (문자발송 대상)", "V",
        [("사번", "A"), ("성명", "B"), ("전화번호", "F"), ("소속부서", "C"), ("근무부서", "D"),
         ("직종", "E"), ("상태", "G"), ("최종입사일", "H"), ("이수기한", "I"),
         ("필수과정1", "O"), ("필수과정2", "P")],
        "재직·모자발령 중 대상 입사기간 입사자로서 미이수인 인원")
extract(ws_out, "온라인 사이트 제외 대상 (퇴직)", "W",
        [("사번", "A"), ("성명", "B"), ("전화번호", "F"), ("소속부서", "C"), ("근무부서", "D"),
         ("최종입사일", "H"), ("학습현황 등록", "X"), ("필수과정1", "O"), ("필수과정2", "P")],
        "'학습현황 등록'이 '등록'인 인원은 온라인 사이트에서 삭제 필요")

# 열 너비
for ws, widths in [
    (ws_calc, [10, 9, 10, 11, 9, 14, 9, 11, 11, 15, 8, 8, 8, 8, 9, 9, 8, 11, 11, 11, 7, 7, 7, 10]),
    (ws_sms, [6, 10, 9, 14, 10, 11, 9, 9, 11, 11, 10, 10]),
    (ws_out, [6, 10, 9, 14, 10, 11, 11, 11, 10, 10]),
]:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
for i in range(1, 23):
    ws_lrn.column_dimensions[L(i)].width = 12
ws_lrn.column_dimensions["O"].width = 26
for col in ("X", "Y", "Z", "AA", "AB", "AC"):
    ws_lrn.column_dimensions[col].width = 11
for i in range(1, 40):
    ws_emp.column_dimensions[L(i)].width = 11

ws_set["B3"].comment = Comment("매월 산출일. 기본값 =TODAY()", "작성")
for ws in (ws_set, ws_sms, ws_out):
    ws.sheet_view.showGridLines = False
ws_emp.sheet_properties.tabColor = "FFC000"
ws_lrn.sheet_properties.tabColor = "FFC000"
ws_sms.sheet_properties.tabColor = "C00000"
ws_out.sheet_properties.tabColor = "7F7F7F"
wb.save(OUT)
print("saved", OUT)
